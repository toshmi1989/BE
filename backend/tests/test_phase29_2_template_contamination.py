"""Phase 29.2 — Template semantic contamination protection."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document

from app.core.config import get_settings
from app.core.db import Base, configure_engine, get_engine
from app.core import db as db_module
from app.domain.docx_contamination import (
    clear_product_specific_contamination,
    scan_docx_contamination,
    scrub_contamination_xml_package,
)
from app.domain.docx_profile import get_template_profile
from app.domain.exceptions import ValidationError
from app.domain.protocol_template_registry import study_is_bosutinib_sample
from app.domain.template_contamination import (
    CONTAMINATION_BLOCKS,
    CONTAMINATION_FINGERPRINTS,
    contamination_preflight,
    is_product_specific_text,
    scan_text_blob_for_contamination,
    text_contamination_hits,
)
from app.domain.decision_store import clear_decision_store
from app.domain.research_evidence_store import clear_research_evidence_store
from app.domain.sample_size_store import reset_sample_size_store
from app.domain.statistics_store import reset_statistics_store
from app.domain.study_input_store import clear_store as clear_study_input
from app.domain.study_workspace import put_protocol_draft_version, reset_workspace_store
from app.domain.workspace_protocol import generate_docx_artifact
import app.models  # noqa: F401


@pytest.fixture(autouse=True)
def _clean(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "sqlite+pysqlite:///:memory:")
    monkeypatch.setenv("AI_ENABLED", "false")
    monkeypatch.setenv("AUTH_REQUIRED", "false")
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    get_settings.cache_clear()
    configure_engine("sqlite+pysqlite:///:memory:")
    engine = get_engine()
    Base.metadata.create_all(bind=engine)
    reset_workspace_store()
    reset_sample_size_store()
    reset_statistics_store()
    clear_research_evidence_store()
    clear_decision_store()
    clear_study_input()
    yield
    reset_workspace_store()
    Base.metadata.drop_all(bind=engine)
    get_settings.cache_clear()


def test_template_product_specific_blocks():
    assert CONTAMINATION_BLOCKS
    assert any(b.category == "PHARMACOLOGY" for b in CONTAMINATION_BLOCKS)
    assert any(b.category == "CHEMICAL_FORMULA" for b in CONTAMINATION_BLOCKS)
    assert CONTAMINATION_FINGERPRINTS, "fingerprint baseline missing — run docs/_phase29_2_audit_template.py"
    assert len(CONTAMINATION_FINGERPRINTS) >= 20


def test_no_bosutinib_semantic_contamination(tmp_path: Path):
    tpl = get_template_profile().template_path()
    out = tmp_path / "t.docx"
    shutil.copy2(tpl, out)
    doc = Document(str(out))
    ctx = {"product": {"trade_name": "Upadacitinib", "inn": "upadacitinib"}}
    assert not study_is_bosutinib_sample(ctx["product"])
    clear = clear_product_specific_contamination(doc, ctx)
    assert clear["cleared"] > 0
    doc.save(str(out))
    scrub_contamination_xml_package(out, ctx)
    scan = scan_docx_contamination(out, ctx)
    assert scan["contaminated"] is False, scan


def test_no_cml_contamination():
    blob = "Исследуемый препарат предназначен для лечения хронического миелоидного лейкоза (ХМЛ)."
    hits = text_contamination_hits(blob)
    assert any(h["category"] == "DISEASE_CML" for h in hits)
    assert is_product_specific_text(blob) is True
    clean = scan_text_blob_for_contamination(
        "Исследуемый препарат: upadacitinib. Дизайн: перекрёстное исследование."
    )
    assert clean["contaminated"] is False


def test_no_unrelated_pharmacology():
    chem = "Брутто-формула: C26H29Cl2N5O3. Молекулярная масса: 530,45 г/моль."
    assert is_product_specific_text(chem) is True
    mech = "Бозутиниб — ингибитор тирозинкиназ Bcr-Abl и Src."
    assert is_product_specific_text(mech) is True


def test_no_unresolved_product_specific_placeholder():
    ctx = {"product": {"trade_name": "Upadacitinib", "inn": "upadacitinib"}, "structured_facts": {}}
    draft = contamination_preflight(ctx, mode="DRAFT")
    assert draft["ok"] is True
    assert draft["code"] == "CRITICAL_TEMPLATE_CONTAMINATION"
    final = contamination_preflight(ctx, mode="FINAL")
    assert final["ok"] is False
    assert any(
        u.get("reason") == "FINAL_REQUIRES_VERIFIED_EVIDENCE" for u in final["unmanaged_blocks"]
    )


def test_draft_not_blocked_when_fingerprint_baseline_missing(monkeypatch):
    """Missing audit fingerprint inventory must not block DRAFT export."""
    monkeypatch.setattr(
        "app.domain.template_contamination.CONTAMINATION_FINGERPRINTS",
        frozenset(),
    )
    ctx = {"product": {"trade_name": "Upadacitinib", "inn": "upadacitinib"}, "structured_facts": {}}
    draft = contamination_preflight(ctx, mode="DRAFT")
    assert draft["ok"] is True
    assert draft.get("fingerprint_baseline_missing") is True
    final = contamination_preflight(ctx, mode="FINAL")
    assert final["ok"] is False
    assert any(u.get("reason") == "MISSING_FINGERPRINT_BASELINE" for u in final["unmanaged_blocks"])


def test_regression_stale_block_blocks_generation(monkeypatch):
    """Known unmanaged product-specific block must block DOCX generation."""
    put_protocol_draft_version(
        "UPDCB-02-BE-2026",
        status="DRAFT",
        created_by="p292",
        based_on={"snapshot": None, "decisions": []},
    )
    monkeypatch.setattr(
        "app.domain.workspace_protocol.build_preflight",
        lambda _sid: {"critical_blockers": [], "warnings": [], "can_generate_docx": True},
    )
    monkeypatch.setattr(
        "app.domain.workspace_protocol.detect_stale_protocol_dependencies",
        lambda *_a, **_k: {"stale": False},
    )
    monkeypatch.setattr(
        "app.domain.workspace_protocol.workspace_docx_blockers",
        lambda *_a, **_k: [],
    )
    monkeypatch.setattr(
        "app.domain.workspace_protocol.build_workspace_assembly_context",
        lambda *_a, **_k: {
            "product": {"trade_name": "Upadacitinib", "inn": "upadacitinib"},
            "study": {"protocol_number": "UPDCB-02-BE-2026", "version": "1"},
            "workspace": True,
        },
    )
    monkeypatch.setattr(
        "app.domain.workspace_protocol.contamination_preflight",
        lambda *_a, **_k: {
            "ok": False,
            "code": "CRITICAL_TEMPLATE_CONTAMINATION",
            "message": "Template contains product-specific content that is not supported by the current study.",
            "action": "Resolve template/content mapping before DOCX generation.",
            "unmanaged_blocks": [
                {
                    "block_id": "fixture.stale_pharmacology",
                    "reason": "MISSING_VERIFIED_EVIDENCE",
                    "action": "BLOCK",
                }
            ],
        },
    )
    db = db_module.SessionLocal()
    try:
        with pytest.raises(ValidationError) as ei:
            generate_docx_artifact(db, "UPDCB-02-BE-2026", created_by="p292", force_warnings_ok=True)
        assert ei.value.field == "CRITICAL_TEMPLATE_CONTAMINATION"
    finally:
        db.close()


def test_updcB_generated_docx_content_audit(tmp_path: Path):
    """Template → clear → scan must leave no CML/chemistry/Bosutinib semantics."""
    tpl = get_template_profile().template_path()
    out = tmp_path / "updcb_clean.docx"
    shutil.copy2(tpl, out)
    doc = Document(str(out))
    ctx = {
        "product": {"trade_name": "upadacitinib", "inn": "upadacitinib"},
        "reference_product": {"trade_name": "РАНВЭК"},
        "study": {"protocol_number": "UPDCB-02-BE-2026"},
    }
    clear_product_specific_contamination(doc, ctx)
    doc.save(str(out))
    scrub_contamination_xml_package(out, ctx)
    scan = scan_docx_contamination(out, ctx)
    assert scan["contaminated"] is False, scan
    blob = "\n".join(p.text or "" for p in Document(str(out)).paragraphs)
    assert "C26H29Cl2N5O3" not in blob
    assert "хронического миелоидного лейкоза" not in blob
    assert "Бозутиниб" not in blob
    assert "Bosulif" not in blob
