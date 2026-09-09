"""Phase 6 research engine / ingestion unit tests — offline only."""

from __future__ import annotations

from app.domain.conflict_engine import ClaimSnippet, detect_conflicts
from app.domain.document_ingest import (
    extract_document,
    extract_txt,
    sha256_hex,
    validate_upload,
)
from app.domain.document_search import search_chunks
from app.domain.evidence_normalize import normalize_claim_value
from app.domain.exceptions import ValidationError
from app.domain.research_completeness import calculate_research_completeness
from app.domain.research_profile import ClientResearchInput, build_search_profile
from app.domain.research_tasks import TASK_DEPENDENCIES, generate_research_tasks
import pytest


def test_search_profile_generation() -> None:
    profile = build_search_profile(
        ClientResearchInput(
            requested_product_name="Bosutinib",
            inn="bosutinib",
            dosage="400 mg",
            dosage_form="tablet",
            route="oral",
            requested_subject_count=46,
            regulatory_jurisdiction="EAEU",
        )
    )
    assert profile.inn == "bosutinib"
    assert any("bioequivalence" in t.lower() or "биоэквивалентность" in t.lower() for t in profile.bioequivalence_search_terms)
    assert profile.pk_search_terms
    assert profile.version == "SEARCH.PROFILE.v1"


def test_task_generation_and_dependencies() -> None:
    search = build_search_profile(ClientResearchInput(inn="bosutinib", dosage="400 mg"))
    tasks = generate_research_tasks(search)
    types = [t.task_type for t in tasks]
    assert types[0] == "REFERENCE_PRODUCT"
    assert "PK" in types
    pk = next(t for t in tasks if t.task_type == "PK")
    assert pk.depends_on_tasks == TASK_DEPENDENCIES["PK"]
    assert "PRODUCT_LABEL" in pk.depends_on_tasks


def test_task_blocked_without_identity() -> None:
    tasks = generate_research_tasks(build_search_profile(ClientResearchInput()))
    assert tasks[0].status == "BLOCKED"


def test_txt_ingestion_checksum_chunks() -> None:
    content = b"Page1 Tmax was 2-3 h.\fPage2 CVintra 25%"
    safe = validate_upload(
        filename="note.txt", mime_type="text/plain", size=len(content), content=content
    )
    assert safe.endswith(".txt")
    digest = sha256_hex(content)
    assert len(digest) == 64
    result = extract_txt(content)
    assert len(result.pages) == 2
    assert result.chunks
    assert any("Tmax" in c.text for c in result.chunks)


def test_docx_xlsx_extraction() -> None:
    from io import BytesIO

    from docx import Document as DocxDocument
    from openpyxl import Workbook

    buf = BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Half-life was 10–12 hours.")
    doc.save(buf)
    docx_bytes = buf.getvalue()
    validate_upload(
        filename="a.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size=len(docx_bytes),
        content=docx_bytes,
    )
    docx_res = extract_document(filename="a.docx", content=docx_bytes)
    assert "Half-life" in docx_res.pages[0].text

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "CV"
    ws["B1"] = "25%"
    xbuf = BytesIO()
    wb.save(xbuf)
    xbytes = xbuf.getvalue()
    xres = extract_document(filename="t.xlsx", content=xbytes)
    assert any("25%" in p.text for p in xres.pages)


def test_pdf_extraction() -> None:
    from io import BytesIO

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()
    validate_upload(
        filename="blank.pdf", mime_type="application/pdf", size=len(pdf_bytes), content=pdf_bytes
    )
    res = extract_document(filename="blank.pdf", content=pdf_bytes)
    assert res.pages
    assert isinstance(res.pages[0].text, str)


def test_forbidden_executable() -> None:
    with pytest.raises(ValidationError):
        validate_upload(
            filename="evil.exe", mime_type="application/octet-stream", size=4, content=b"MZ\x00\x00"
        )


def test_normalize_tmax_and_cv() -> None:
    n = normalize_claim_value("tmax", "Tmax was 2–3 часа")
    assert n.normalized["min"] == 2
    assert n.normalized["max"] == 3
    assert n.unit == "h"
    cv = normalize_claim_value("cv_cmax", "CVintra 25%")
    assert cv.normalized["value"] == 25


def test_conflict_detection() -> None:
    conflicts = detect_conflicts(
        [
            ClaimSnippet("e1", "c1", "tmax", "2 h", {"min": 2, "max": 2, "unit": "h"}, "PROPOSED"),
            ClaimSnippet("e2", "c2", "tmax", "2–3 h", {"min": 2, "max": 3, "unit": "h"}, "PROPOSED"),
        ]
    )
    assert len(conflicts) == 1
    assert conflicts[0].field_name == "tmax"


def test_search_chunks() -> None:
    hits = search_chunks(
        [
            {
                "chunk_id": "1",
                "document_id": "d1",
                "page_number": 1,
                "text": "Tmax was observed at 2-3 hours after dosing",
                "source_id": "s1",
                "source_type": "SmPC",
            }
        ],
        "Tmax 2-3",
        phrase=False,
    )
    assert hits
    assert hits[0].page == 1


def test_completeness_score() -> None:
    result = calculate_research_completeness(
        tasks=[
            {"task_type": "REFERENCE_PRODUCT", "status": "DONE"},
            {"task_type": "PRODUCT_LABEL", "status": "DONE"},
            {"task_type": "PK", "status": "TODO"},
            {"task_type": "BIOEQUIVALENCE_STUDIES", "status": "BLOCKED"},
            {"task_type": "CV", "status": "TODO"},
            {"task_type": "FOOD_CONDITION", "status": "TODO"},
        ],
        evidence_field_codes=["tmax", "reference_product"],
        unresolved_conflicts=1,
    )
    assert 0 <= result.score <= 100
    assert "PK" in result.open_tasks
    assert "half_life" in result.missing_evidence
