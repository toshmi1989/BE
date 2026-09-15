"""Phase 30.3 — DOCX semantic integrity: placeholders, stale content, consistency.

Stronger than table/paragraph counts. Fail-closed for FINAL.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document

from app.domain.cover_mapping import detect_wrong_cover_mapping
from app.domain.protocol_template_registry import study_is_bosutinib_sample

UNRESOLVED_RE = re.compile(r"\{\{[A-Z0-9_.]+\}\}")
INTERNAL_TABLE_ID_RE = re.compile(
    r"таблица\s+(PK_PARAMETERS|BLOOD_SAMPLING|STUDY_METADATA|TEST_PRODUCT|REFERENCE_PRODUCT|CV_EVIDENCE)\b",
    re.IGNORECASE,
)
TEMPLATE_DATE_RE = re.compile(r"18\.04\.2025")
STALE_400_RE = re.compile(r"\b400\s*(mg|мг)\b", re.IGNORECASE)
WASHOUT_DAY_RE = re.compile(
    r"(?:отмывк\w*|washout)[^\d]{0,40}?(\d+(?:[.,]\d+)?)\s*(?:дн|день|дня|дней|day|days)",
    re.IGNORECASE,
)
WASHOUT_DAY_ALT_RE = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*(?:дн|день|дня|дней|day|days)[^\d]{0,40}?(?:отмывк\w*|washout)",
    re.IGNORECASE,
)

# Placeholders intentionally allowed in FINAL (empty by default = block all)
FINAL_PLACEHOLDER_ALLOWLIST: frozenset[str] = frozenset()


@dataclass
class IntegrityIssue:
    severity: str
    code: str
    message: str
    details: dict[str, Any] = field(default_factory=dict)


@dataclass
class SemanticIntegrityReport:
    ok: bool
    mode: str
    issues: list[IntegrityIssue] = field(default_factory=list)
    placeholder_count: int = 0
    stale_content_count: int = 0
    wrong_mappings: int = 0
    cross_section_failures: int = 0
    toc_refreshed: bool = False
    unresolved_required_fields: int = 0

    @property
    def blocking(self) -> bool:
        return any(i.severity == "CRITICAL" for i in self.issues)

    def to_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "mode": self.mode,
            "blocking": self.blocking,
            "placeholder_count": self.placeholder_count,
            "stale_content_count": self.stale_content_count,
            "wrong_mappings": self.wrong_mappings,
            "cross_section_failures": self.cross_section_failures,
            "toc_refreshed": self.toc_refreshed,
            "unresolved_required_fields": self.unresolved_required_fields,
            "issues": [
                {"severity": i.severity, "code": i.code, "message": i.message, "details": i.details}
                for i in self.issues
            ],
        }


def _iter_all_texts(doc: Document) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for i, p in enumerate(doc.paragraphs):
        style = (p.style.name if p.style else "") or ""
        out.append((f"para[{i}:{style}]", p.text or ""))
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                out.append((f"T{ti}R{ri}C{ci}", cell.text or ""))
    for si, section in enumerate(doc.sections):
        for attr in ("header", "footer", "first_page_header", "first_page_footer"):
            part = getattr(section, attr, None)
            if part is None:
                continue
            for pi, p in enumerate(part.paragraphs):
                out.append((f"S{si}.{attr}.p{pi}", p.text or ""))
            for ti, table in enumerate(part.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        out.append((f"S{si}.{attr}.T{ti}R{ri}C{ci}", cell.text or ""))
    return out


def _scan_xml_placeholders(path: Path) -> list[str]:
    found: list[str] = []
    try:
        with ZipFile(path) as zf:
            for name in zf.namelist():
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                raw = zf.read(name).decode("utf-8", errors="ignore")
                found.extend(UNRESOLVED_RE.findall(raw))
    except OSError:
        pass
    return sorted(set(found))


def _canonical_washout_days(study_ctx: dict[str, Any]) -> float | None:
    wo = study_ctx.get("washout") or {}
    val = wo.get("selected_value")
    unit = str(wo.get("unit") or "day").lower()
    if val is None:
        return None
    try:
        num = float(val)
    except (TypeError, ValueError):
        return None
    if unit.startswith("h"):
        return num / 24.0
    return num


def _canonical_dose_tokens(study_ctx: dict[str, Any]) -> list[str]:
    tokens: list[str] = []
    for key in ("product", "reference_product"):
        dose = (study_ctx.get(key) or {}).get("dosage")
        if dose:
            tokens.append(str(dose).strip().lower())
    return tokens


def assess_docx_semantic_integrity(
    path: Path,
    study_ctx: dict[str, Any],
    *,
    mode: str = "DRAFT",
    toc_refreshed: bool = False,
) -> SemanticIntegrityReport:
    mode_u = str(mode or "DRAFT").upper()
    issues: list[IntegrityIssue] = []
    report = SemanticIntegrityReport(ok=True, mode=mode_u, toc_refreshed=toc_refreshed)

    if not path.exists():
        issues.append(IntegrityIssue("CRITICAL", "FILE_MISSING", f"Missing DOCX: {path}"))
        report.issues = issues
        report.ok = False
        return report

    doc = Document(str(path))
    texts = _iter_all_texts(doc)
    blob = "\n".join(t for _, t in texts)

    # Placeholders (body + XML w:t via zip)
    placeholders = sorted(set(UNRESOLVED_RE.findall(blob) + _scan_xml_placeholders(path)))
    blocked_ph = [p for p in placeholders if p not in FINAL_PLACEHOLDER_ALLOWLIST]
    report.placeholder_count = len(blocked_ph)
    if blocked_ph and mode_u == "FINAL":
        issues.append(
            IntegrityIssue(
                "CRITICAL",
                "UNRESOLVED_TEMPLATE_PLACEHOLDER",
                f"FINAL DOCX has unresolved placeholders: {blocked_ph[:25]}",
                {"placeholders": blocked_ph},
            )
        )
        report.unresolved_required_fields = len(blocked_ph)
    elif blocked_ph and mode_u == "DRAFT":
        issues.append(
            IntegrityIssue(
                "WARNING",
                "UNRESOLVED_TEMPLATE_PLACEHOLDER",
                f"DRAFT still has placeholders (allowed for draft only): {blocked_ph[:15]}",
                {"placeholders": blocked_ph},
            )
        )

    bosutinib = study_is_bosutinib_sample(study_ctx.get("product") or {})

    # Stale 400 mg when study dose is not 400
    dose_tokens = _canonical_dose_tokens(study_ctx)
    study_is_400 = any("400" in d for d in dose_tokens)
    if not bosutinib and not study_is_400 and STALE_400_RE.search(blob):
        report.stale_content_count += len(STALE_400_RE.findall(blob))
        issues.append(
            IntegrityIssue(
                "CRITICAL" if mode_u == "FINAL" else "WARNING",
                "STALE_TEMPLATE_DOSE_400",
                "Stale template dose 400 mg/мг remains in DOCX",
                {"canonical_doses": dose_tokens},
            )
        )

    # Template date leakage
    if TEMPLATE_DATE_RE.search(blob):
        study_date = str((study_ctx.get("study") or {}).get("version_date") or "")
        if "18.04.2025" not in study_date:
            report.stale_content_count += 1
            issues.append(
                IntegrityIssue(
                    "CRITICAL" if mode_u == "FINAL" else "WARNING",
                    "STALE_TEMPLATE_DATE",
                    "Template date 18.04.2025 remains; study version_date missing or different",
                    {"study_version_date": study_date},
                )
            )

    # Internal table IDs
    if INTERNAL_TABLE_ID_RE.search(blob):
        issues.append(
            IntegrityIssue(
                "CRITICAL" if mode_u == "FINAL" else "WARNING",
                "INTERNAL_TABLE_ID_LEAK",
                "Internal table key leaked into document prose",
                {"matches": INTERNAL_TABLE_ID_RE.findall(blob)[:10]},
            )
        )

    # Wrong cover mapping: dosage_form = subject count
    if len(doc.tables) > 0:
        for row in doc.tables[0].rows:
            if len(row.cells) < 2:
                continue
            label, value = row.cells[0].text or "", row.cells[1].text or ""
            code = detect_wrong_cover_mapping(label, value)
            if code:
                report.wrong_mappings += 1
                issues.append(
                    IntegrityIssue(
                        "CRITICAL",
                        code,
                        f"Cover cell '{label.strip()}' has subject-count-like value '{value.strip()}'",
                        {"label": label, "value": value},
                    )
                )

    # Washout cross-section consistency
    canonical_wo = _canonical_washout_days(study_ctx)
    if canonical_wo is not None:
        found_days: list[float] = []
        for loc, text in texts:
            for rx in (WASHOUT_DAY_RE, WASHOUT_DAY_ALT_RE):
                for m in rx.finditer(text):
                    try:
                        found_days.append(float(m.group(1).replace(",", ".")))
                    except ValueError:
                        pass
        mismatched = [d for d in found_days if abs(d - float(canonical_wo)) > 0.01]
        if mismatched:
            report.cross_section_failures += 1
            issues.append(
                IntegrityIssue(
                    "CRITICAL" if mode_u == "FINAL" else "WARNING",
                    "WASHOUT_INCONSISTENT",
                    f"Rendered washout days {sorted(set(mismatched))} != canonical {canonical_wo}",
                    {"canonical_days": canonical_wo, "found": sorted(set(found_days))},
                )
            )

    # TOC
    if mode_u == "FINAL" and not toc_refreshed:
        issues.append(
            IntegrityIssue(
                "CRITICAL",
                "TOC_NOT_REFRESHED",
                "TOC fields were not refreshed after render",
            )
        )

    # Sampling empty placeholders
    if mode_u == "FINAL" and ("{{SAMPLING.POINTS}}" in blob or "{{SAMPLING.TIME_H}}" in blob or "{{SAMPLING.POINT}}" in blob):
        issues.append(
            IntegrityIssue(
                "CRITICAL",
                "MISSING_SAMPLING_PLAN",
                "Sampling placeholders remain — SamplingPlan required",
            )
        )

    report.issues = issues
    report.ok = not any(i.severity == "CRITICAL" for i in issues)
    return report


def semantic_preflight_from_context(study_ctx: dict[str, Any], *, mode: str = "DRAFT") -> dict[str, Any]:
    """Pre-render gates for missing required sources (FINAL fail-closed)."""
    mode_u = str(mode or "DRAFT").upper()
    blockers: list[dict[str, str]] = []
    study = study_ctx.get("study") or {}
    product = study_ctx.get("product") or {}
    subjects = study_ctx.get("subjects") or {}
    sampling = study_ctx.get("sampling") or {}
    eligibility = study_ctx.get("eligibility") or {}
    bio = study_ctx.get("bioanalysis_plan") or study_ctx.get("bioanalysis") or {}
    stats = study_ctx.get("statistical_config") or {}
    washout = study_ctx.get("washout") or {}
    observation = study_ctx.get("observation") or {}

    def add(code: str, reason: str, field: str) -> None:
        blockers.append({"code": code, "reason": reason, "field": field})

    if mode_u != "FINAL":
        return {"ok": True, "mode": mode_u, "blockers": [], "message": "DRAFT semantic preflight soft"}

    if not study.get("version_date"):
        add("MISSING_PROTOCOL_DATE", "Study/ProtocolDraft version_date required for FINAL", "study.version_date")
    if not product.get("dosage_form"):
        add("MISSING_DOSAGE_FORM", "dosage_form required (typed string)", "product.dosage_form")
    if not product.get("dosage"):
        add("MISSING_DOSE", "dose required (typed Dose)", "product.dosage")
    if subjects.get("target_evaluable_n") is None and subjects.get("planned_randomized_n") is None:
        add("MISSING_SUBJECT_COUNTS", "evaluable/randomized targets required", "subjects")
    if not (sampling.get("points") or []):
        add("MISSING_SAMPLING_PLAN", "SamplingPlan points required", "sampling.points")
    if not washout.get("selected_value"):
        add("MISSING_WASHOUT", "Canonical washout required", "washout")
    if observation.get("selected_duration") is None and observation.get("final_sampling_time") is None:
        add("MISSING_OBSERVATION_DURATION", "Observation duration required", "observation")
    inc = eligibility.get("inclusion") or []
    if not inc:
        add("MISSING_ELIGIBILITY", "Eligibility inclusion criteria required", "eligibility.inclusion")
    if not bio:
        add("MISSING_BIOANALYSIS", "Bioanalysis plan required or explicitly blocked", "bioanalysis_plan")
    if not stats or str(stats.get("status") or "").upper() not in {"APPROVED", "ACCEPTED"}:
        # PRIMARY_BE approval lives in statistics plan
        add("MISSING_STATISTICS_PLAN", "APPROVED StatisticsPlan required for FINAL", "statistical_config")

    return {
        "ok": len(blockers) == 0,
        "mode": mode_u,
        "blockers": blockers,
        "message": "FINAL semantic preflight passed" if not blockers else "FINAL blocked by semantic gaps",
    }
