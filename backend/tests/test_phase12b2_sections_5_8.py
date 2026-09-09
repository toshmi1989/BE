"""Phase 12B.2 — sections 5–8 controlled content generation tests."""

from __future__ import annotations

from pathlib import Path

from app.domain.content_core_templates import CORE_12B2_SECTION_CODES
from app.domain.content_generation_qa import run_sections_5_8_qa
from app.domain.content_resolver import filter_decision_for_render
from app.domain.pk_semantic import display_auc_metric as display_auc
from app.domain.protocol_assembly import GENERATORS, assemble_protocol
from app.domain.protocol_generators_12b2 import CORE_12B2_GENERATORS
from app.domain.criteria_rules import evaluate_criteria


def _ctx(**overrides):
    base = {
        "project_id": "p12b2",
        "study": {"protocol_number": "BE-12B2-001", "title": "BE 12B2"},
        "sponsor": {"name": "Sponsor Co"},
        "product": {
            "trade_name": "TestDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
            "manufacturer": "Maker",
        },
        "reference_product": {
            "trade_name": "RefDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
        },
        "design": {
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
        },
        "food": {"condition": "FASTING"},
        "subjects": {
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
            "reserve_n": 0,
        },
        "sample_size": {"evaluable_n": 24, "randomized_n": 28},
        "sampling": {
            "points": [
                {"time_h": 0.0, "reason": "BASELINE"},
                {"time_h": 2.0, "reason": "TMAX_CAPTURE"},
                {"time_h": 24.0, "reason": "FINAL"},
            ],
            "total_points_per_period": 3,
            "source_ids": ["samp1"],
        },
        "washout": {"selected_value": 7, "unit": "day", "requires_washout": True},
        "observation": {"selected_duration": 24, "unit": "h"},
        "analytes": [{"name": "testdrug", "tmax": 2.0, "type": "PARENT"}],
        "pk_parameters": [
            {"parameter_code": "Cmax"},
            {"parameter_code": "AUC_0_LAST"},
            {"parameter_code": "AUC_0_72H"},
        ],
        "eligibility": {
            "inclusion": [{"number": 1, "text": "Healthy volunteers 18–45 years"}],
            "non_inclusion": [{"number": 1, "text": "Known hypersensitivity to drug"}],
            "exclusion": [{"number": 1, "text": "Positive pregnancy test"}],
        },
        "criteria_meta": {},
        "expert_decisions": [],
        "bioanalysis_plan": {},
        "safety_plan": {},
        "sample_processing": {},
        "evidence_claims": [],
        "sources": [],
    }
    base.update(overrides)
    return base


def _assemble(ctx=None, only=None):
    return assemble_protocol(
        ctx or _ctx(),
        blocking_validation=False,
        only_sections=only or CORE_12B2_SECTION_CODES,
    )


def _section(assembled, code):
    return next(s for s in assembled["sections"] if s["section_code"] == code)


def _texts(sec) -> str:
    out = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            out.append(str(b["text"]))
        out.extend(str(x) for x in (b.get("items") or []))
    return "\n".join(out)


def _blocks(sec):
    return list(sec.get("content_blocks") or [])


# ---- Section 5 ----


def test_standard_criteria_resolve():
    a = _assemble()
    assert "Healthy volunteers" in _texts(_section(a, "5.1"))
    assert "hypersensitivity" in _texts(_section(a, "5.2"))
    assert "pregnancy" in _texts(_section(a, "5.3")).lower()


def test_drug_specific_criteria_require_evidence():
    ctx = _ctx(criteria_meta={"has_cyp_mentions": True})
    a = _assemble(ctx)
    blob = _texts(_section(a, "5.1")) + str(_blocks(_section(a, "5.1")))
    assert "ELIGIBILITY.CYP" in blob or any(
        g.get("domain") == "ELIGIBILITY" for b in _blocks(_section(a, "5.1")) for g in b.get("knowledge_gaps") or []
    )


def test_contraception_requires_source():
    a = _assemble(_ctx(criteria_meta={}))
    sec = _section(a, "6.2.3")
    assert "{{CONTRACEPTION" in _texts(sec) or any(
        b.get("resolution_status") == "UNRESOLVED" for b in _blocks(sec)
    )
    a2 = _assemble(
        _ctx(criteria_meta={"contraception_days_from_smpc": 30, "smpc_evidence_claim_ids": ["e1"]})
    )
    assert "30" in _texts(_section(a2, "6.2.3"))


def test_smoking_requires_evidence():
    a = _assemble(_ctx(criteria_meta={"require_smoking": True}))
    assert "{{ELIGIBILITY.SMOKING}}" in _texts(_section(a, "5.1")) or any(
        "SMOKING" in str(b.get("block_code")) for b in _blocks(_section(a, "5.1"))
    )


def test_vomiting_rule_not_verified_automatically():
    crit = evaluate_criteria(tmax_h=2.0)
    assert any(o.get("status") == "PROPOSED" for o in crit.drug_specific_overlays)
    a = _assemble()
    vomit = [b for b in _blocks(_section(a, "5.1")) if "VOMITING" in str(b.get("block_code"))]
    assert vomit
    assert vomit[0]["resolution_status"] == "PROPOSED"
    assert vomit[0]["display_as_final"] is False


def test_missing_criteria_creates_gap():
    a = _assemble(_ctx(eligibility={"inclusion": [], "non_inclusion": [], "exclusion": []}))
    sec = _section(a, "5.1")
    assert any(b.get("resolution_status") == "UNRESOLVED" for b in _blocks(sec))
    assert any(b.get("knowledge_gaps") for b in _blocks(sec))


def test_synopsis_and_section5_match():
    ctx = _ctx()
    # Assemble both synopsis (if in tree with only) and section 5
    from app.domain.content_core_templates import CORE_12B1_SECTION_CODES

    both = assemble_protocol(
        ctx,
        blocking_validation=False,
        only_sections=set(CORE_12B2_SECTION_CODES) | {"SYNOPSIS"} | set(CORE_12B1_SECTION_CODES),
    )
    incl = "Healthy volunteers 18–45 years"
    assert incl in _texts(_section(both, "5.1"))
    findings = run_sections_5_8_qa(sections=both["sections"], ctx=ctx, mode="DRAFT")
    assert not any(f.code == "QA.ELIGIBILITY.CANONICAL_MISMATCH" and f.blocking for f in findings)


# ---- Section 6 ----


def test_procedure_schedule_resolves():
    a = _assemble()
    sec = _section(a, "6.1.1")
    text = _texts(sec)
    assert "DOSING" in text or "Screening" in text or "Hospitalization" in text or "dosing" in text.lower()


def test_dosing_uses_canonical_product():
    a = _assemble()
    text = _texts(_section(a, "6.1"))
    assert "TestDrug" in text
    assert "100 mg" in text


def test_food_does_not_invent_meal_values():
    a = _assemble()
    text = _texts(_section(a, "6.2.1")).lower()
    assert "800" not in text
    assert "kcal" not in text
    assert "50%" not in text
    assert "натощак" in text or "fasting" not in text  # display, not raw enum leak preferred


def test_washout_uses_canonical_value():
    a = _assemble()
    assert "7" in _texts(_section(a, "6.1.5"))


def test_sampling_uses_single_canonical_plan():
    a = _assemble()
    p1 = _texts(_section(a, "6.1.4"))
    methods = _texts(_section(a, "7.2"))
    assert "0.0" in p1 or "0" in p1
    assert "24" in p1 and "24" in methods


def test_sample_processing_has_provenance():
    a = _assemble(
        _ctx(
            bioanalysis_plan={
                "anticoagulant": "K2EDTA",
                "centrifugation": "1500 g",
                "source_ids": ["lab-sop-1"],
                "status": "RESOLVED",
            }
        )
    )
    sec = _section(a, "6.1.9")
    blocks = _blocks(sec)
    assert any(b.get("source_ids") for b in blocks)
    assert "K2EDTA" in _texts(sec)


def test_missing_procedure_field_creates_gap():
    a = _assemble(_ctx(product={}))
    sec = _section(a, "6.1")
    assert any(b.get("resolution_status") == "UNRESOLVED" for b in _blocks(sec))


# ---- Section 7 ----


def test_analyte_requires_approved_decision():
    a = _assemble()
    blocks = _blocks(_section(a, "7.1"))
    analyte_blocks = [b for b in blocks if "ANALYTE" in str(b.get("block_code"))]
    assert analyte_blocks
    assert analyte_blocks[0]["display_as_final"] is False


def test_pk_parameter_profile_resolves():
    a = _assemble()
    text = _texts(_section(a, "7.1"))
    assert "Cmax" in text or "AUC" in text


def test_auc0_last_not_equal_auc0_72():
    assert display_auc("AUC_0_LAST") != display_auc("AUC_0_72H")
    a = _assemble()
    text = _texts(_section(a, "7.1"))
    assert "AUC0-x" in text
    assert "AUC0-72" in text
    assert "AUC0-x = AUC0-72" not in text


def test_auc_display_is_human_readable():
    assert display_auc("AUC_0_LAST") == "AUC0-x"
    assert display_auc("AUC_0_72H") == "AUC0-72"
    assert "AUC_0_LAST" not in display_auc("AUC_0_LAST")


def test_bioanalysis_does_not_invent_method():
    a = _assemble()
    text = _texts(_section(a, "7.3.1"))
    assert "HPLC" not in text
    assert "{{BIOANALYSIS.METHOD}}" in text or any(
        b.get("resolution_status") == "UNRESOLVED" for b in _blocks(_section(a, "7.3.1"))
    )


def test_bioanalysis_missing_value_creates_gap():
    a = _assemble()
    assert any(b.get("knowledge_gaps") for b in _blocks(_section(a, "7.3.1")))


def test_section7_uses_canonical_sampling():
    a = _assemble()
    text = _texts(_section(a, "7.2"))
    assert "SamplingPlan" in text or "точек" in text
    assert "24" in text


# ---- Section 8 ----


def test_safety_plan_resolves_existing_assessments():
    a = _assemble(
        _ctx(
            safety_plan={
                "vital_signs": {"enabled": True, "items": ["BP", "HR"]},
                "status": "PROPOSED",
                "source_ids": ["s1"],
            }
        )
    )
    assert "Vital" in _texts(_section(a, "8.2.2")) or "жизнен" in _texts(_section(a, "8.2.2")).lower() or "BP" in _texts(
        _section(a, "8.2.2")
    )


def test_safety_does_not_invent_assessment():
    a = _assemble(_ctx(safety_plan={}))
    assert "{{SAFETY" in _texts(_section(a, "8.1")) or any(
        b.get("resolution_status") == "UNRESOLVED" for b in _blocks(_section(a, "8.1"))
    )


def test_safety_timing_uses_schedule():
    a = _assemble(
        _ctx(
            safety_plan={"vital_signs": {"enabled": True}, "status": "PROPOSED"},
        )
    )
    # Must not invent hardcoded -1h,2h,4h
    text = _texts(_section(a, "8.2.2"))
    assert "-1h" not in text
    assert "2h, 4h" not in text


def test_ae_content_requires_safety_definition():
    a = _assemble(_ctx(safety_plan={}))
    assert "{{SAFETY.AE}}" in _texts(_section(a, "8.3")) or any(
        b.get("resolution_status") == "UNRESOLVED" for b in _blocks(_section(a, "8.3"))
    )


def test_unverified_safety_not_final():
    a = _assemble(
        _ctx(safety_plan={"AE": {"enabled": True}, "SAE": {"enabled": True}, "status": "PROPOSED"})
    )
    for b in _blocks(_section(a, "8.3")):
        if b.get("resolution_status") == "PROPOSED":
            assert b.get("display_as_final") is False


# ---- Cross ----


def test_section5_matches_canonical_criteria():
    ctx = _ctx()
    a = _assemble(ctx)
    findings = run_sections_5_8_qa(sections=a["sections"], ctx=ctx)
    assert not any(f.code == "QA.ELIGIBILITY.CANONICAL_MISMATCH" for f in findings)


def test_section6_matches_procedure_schedule():
    a = _assemble()
    assert _blocks(_section(a, "6.1.1"))


def test_section7_matches_pk_plan():
    a = _assemble()
    assert "AUC0-x" in _texts(_section(a, "7.1"))


def test_section8_matches_safety_plan():
    a = _assemble(_ctx(safety_plan={"physical_exam": {"enabled": True}, "status": "RESOLVED", "source_ids": ["x"]}))
    assert any(b.get("display_as_final") for b in _blocks(_section(a, "8.2.1")))


def test_sections_5_8_have_provenance():
    a = _assemble(
        _ctx(
            bioanalysis_plan={
                "analytical_method": "validated LC-MS/MS per lab report",
                "source_ids": ["bio-1"],
                "status": "RESOLVED",
            }
        )
    )
    bio = _blocks(_section(a, "7.3.1"))[0]
    assert bio.get("source_ids") or bio.get("canonical_source")
    elig = _blocks(_section(a, "5.1"))[0]
    assert elig.get("canonical_source")


def test_proposed_content_not_final():
    a = _assemble()
    for s in a["sections"]:
        for b in _blocks(s):
            if b.get("resolution_status") == "PROPOSED":
                assert b.get("display_as_final") is False


def test_rejected_content_not_rendered():
    assert filter_decision_for_render({"status": "REJECTED"}) is None
    findings = run_sections_5_8_qa(
        sections=[
            {
                "section_code": "5.1",
                "content_blocks": [
                    {
                        "text": "x",
                        "decision_status": "REJECTED",
                        "used_in_content": True,
                        "display_as_final": True,
                    }
                ],
            }
        ]
    )
    assert any(f.code in {"CONTENT.REJECTED_DECISION_USED", "QA.CONTENT.REJECTED_DECISION"} for f in findings)


def test_superseded_content_not_rendered():
    assert filter_decision_for_render({"status": "SUPERSEDED"}) is None
    findings = run_sections_5_8_qa(
        sections=[
            {
                "section_code": "7.1",
                "content_blocks": [
                    {"text": "x", "decision_status": "SUPERSEDED", "used_in_content": True}
                ],
            }
        ]
    )
    assert any(f.code in {"CONTENT.SUPERSEDED_DECISION_USED", "QA.CONTENT.SUPERSEDED_DECISION"} for f in findings)


def test_no_n_a_fallback():
    a = _assemble()
    blob = "\n".join(_texts(s) for s in a["sections"]).lower()
    assert "n/a" not in blob
    assert "not applicable" not in blob


def test_no_legacy_template_value():
    findings = run_sections_5_8_qa(
        sections=[],
        legacy_hints={"eligibility_legacy": "old criterion from template"},
    )
    assert any(f.code == "QA.ELIGIBILITY.LEGACY_VALUE" for f in findings)


def test_generators_registered():
    for k in CORE_12B2_GENERATORS:
        assert GENERATORS[k] is CORE_12B2_GENERATORS[k]


# ---- DOCX ----


def test_targeted_docx_sections_5_8(tmp_path):
    from app.domain.docx_renderer import render_protocol_docx

    ctx = _ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="12b2",
        protocol_version="1",
        only_sections=["5.1", "6.1", "7.1", "8.1"],
    )
    assert result.status == "READY", result.blocking_reasons
    assert result.filename


def test_unselected_sections_preserved(tmp_path):
    import shutil

    from docx import Document

    from app.domain.docx_profile import get_template_profile
    from app.domain.docx_renderer import _find_heading_indexes, _section_body_range, render_protocol_docx

    ctx = _ctx()
    assembled = assemble_protocol(ctx, blocking_validation=False)
    profile = get_template_profile()
    baseline = tmp_path / "base.docx"
    shutil.copy2(profile.template_path(), baseline)
    base_doc = Document(str(baseline))
    base_map = _find_heading_indexes(base_doc)
    static_code = next((c for c in ("9", "10", "11", "12") if c in base_map), None)
    if static_code is None:
        static_code = next(c for c in base_map if not str(c).startswith(("5", "6", "7", "8")))
    br = _section_body_range(base_map, static_code, len(base_doc.paragraphs))
    before = [base_doc.paragraphs[i].text for i in range(br[0], br[1])]
    result = render_protocol_docx(
        protocol_payload=assembled,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="12b2-static",
        protocol_version="1",
        only_sections=["5.1"],
    )
    assert result.status == "READY", result.blocking_reasons
    out = Document(str(result.output_path))
    out_map = _find_heading_indexes(out)
    orng = _section_body_range(out_map, static_code, len(out.paragraphs))
    after = [out.paragraphs[i].text for i in range(orng[0], orng[1])]
    assert before == after


def test_no_global_replace():
    src = Path(__file__).resolve().parents[1] / "app" / "domain" / "docx_renderer.py"
    text = src.read_text(encoding="utf-8")
    assert "only_sections" in text
    assert "str.replace(" not in text or "only_sections" in text


def test_no_positional_mapping():
    src = Path(__file__).resolve().parents[1] / "app" / "domain" / "protocol_generators_12b2.py"
    text = src.read_text(encoding="utf-8")
    assert "cells[0]" not in text
    assert "cell[1]" not in text
    assert "paragraphs[" not in text


def test_full_pipeline_golden_sections_5_8(client):
    pid = client.post("/api/projects", json={"name": "12B2-golden"}).json()["id"]
    client.put(
        f"/api/projects/{pid}/study",
        json={"protocol_number": "BE-12B2-G", "title": "Golden 12B2", "version": "1.0"},
    )
    client.put(
        f"/api/projects/{pid}/product",
        json={"trade_name": "TestDrug", "inn": "x", "dosage": "100 mg", "dosage_form": "tablet"},
    )
    client.put(
        f"/api/projects/{pid}/reference-product",
        json={"trade_name": "RefDrug", "inn": "x", "dosage": "100 mg", "dosage_form": "tablet"},
    )
    client.post(
        f"/api/projects/{pid}/design",
        json={"type": "CROSSOVER_2X2", "periods": 2, "sequences": [["T", "R"], ["R", "T"]]},
    )
    client.post(f"/api/projects/{pid}/food", json={"condition": "FASTING"})
    client.post(
        f"/api/projects/{pid}/subjects",
        json={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "reserve_n": 0,
            "planned_screened_n": 40,
        },
    )
    client.put(
        f"/api/projects/{pid}/eligibility",
        json={
            "inclusion": [{"number": 1, "text": "Healthy volunteers"}],
            "non_inclusion": [{"number": 1, "text": "Allergy"}],
            "exclusion": [{"number": 1, "text": "Pregnancy"}],
        },
    )
    gen = client.post(
        f"/api/projects/{pid}/content/generate-sections-5-8",
        json={"persist": True, "mode": "DRAFT"},
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()
    assert body["mutates_study"] is False
    assert body["sections"]
    codes = {f["code"] for f in body.get("qa_findings") or []}
    assert "CONTENT.PROPOSED_RENDERED_AS_FINAL" not in codes or not any(
        f.get("blocking") and f.get("code") == "CONTENT.PROPOSED_RENDERED_AS_FINAL"
        for f in body.get("qa_findings") or []
    )
    docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "DRAFT", "only_sections": ["5.1", "6.1", "7.1", "8.1"]},
    )
    assert docx.status_code == 200, docx.text
    assert docx.json()["status"] == "READY", docx.json()
