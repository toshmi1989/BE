"""Phase 11B — P0 DOCX coverage: sections, legacy removal, no raw enums."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.domain.canonical_sampling import get_canonical_sampling_plan
from app.domain.canonical_subjects import get_canonical_subject_counts
from app.domain.display_value_registry import find_raw_enums_in_text
from app.domain.docx_profile import ACTIVE_SECTION_CODES, TABLE_KEY_TO_INDEX
from app.domain.docx_renderer import _find_heading_indexes, _gate_mode, render_protocol_docx
from app.domain.protocol_assembly import assemble_protocol
from app.domain.protocol_sections import SECTION_TREE
from app.domain.static_blocks import legacy_uncontrolled_blocks
from app.domain.study_snapshot import build_canonical_snapshot


P0_CODES = (
    "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9",
    "2.2", "2.3", "2.4", "2.7", "2.8", "2.9", "2.10", "2.11", "2.12",
    "3",
    "4.3", "4.4", "4.4.1", "4.4.2", "4.5", "4.6", "4.7", "4.7.1", "4.7.2", "4.7.3",
    "4.8", "4.8.1", "4.8.2", "4.8.3", "4.9",
    "6.1", "6.1.1", "6.1.2", "6.1.3", "6.1.4", "6.1.5", "6.1.6", "6.1.7", "6.1.8",
    "6.1.9", "6.1.10", "6.2", "6.2.1", "6.2.2", "6.2.3", "6.3", "6.3.1",
    "7.1", "7.2", "7.3", "7.3.1", "7.3.2", "7.3.3", "7.3.4",
    "8.1", "8.2", "8.2.1", "8.2.2", "8.2.3", "8.3", "8.4", "8.5",
    "9.1", "9.2", "9.3", "9.4", "9.5", "9.6", "9.7", "9.7.1", "9.7.1.1", "9.7.1.2",
    "9.7.2", "9.7.3", "9.7.4",
    "18",
)

LEGACY_TEMPLATE_N = 46


def _ctx(**overrides):
    ctx = {
        "project_id": "p11b",
        "study": {
            "protocol_number": "BE-BOS-001",
            "title": "Bosutinib BE",
            "primary_objective": "Оценить биоэквивалентность",
        },
        "sponsor": {"name": "Sponsor LLC", "country": "RU"},
        "organizations": [
            {"name": "Site A", "role": "INVESTIGATOR", "address": "Moscow"},
            {"name": "Lab B", "role": "BIOANALYTICAL_LAB"},
            {"name": "CRO C", "role": "CRO"},
            {"name": "Expert D", "role": "MEDICAL_EXPERT"},
        ],
        "product": {
            "trade_name": "Бозутиниб",
            "inn": "bosutinib",
            "manufacturer": "ООО Тест",
            "dosage": "400 mg",
            "dosage_form": "tablet",
            "source_ids": ["s1"],
        },
        "reference_product": {
            "trade_name": "Бозулиф",
            "inn": "bosutinib",
            "manufacturer": "Pfizer",
            "dosage": "400 mg",
            "purchased_status": "UNKNOWN",
            "source_ids": ["s2"],
        },
        "design": {
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
            "blinding": "OPEN",
        },
        "food": {"condition": "FED", "meal_type": "HIGH_CALORIE", "calories": 900},
        "subjects": {
            "target_evaluable_n": 32,
            "planned_randomized_n": 36,
            "planned_screened_n": 36,
        },
        "sample_size": {"evaluable_n": 32, "randomized_n": 36, "selected_cv": 30},
        "sampling": {
            "points": [
                {"time_h": 0.0, "reason": "BASELINE"},
                {"time_h": 6.0, "reason": "TMAX_CAPTURE"},
                {"time_h": 72.0, "reason": "FINAL"},
            ],
            "total_points_per_period": 3,
        },
        "observation": {"selected_duration": 72, "unit": "h"},
        "washout": {"selected_value": 14, "unit": "day"},
        "analytes": [{"id": "a1", "name": "bosutinib", "type": "PARENT"}],
        "pk_parameters": [
            {"analyte_id": "a1", "parameter_code": "Cmax"},
            {"analyte_id": "a1", "parameter_code": "AUC0-t"},
        ],
        "eligibility": {
            "inclusion": [{"number": 1, "text": "Healthy volunteers"}],
            "non_inclusion": [{"number": 1, "text": "Pregnancy"}],
            "exclusion": [{"number": 1, "text": "AE leading to withdrawal"}],
        },
        "sources": [
            {
                "id": "s1",
                "type": "SmPC",
                "title": "Bosulif SmPC",
                "authors": ["Pfizer"],
                "year": 2023,
                "url": "https://example.com/smpc",
            }
        ],
        "evidence_claims": [
            {
                "field_name": "clinical_summary",
                "value": "Bosutinib studied in healthy volunteers.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
            {
                "field_name": "risk_benefit",
                "value": "Risk-benefit acceptable for HV BE study.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
            {
                "field_name": "pharmacology",
                "value": "TKI; oral absorption.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
        ],
        "statistical_config": {"alpha": 0.05, "be_lower": 0.8, "be_upper": 1.25},
        "cv_studies": [],
        "cv_selection": {"selected_cv": 30, "selection_method": "SINGLE_STUDY"},
        "evidence_summary": {"count": 3},
    }
    ctx.update(overrides)
    return ctx


def _section_text(sec: dict) -> str:
    parts: list[str] = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            parts.append(str(b["text"]))
        if b.get("items"):
            parts.extend(str(x) for x in b["items"])
        if b.get("display_text"):
            parts.append(str(b["display_text"]))
    return "\n".join(parts)


def _table_cell_texts(table: dict) -> list[str]:
    out: list[str] = []
    for row in table.get("rows") or []:
        for c in row:
            if c is None:
                continue
            out.append(str(c))
    return out


def _assert_no_legacy_n(blob: str, *, canonical_n: int, label: str) -> None:
    if canonical_n == LEGACY_TEMPLATE_N:
        return
    # whole-number token 46 must not appear in generated canonical content
    assert not re.search(r"(?<!\d)46(?!\d)", blob), f"{label}: legacy N=46 present; blob={blob[:400]!r}"


def _n_value_from_rows(rows: list, label_substr: str) -> int | None:
    for row in rows:
        if len(row) < 2:
            continue
        lab = str(row[0] or "").lower()
        if label_substr.lower() in lab:
            try:
                return int(row[1])
            except (TypeError, ValueError):
                return None
    return None


def test_p0_sections_in_tree_and_active() -> None:
    tree_codes = {s.section_code for s in SECTION_TREE}
    for code in P0_CODES:
        assert code in tree_codes, f"missing from SECTION_TREE: {code}"
        assert code in ACTIVE_SECTION_CODES, f"missing from ACTIVE_SECTION_CODES: {code}"


def test_p0_sections_generated_no_duplicates() -> None:
    out = assemble_protocol(_ctx(), blocking_validation=False)
    by_code = {}
    for sec in out["sections"]:
        code = sec["section_code"]
        assert code not in by_code, f"duplicate section {code}"
        by_code[code] = sec
    for code in P0_CODES:
        assert code in by_code
        assert by_code[code]["status"] in {"GENERATED", "UNRESOLVED", "SKIPPED"}
        assert by_code[code].get("content_blocks") is not None


def test_no_raw_enums_in_p0_text() -> None:
    out = assemble_protocol(_ctx(), blocking_validation=False)
    blob_parts = []
    for sec in out["sections"]:
        if sec["section_code"] not in P0_CODES and sec["section_code"] not in {"2.5", "2.6", "5.1", "SYNOPSIS"}:
            continue
        blob_parts.append(_section_text(sec))
    for t in out.get("tables") or []:
        blob_parts.extend(_table_cell_texts(t))
    blob = "\n".join(blob_parts)
    found = find_raw_enums_in_text(blob)
    assert not found, f"raw enums in protocol text: {found}"


def test_canonical_n_preserved() -> None:
    ctx = _ctx()
    snap = build_canonical_snapshot(ctx)
    subjects = get_canonical_subject_counts(ctx)
    assert subjects.randomized_n == 36
    assert subjects.evaluable_n == 32
    assert snap.randomized_n == 36
    assert snap.evaluable_n == 32

    out = assemble_protocol(ctx, blocking_validation=False)
    cons = out.get("consistency_snapshot") or {}
    assert cons.get("randomized_n") == 36
    assert cons.get("evaluable_n") == 32

    synopsis = next(s for s in out["sections"] if s["section_code"] == "SYNOPSIS")
    syn_text = _section_text(synopsis)
    assert f"рандомизированных: {subjects.randomized_n}" in syn_text
    assert f"Оцениваемых субъектов: {subjects.evaluable_n}" in syn_text
    _assert_no_legacy_n(syn_text, canonical_n=subjects.randomized_n, label="SYNOPSIS text")

    by_key = {t["table_key"]: t for t in out["tables"]}
    syn_n = by_key["SYNOPSIS_N"]
    assert _n_value_from_rows(syn_n["rows"], "Рандомизированные") == 36
    assert _n_value_from_rows(syn_n["rows"], "Оцениваемые") == 32
    _assert_no_legacy_n(
        " ".join(_table_cell_texts(syn_n)),
        canonical_n=36,
        label="SYNOPSIS_N table",
    )

    sample = next(s for s in out["sections"] if s["section_code"] == "9.2")
    stext = _section_text(sample)
    assert f"рандомизированных N={subjects.randomized_n}" in stext
    assert f"оцениваемых субъектов N={subjects.evaluable_n}" in stext
    _assert_no_legacy_n(stext, canonical_n=36, label="section 9.2")


def test_canonical_randomized_n_36_everywhere_no_legacy_46() -> None:
    """canonical randomized_n=36 → Synopsis / 9.2 / stats tables; never legacy 46."""
    ctx = _ctx()
    assert get_canonical_subject_counts(ctx).randomized_n == 36
    out = assemble_protocol(ctx, blocking_validation=False)

    synopsis = next(s for s in out["sections"] if s["section_code"] == "SYNOPSIS")
    syn_text = _section_text(synopsis)
    assert "рандомизированных: 36" in syn_text

    sec_92 = next(s for s in out["sections"] if s["section_code"] == "9.2")
    assert "рандомизированных N=36" in _section_text(sec_92)

    by_key = {t["table_key"]: t for t in out["tables"]}
    assert _n_value_from_rows(by_key["SYNOPSIS_N"]["rows"], "Рандомизированные") == 36
    assert _n_value_from_rows(by_key["STUDY_METADATA"]["rows"], "Randomized N") == 36

    generated_blobs = [
        syn_text,
        _section_text(sec_92),
        " ".join(_table_cell_texts(by_key["SYNOPSIS_N"])),
        " ".join(_table_cell_texts(by_key["STUDY_METADATA"])),
    ]
    for blob in generated_blobs:
        _assert_no_legacy_n(blob, canonical_n=36, label="N consistency")


def test_sampling_plan_matches_synopsis_and_docx_table(tmp_path: Path) -> None:
    ctx = _ctx()
    plan = get_canonical_sampling_plan(ctx)
    plan_times = [float(p["time_h"]) for p in plan.points]
    assert plan_times == [0.0, 6.0, 72.0]
    assert plan.points_per_period == 3

    out = assemble_protocol(ctx, blocking_validation=False)
    by_key = {t["table_key"]: t for t in out["tables"]}
    blood = by_key["BLOOD_SAMPLING"]
    draft_times = [float(row[1]) for row in blood["rows"]]
    assert draft_times == plan_times

    # Synopsis / sampling narrative must use the same point count (single source)
    sampling_sec = next(s for s in out["sections"] if s["section_code"] == "4.4.2")
    samp_text = _section_text(sampling_sec)
    assert str(plan.points_per_period) in samp_text
    assert any(b.get("table_key") == "BLOOD_SAMPLING" for b in sampling_sec.get("content_blocks") or [])

    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11b-samp",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    assert result.output_path is not None

    doc = Document(str(result.output_path))
    t10_idx = TABLE_KEY_TO_INDEX["BLOOD_SAMPLING"]
    docx_times = []
    for row in doc.tables[t10_idx].rows[1:]:  # skip header
        cells = row.cells
        if len(cells) < 2:
            continue
        raw = (cells[1].text or "").strip()
        if not raw:
            continue
        docx_times.append(float(raw.replace(",", ".")))
    assert docx_times == plan_times == draft_times


def test_no_raw_enums_in_generated_p0_docx(tmp_path: Path) -> None:
    ctx = _ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11b-enums",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))

    # Collect generated P0 section bodies (between mapped headings) + dynamic tables we overwrite
    heading_map = _find_heading_indexes(doc)
    parts: list[str] = []
    paras = list(doc.paragraphs)
    sorted_heads = sorted(heading_map.items(), key=lambda kv: kv[1])
    for i, (code, start) in enumerate(sorted_heads):
        if code not in P0_CODES and code not in {"2.5", "2.6", "5.1", "5.2", "5.3"}:
            continue
        end = sorted_heads[i + 1][1] if i + 1 < len(sorted_heads) else len(paras)
        for p in paras[start + 1 : end]:
            style = p.style.name if p.style else ""
            if style.startswith("Heading"):
                continue
            if p.text:
                parts.append(p.text)

    for key in ("BLOOD_SAMPLING", "PK_PARAMETERS", "TEST_PRODUCT", "REFERENCE_PRODUCT", "STUDY_METADATA", "SYNOPSIS_N"):
        idx = TABLE_KEY_TO_INDEX.get(key)
        if idx is None or idx >= len(doc.tables):
            continue
        for row in doc.tables[idx].rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    blob = "\n".join(parts)
    found = find_raw_enums_in_text(blob)
    assert not found, f"raw enums in generated P0 DOCX: {found}"


def test_evidence_missing_leaves_unresolved_not_invented() -> None:
    out = assemble_protocol(_ctx(evidence_claims=[]), blocking_validation=False)
    sec = next(s for s in out["sections"] if s["section_code"] == "2.2")
    texts = [str(b.get("text") or "") for b in sec.get("content_blocks") or []]
    joined = " ".join(texts)
    assert "{{EVIDENCE" in joined or sec["status"] == "UNRESOLVED"
    assert "bosutinib is safe and effective" not in joined.lower()


def test_dose_from_product_without_evidence() -> None:
    out = assemble_protocol(_ctx(evidence_claims=[]), blocking_validation=False)
    sec = next(s for s in out["sections"] if s["section_code"] == "2.4")
    text = " ".join(str(b.get("text") or "") for b in sec.get("content_blocks") or [])
    assert "400 mg" in text


def test_literature_uses_sources_provenance() -> None:
    out = assemble_protocol(_ctx(), blocking_validation=False)
    sec = next(s for s in out["sections"] if s["section_code"] == "18")
    items = []
    for b in sec.get("content_blocks") or []:
        items.extend(b.get("items") or [])
        if b.get("text"):
            items.append(b["text"])
    joined = " ".join(str(x) for x in items)
    assert "Bosulif" in joined or "SmPC" in joined
    assert "2023" in joined or "example.com" in joined


def test_tables_classified() -> None:
    out = assemble_protocol(_ctx(), blocking_validation=False)
    by_key = {t["table_key"]: t for t in out["tables"]}
    assert by_key["SCHEDULE_OF_ASSESSMENTS"]["classification"] == "STATIC_VERIFIED"
    assert by_key["SCHEDULE_OF_ASSESSMENTS"]["fill_mode"] == "PRESERVE"
    assert by_key["MEAL_TIMING"]["fill_mode"] == "PRESERVE"
    assert "LAB_PARAMETERS" not in TABLE_KEY_TO_INDEX  # T09 static, not overwritten
    assert "SIGNATURES" in TABLE_KEY_TO_INDEX


def test_template_headings_resolve_p0() -> None:
    path = Path(__file__).resolve().parents[2] / "templates" / "protocol" / "BE_Protocol_Template_v2.0.docx"
    doc = Document(str(path))
    mapping = _find_heading_indexes(doc)
    for code in ("1.1", "2.2", "3", "4.3", "6.1.4", "7.1", "7.2", "7.3.1", "8.1", "9.2", "18"):
        assert code in mapping, f"heading not found for {code}"


def test_no_duplicate_p0_in_active() -> None:
    assert len(ACTIVE_SECTION_CODES) == len(set(ACTIVE_SECTION_CODES))


def test_legacy_synopsis_n46_controlled() -> None:
    blocks = legacy_uncontrolled_blocks()
    ids = {b.block_id for b in blocks}
    assert "LEGACY.STATS_N46" not in ids
    assert "LEGACY.SYNOPSIS_N46" not in ids
    from app.domain.static_blocks import list_static_blocks

    dyn = {b["block_id"] for b in list_static_blocks(status="DYNAMIC")}
    assert "DYN.SYNOPSIS_SUBJECT_N" in dyn


def test_final_blocks_without_sponsor_and_orgs() -> None:
    reasons = _gate_mode(
        "FINAL",
        draft_status="READY_FOR_REVIEW",
        unresolved=["{{SPONSOR.NAME}}", "{{INVESTIGATORS.DETAILS}}"],
        consistency={"evaluable_n": 32, "design": "CROSSOVER_2X2"},
        study={"protocol_number": "X"},
        product={"trade_name": "T"},
        reference={"trade_name": "R"},
        sponsor=None,
    )
    assert any("sponsor" in r for r in reasons)
    assert any("unresolved" in r for r in reasons)
