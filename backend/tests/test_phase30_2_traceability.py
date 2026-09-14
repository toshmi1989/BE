"""Phase 30.2 — ACCEPTED_CALCULATION leak + section traceability + DOCX scrub."""

from __future__ import annotations

import pytest
from docx import Document

from app.core.config import Settings, get_settings
from app.domain.display_value_registry import (
    find_raw_enums_in_text,
    resolve_display,
    scrub_technical_tokens_in_text,
)
from app.domain.docx_profile import DOCX_GENERATOR_VERSION
from app.domain.protocol_constants import PROTOCOL_GENERATOR_VERSION
from app.domain.protocol_tables import build_tables
from app.domain.protocol_traceability import (
    build_section_traceability,
    enrich_field_bindings_with_sources,
    scrub_docx_technical_enums,
)
from app.domain.research_evidence_models import RegisteredSource, ResearchClaim
from app.domain.research_evidence_store import clear_research_evidence_store, put_claim, put_source


def test_phase30_2_version():
    assert Settings().app_version == "0.35.2"
    assert PROTOCOL_GENERATOR_VERSION == "0.35.2"
    assert DOCX_GENERATOR_VERSION == "0.35.2"
    assert get_settings().app_version == "0.35.2"


def test_accepted_calculation_never_human_facing_raw():
    human = resolve_display("ACCEPTED_CALCULATION", context="selection_method")
    assert human == "Расчёт подтверждён"
    assert "ACCEPTED_CALCULATION" not in human
    assert "ACCEPTED_CALCULATION" in find_raw_enums_in_text("method ACCEPTED_CALCULATION used")


def test_accepted_calculation_scrubbed_from_text():
    text, replaced = scrub_technical_tokens_in_text("CV via ACCEPTED_CALCULATION")
    assert "ACCEPTED_CALCULATION" not in text
    assert "Расчёт подтверждён" in text
    assert "ACCEPTED_CALCULATION" in replaced


def test_protocol_tables_selection_method_human():
    tables = build_tables(
        {
            "cv_studies": [],
            "cv_selection": {
                "selected_cv": 25.0,
                "selection_method": "ACCEPTED_CALCULATION",
            },
        },
        {},
    )
    blob = " ".join(
        str(cell) for t in tables for row in (t.get("rows") or []) for cell in row
    )
    assert "ACCEPTED_CALCULATION" not in blob
    assert "Расчёт подтверждён" in blob


def test_docx_scrub_removes_accepted_calculation():
    doc = Document()
    doc.add_paragraph("Selection: ACCEPTED_CALCULATION")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ACCEPTED_CALCULATION"
    table.rows[0].cells[1].text = "ok"
    result = scrub_docx_technical_enums(doc)
    assert "ACCEPTED_CALCULATION" in result["replaced"]
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    blob = "\n".join(parts)
    assert "ACCEPTED_CALCULATION" not in blob
    assert "Расчёт подтверждён" in blob


def test_section_traceability_emits_claim_source_location():
    clear_research_evidence_store()
    src = RegisteredSource(
        locator="smpc://upadacitinib",
        source_type="SmPC",
        title="SmPC Upadacitinib",
        study_id="UPDCB-TRACE-01",
    )
    put_source(src)
    put_claim(
        ResearchClaim(
            claim_text="Tmax 2-4 h",
            field_path="pk.expected_tmax",
            value="2–4 h",
            source_id=src.id,
            excerpt="Tmax is 2 to 4 hours",
            location="section 5.2",
            verification_status="VERIFIED",
            extraction_method="AI",
            confidence="HIGH",
            study_id="UPDCB-TRACE-01",
        )
    )
    rows = build_section_traceability(
        "UPDCB-TRACE-01",
        facts={"pk.expected_tmax": "2–4 h"},
    )
    assert rows
    hit = next(r for r in rows if r["canonical_field"] == "pk.expected_tmax")
    assert hit["protocol_section_id"] in {"PK", "SAMPLING"}
    assert hit["claim_id"]
    assert hit["source_id"] == src.id
    assert hit["source_label"]
    assert hit["location"]
    assert hit["provenance"]["expert_verified"] is True
    assert hit["provenance"]["ai_proposal"] is False
    assert "Verified by expert" in (hit["provenance"]["labels"]["expert_verified"] or "")

    bindings = enrich_field_bindings_with_sources(
        {
            "pk.expected_tmax": {
                "section": "PK",
                "label": "Tmax",
                "value": "2–4 h",
            }
        },
        rows,
    )
    assert bindings["pk.expected_tmax"]["source"]["claim_id"]
    assert bindings["pk.expected_tmax"]["source"]["status"] == "VERIFIED"


def test_ai_proposal_provenance_distinct_from_verified():
    clear_research_evidence_store()
    put_claim(
        ResearchClaim(
            claim_text="mechanism proposed",
            field_path="product.pharmacology.mechanism",
            value="JAK inhibitor",
            source_id="SRC-x",
            excerpt="JAK",
            verification_status="PROPOSED",
            extraction_method="AI",
            confidence="MEDIUM",
            study_id="UPDCB-TRACE-02",
        )
    )
    rows = build_section_traceability("UPDCB-TRACE-02", facts={})
    hit = next(r for r in rows if r["canonical_field"] == "product.pharmacology.mechanism")
    assert hit["provenance"]["ai_proposal"] is True
    assert hit["provenance"]["expert_verified"] is False
    assert hit["provenance"]["ai_confidence"] == "MEDIUM"


def test_missing_formula_not_invented_in_traceability():
    clear_research_evidence_store()
    rows = build_section_traceability(
        "UPDCB-TRACE-03",
        facts={"product.name": "X"},
    )
    chem = [r for r in rows if "chemistry" in (r["canonical_field"] or "")]
    assert chem == []


@pytest.mark.parametrize(
    "token",
    ["ACCEPTED_CALCULATION", "EXPERT_INPUT", "EXPLICIT_CONFIGURATION"],
)
def test_technical_tokens_mapped(token: str):
    label = resolve_display(token, context="selection_method")
    assert label
    assert token not in label
    assert token in find_raw_enums_in_text(f"x {token} y")
