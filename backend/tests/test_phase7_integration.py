"""Phase 7 AI extraction integration with MockAIProvider — no network LLM."""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient


def test_ai_status_disabled_by_default(client: TestClient) -> None:
    resp = client.get("/api/ai/status")
    assert resp.status_code == 200
    body = resp.json()
    assert body["enabled"] is False
    assert body["available"] is False


def test_disabled_ai_extract_fails_safely(client: TestClient) -> None:
    pid = client.post("/api/projects", json={"name": "ai-off"}).json()["id"]
    resp = client.post(
        f"/api/projects/{pid}/ai/extract",
        json={"task_type": "EXTRACT_PK", "force_mock": False},
    )
    assert resp.status_code == 422


def test_document_search_mock_ai_verify_apply(client: TestClient, tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    from app.core.config import get_settings

    get_settings.cache_clear()

    project = client.post("/api/projects", json={"name": "ai-flow"}).json()
    pid = project["id"]

    analyte = client.post(
        f"/api/projects/{pid}/analytes",
        json={
            "name": "bosutinib",
            "type": "PARENT",
            "tmax_min": 1,
            "tmax_max": 1,
            "half_life_min": 10,
            "half_life_max": 10,
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    ).json()

    client.post(f"/api/projects/{pid}/research-case", json={})

    buf = BytesIO()
    doc = DocxDocument()
    doc.add_paragraph(
        "Tmax was 2-3 hours. Elimination half-life ranged from 35 to 50 hours. "
        "Within-subject coefficient of variation for Cmax was 28.4%. "
        "The reference product should be administered with food."
    )
    # second conflicting tmax source chunk via second paragraph (same doc)
    doc.add_paragraph("In another study Tmax was 4 hours under fed conditions.")
    doc.save(buf)

    up = client.post(
        f"/api/projects/{pid}/documents",
        files={
            "file": (
                "pk.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"source_type": "SmPC"},
    )
    assert up.status_code == 201, up.text
    doc_id = up.json()["id"]
    assert client.post(f"/api/projects/{pid}/documents/{doc_id}/ingest").status_code == 200

    extract = client.post(
        f"/api/projects/{pid}/ai/extract",
        json={"task_type": "EXTRACT_PK", "force_mock": True, "query": "Tmax half-life"},
    )
    assert extract.status_code == 200, extract.text
    body = extract.json()
    assert body["run"]["status"] in {"NEEDS_REVIEW", "COMPLETED"}
    assert body["run"]["prompt_version"].startswith("EXTRACT_PK")
    assert body["run"]["provider"] == "mock"
    assert body["claims"]
    assert all(c["status"] == "PROPOSED" for c in body["claims"])
    assert all(c["origin"] == "AI_PROPOSED" for c in body["claims"])
    assert all(c["evidence_text"] for c in body["claims"])
    assert all(c["confidence"] > 0 for c in body["claims"])

    proposed = client.get(f"/api/projects/{pid}/ai/proposed")
    assert proposed.status_code == 200
    assert proposed.json()

    tmax_claim = next(c for c in body["claims"] if c["field_name"] == "tmax")
    claim_id = tmax_claim["claim_id"]

    # reject path
    cv_extract = client.post(
        f"/api/projects/{pid}/ai/extract",
        json={"task_type": "EXTRACT_CV", "force_mock": True},
    )
    assert cv_extract.status_code == 200
    if cv_extract.json()["claims"]:
        reject_id = cv_extract.json()["claims"][0]["claim_id"]
        rejected = client.post(
            f"/api/projects/{pid}/ai/claims/{reject_id}/review",
            json={"action": "reject"},
        )
        assert rejected.status_code == 200
        assert rejected.json()["status"] == "REJECTED"

    # verify + apply
    verified = client.post(
        f"/api/projects/{pid}/ai/claims/{claim_id}/review",
        json={"action": "verify", "analyte_id": analyte["id"]},
    )
    assert verified.status_code == 200, verified.text
    assert verified.json()["status"] == "VERIFIED"
    assert verified.json()["apply"]["applied"]

    refreshed = client.get(f"/api/projects/{pid}").json()
    a = refreshed["analytes"][0]
    assert a["tmax_min"] == 2
    assert a["tmax_max"] == 3
    assert a["provenance"]["origin"] == "SOURCE_DERIVED"

    runs = client.get(f"/api/projects/{pid}/ai/runs")
    assert runs.status_code == 200
    assert len(runs.json()) >= 1
