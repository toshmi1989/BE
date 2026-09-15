"""Phase 29 — Workspace → template DOCX (reuse render_protocol_docx).

Acceptance: no second renderer; BE_Protocol_Template_v2.0; no Bosutinib leftovers
for non-bosutinib studies; artifact binds snapshot/decision/SS/stats versions.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from sqlalchemy.orm import Session

from app.core.config import Settings, get_settings
from app.core.db import Base, configure_engine, get_engine
from app.core import db as db_module
from app.domain.docx_profile import DOCX_GENERATOR_VERSION, TEMPLATE_VERSION, get_template_profile
from app.domain.docx_renderer import render_protocol_docx
from app.domain.docx_stale_scrub import scrub_stale_template_products
from app.domain.exceptions import ValidationError
from app.domain.protocol_assembly import assemble_protocol
from app.domain.protocol_constants import PROTOCOL_GENERATOR_VERSION
from app.domain.protocol_template_registry import (
    STALE_TEMPLATE_PRODUCT_TOKENS,
    registry_as_list,
    study_is_bosutinib_sample,
)
from app.domain.decision_store import clear_decision_store
from app.domain.research_evidence_store import clear_research_evidence_store
from app.domain.sample_size_store import reset_sample_size_store
from app.domain.statistics_store import reset_statistics_store
from app.domain.study_input_store import clear_store as clear_study_input
from app.domain.study_workspace import put_protocol_draft_version, reset_workspace_store
from app.domain.workspace_assembly_context import (
    build_workspace_assembly_context,
    workspace_docx_blockers,
)
from app.domain.workspace_protocol import generate_docx_artifact, read_artifact_bytes
from app.models.workspace_persistence import WorkspaceSnapshotRecord
import app.models  # noqa: F401

STUDY = "UPDCB-02-BE-2026"


def _minimal_ctx(*, product_name: str = "Upadacitinib") -> dict:
    return {
        "project_id": STUDY,
        "study_id": STUDY,
        "workspace": True,
        "study": {
            "protocol_number": STUDY,
            "title": f"BE {product_name}",
            "version": "1.0",
        },
        "product": {
            "trade_name": product_name,
            "inn": product_name.lower(),
            "dosage": "15 mg",
            "source_ids": [],
        },
        "reference_product": {
            "trade_name": "Reference",
            "inn": product_name.lower(),
            "dosage": "15 mg",
            "purchased_status": "UNKNOWN",
            "source_ids": [],
        },
        "design": {
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
            "treatments": ["T", "R"],
            "food_condition": "FED",
            "source_ids": [],
            "display": "2×2 crossover",
        },
        "food": {
            "condition": "FED",
            "meal_type": "HIGH_CALORIE",
            "calorie_target": None,
            "display": "после еды",
        },
        "subjects": {
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
        },
        "eligibility": {"inclusion": [], "non_inclusion": [], "exclusion": []},
        "analytes": [
            {
                "id": "a1",
                "name": product_name.lower(),
                "type": "PARENT",
                "tmax_min": 1,
                "tmax_max": 2,
                "tmax_unit": "h",
                "half_life_min": 9,
                "half_life_max": 14,
            }
        ],
        "pk_parameters": [
            {"parameter_code": "Cmax", "role": "PRIMARY_BE"},
            {"parameter_code": "AUC0-t", "role": "PRIMARY_BE"},
        ],
        "washout": {"selected_value": 7, "unit": "day"},
        "observation": {"selected_duration": 48, "selected_unit": "h"},
        "sampling": {"points": [{"time_h": 0}, {"time_h": 1}, {"time_h": 2}]},
        "sample_size": {
            "id": "SS-test",
            "evaluable_n": 24,
            "randomized_n": 28,
            "dropout_pct": 15,
            "cv_used": 30,
            "parameter": "Cmax",
            "power": 0.8,
            "alpha": 0.05,
            "status": "ACCEPTED",
        },
        "cv_selection": {"selected_cv": 30, "parameter": "Cmax"},
        "cv_studies": [],
        "sources": [],
        "blood_volume": {},
        "statistical_config": {
            "id": "ST-test",
            "status": "APPROVED",
            "primary_parameters": ["Cmax", "AUC0-t"],
            "confidence_level": 0.9,
        },
        "evidence_summary": {"count": 0, "workspace": True},
        "sponsor": {"name": "Test Sponsor"},
        "organizations": [],
        "persons": [],
        "approved_decisions": {},
        "fact_sources": {},
        "structured_facts": {},
        "is_bosutinib_template_sample": study_is_bosutinib_sample(
            {"trade_name": product_name, "inn": product_name.lower()}
        ),
    }


def _xml_stale_hits(path: Path) -> dict[str, dict[str, int]]:
    out: dict[str, dict[str, int]] = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(z.read(name))
            blob = "".join((t.text or "") for t in root.xpath("//*[local-name()='t']"))
            hits = {tok: blob.count(tok) for tok in STALE_TEMPLATE_PRODUCT_TOKENS if tok in blob}
            if hits:
                out[name] = hits
    return out


@pytest.fixture(autouse=True)
def _clean(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "sqlite+pysqlite:///:memory:")
    monkeypatch.setenv("AI_ENABLED", "false")
    monkeypatch.setenv("AUTH_REQUIRED", "false")
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    get_settings.cache_clear()
    configure_engine("sqlite+pysqlite:///:memory:")
    engine = get_engine()
    Base.metadata.create_all(bind=engine)
    reset_workspace_store()
    reset_sample_size_store()
    reset_statistics_store()
    clear_research_evidence_store()
    clear_decision_store()
    clear_study_input()
    yield
    reset_workspace_store()
    reset_sample_size_store()
    reset_statistics_store()
    clear_research_evidence_store()
    clear_decision_store()
    clear_study_input()
    Base.metadata.drop_all(bind=engine)
    get_settings.cache_clear()


def test_phase29_version():
    assert Settings().app_version == "0.35.5"
    assert PROTOCOL_GENERATOR_VERSION == "0.35.5"
    assert DOCX_GENERATOR_VERSION == "0.35.5"
    assert get_settings().app_version == "0.35.5"


def test_template_registry_covers_dynamic_fields():
    fields = registry_as_list()
    assert len(fields) >= 10
    ids = {f["template_id"] for f in fields}
    assert "header.test_product" in ids
    assert "T05.test_product" in ids
    assert "T06.reference_product" in ids
    assert "T17.cv_sample_size" in ids
    assert any(f["required"] for f in fields)
    assert "Бозутиниб" in STALE_TEMPLATE_PRODUCT_TOKENS
    assert "Bosulif" in STALE_TEMPLATE_PRODUCT_TOKENS


def test_workspace_docx_blockers_missing_product():
    ctx = _minimal_ctx()
    ctx["product"] = {"trade_name": None, "inn": None}
    blockers = workspace_docx_blockers(ctx, study_id=STUDY)
    codes = {b["code"] for b in blockers}
    assert "MISSING_TEST_PRODUCT" in codes


def test_workspace_docx_blockers_dose_conflict():
    ctx = _minimal_ctx()
    ctx["reference_product"]["dose_conflict"] = True
    blockers = workspace_docx_blockers(ctx, study_id=STUDY)
    assert any(b["code"] == "UNRESOLVED_REFERENCE_DOSE_CONFLICT" for b in blockers)


def test_scrub_clears_header_table_bosutinib(tmp_path: Path):
    profile = get_template_profile()
    src = profile.template_path()
    assert src.is_file()
    out = tmp_path / "copy.docx"
    out.write_bytes(src.read_bytes())
    doc = Document(str(out))
    # Confirm sample product sits in header table before scrub
    header_blob = "\n".join(
        cell.text
        for table in doc.sections[0].header.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Бозутиниб" in header_blob
    scrub = scrub_stale_template_products(doc, _minimal_ctx(product_name="Upadacitinib"))
    doc.save(str(out))
    assert scrub["remaining"] == 0
    assert scrub["scrubbed"] > 0
    assert not _xml_stale_hits(out)


def test_assemble_render_full_protocol_structure(tmp_path: Path):
    ctx = _minimal_ctx()
    assembled = assemble_protocol(
        ctx,
        blocking_validation=False,
        validation_issues=[],
        rules_version="workspace-1",
        protocol_version="1",
    )
    assert len(assembled.get("sections") or []) >= 50
    payload = {
        "status": "DRAFT",
        "protocol_version": "1",
        "template_version": TEMPLATE_VERSION,
        "rules_version": "workspace-1",
        "generator_version": PROTOCOL_GENERATOR_VERSION,
        "consistency_snapshot": assembled.get("consistency_snapshot") or {},
        "canonical_fingerprint": assembled.get("canonical_fingerprint"),
        "sections": assembled.get("sections") or [],
        "tables": assembled.get("tables") or [],
        "references": assembled.get("references") or [],
        "build_report": assembled.get("build_report") or {},
    }
    result = render_protocol_docx(
        protocol_payload=payload,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="updcb-p29",
        protocol_version="1",
    )
    assert result.status == "READY"
    assert result.output_path and result.output_path.is_file()
    doc = Document(str(result.output_path))
    scrub = scrub_stale_template_products(doc, ctx)
    assert scrub["remaining"] == 0
    doc.save(str(result.output_path))

    assert len(doc.tables) >= 30
    assert len(doc.paragraphs) >= 400
    assert result.output_path.stat().st_size > 500_000

    blob = "\n".join(p.text or "" for p in doc.paragraphs)
    assert "Protocol Draft —" not in blob
    assert "From approved statistics plan when present." not in blob
    assert "Upadacitinib" in blob or "upadacitinib" in blob.lower()
    assert not _xml_stale_hits(result.output_path)


def test_generate_docx_artifact_uses_canonical_renderer(tmp_path, monkeypatch):
    db: Session = db_module.SessionLocal()
    try:
        snap = WorkspaceSnapshotRecord(
            snapshot_id=f"SNAP-{STUDY}-1",
            study_key=STUDY,
            version=1,
            status="ACTIVE",
            created_by="p29",
            payload={
                "candidates": [
                    {"field_path": "test_product.name", "value": "Upadacitinib"},
                    {"field_path": "test_product.inn", "value": "upadacitinib"},
                    {"field_path": "study.protocol_number", "value": STUDY},
                ]
            },
            based_on_decision_ids=["PD-1"],
            content_hash="abc",
        )
        db.add(snap)
        db.commit()

        put_protocol_draft_version(
            STUDY,
            status="DRAFT",
            created_by="p29",
            based_on={
                "snapshot": snap.snapshot_id,
                "decisions": ["PD-1"],
                "statistics": "ST-test",
                "sample_size": "SS-test",
            },
            sections_summary=["GENERAL"],
        )

        monkeypatch.setattr(
            "app.domain.workspace_protocol.build_preflight",
            lambda _sid: {"critical_blockers": [], "warnings": [], "can_generate_docx": True},
        )
        monkeypatch.setattr(
            "app.domain.workspace_protocol.detect_stale_protocol_dependencies",
            lambda *_a, **_k: {"stale": False, "reasons": []},
        )
        monkeypatch.setattr(
            "app.domain.workspace_protocol.build_workspace_assembly_context",
            lambda *_a, **_k: _minimal_ctx(),
        )
        monkeypatch.setattr(
            "app.domain.workspace_protocol.workspace_docx_blockers",
            lambda *_a, **_k: [],
        )

        art = generate_docx_artifact(db, STUDY, created_by="p29", force_warnings_ok=True)
        assert art["sha256"]
        assert art["snapshot_id"] == snap.snapshot_id
        assert art["statistics_version"] == "ST-test"
        assert art["sample_size_version"] == "SS-test"
        assert art["renderer"] == "docx_renderer.render_protocol_docx"
        assert art["template_version"] == TEMPLATE_VERSION
        assert art["generator_version"] == DOCX_GENERATOR_VERSION
        assert art["rebuild_on_download"] is False
        assert art["sections_assembled"] >= 50
        assert art["tables_assembled"] >= 1
        assert art["scrub"]["remaining"] == 0

        data, meta = read_artifact_bytes(db, art["artifact_id"])
        assert len(data) > 500_000
        assert meta["sha256"] == art["sha256"]
        out = tmp_path / "art.docx"
        out.write_bytes(data)
        assert not _xml_stale_hits(out)
        doc = Document(str(out))
        assert len(doc.tables) >= 30
        blob = "\n".join(p.text or "" for p in doc.paragraphs)
        assert "Protocol Draft —" not in blob
    finally:
        db.close()


def test_generate_docx_blocks_when_product_missing(monkeypatch):
    db: Session = db_module.SessionLocal()
    try:
        put_protocol_draft_version(
            STUDY,
            status="DRAFT",
            created_by="p29",
            based_on={"snapshot": None, "decisions": []},
        )
        monkeypatch.setattr(
            "app.domain.workspace_protocol.build_preflight",
            lambda _sid: {"critical_blockers": [], "warnings": [], "can_generate_docx": True},
        )
        monkeypatch.setattr(
            "app.domain.workspace_protocol.detect_stale_protocol_dependencies",
            lambda *_a, **_k: {"stale": False},
        )
        empty = _minimal_ctx()
        empty["product"] = {"trade_name": None, "inn": None}
        monkeypatch.setattr(
            "app.domain.workspace_protocol.build_workspace_assembly_context",
            lambda *_a, **_k: empty,
        )
        with pytest.raises(ValidationError) as ei:
            generate_docx_artifact(db, STUDY, created_by="p29", force_warnings_ok=True)
        assert ei.value.field == "workspace_docx"
        assert any(b["code"] == "MISSING_TEST_PRODUCT" for b in ei.value.details["blockers"])
    finally:
        db.close()


def test_assembly_context_does_not_require_legacy_project():
    # Empty stores → still returns workspace-shaped ctx (product may be empty)
    ctx = build_workspace_assembly_context(STUDY)
    assert ctx.get("workspace") is True
    assert ctx.get("study_id") == STUDY
    assert "product" in ctx
    assert "Legacy" not in str(ctx.get("evidence_summary") or {})
