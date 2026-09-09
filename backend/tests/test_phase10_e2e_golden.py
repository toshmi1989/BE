"""Phase 10 end-to-end golden Bosutinib flow — no invented regulatory values."""

from __future__ import annotations

import io
import time
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.domain.docx_profile import TEMPLATE_CHECKSUM_SHA256, get_template_profile
from app.domain.document_ingest import sha256_hex
from app.domain.docx_validation import validate_docx_file
from tests.phase10_golden_helpers import (
    GOLDEN_ROOT,
    TEMPLATE_EXCERPTS,
    build_fixture_docx_bytes,
    copy_file,
    new_run_dir,
    write_json,
)


def _perf() -> dict:
    return {}


def test_phase10_e2e_ai_disabled_golden(client: TestClient, tmp_path, monkeypatch) -> None:
    """Minimal client input → research → study → protocol → DOCX (AI off)."""
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    monkeypatch.setenv("AI_ENABLED", "false")
    from app.core.config import get_settings

    get_settings.cache_clear()

    timings: dict[str, float] = {}
    run_dir = new_run_dir("bosutinib-ai-off")
    missing: list[str] = []
    notes: list[str] = []

    # --- 1. Minimal client input ---
    t0 = time.perf_counter()
    project = client.post("/api/projects", json={"name": "Golden Bosutinib 400mg"}).json()
    pid = project["id"]
    ci = client.post(
        f"/api/projects/{pid}/client-input",
        json={
            "requested_product_name": "Бозутиниб",
            "inn": "bosutinib",
            "dosage": "400 mg",
            "dosage_form": "tablet",
            "route": "oral",
            "requested_subject_count": 28,
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    )
    assert ci.status_code in {200, 201}, ci.text
    timings["minimal_client_input_s"] = time.perf_counter() - t0

    # --- 2. Research profile + tasks ---
    t0 = time.perf_counter()
    assert client.post(f"/api/projects/{pid}/research-case", json={}).status_code in {200, 201}
    assert (
        client.post(
            f"/api/projects/{pid}/research-profile",
            json={"sync_from_client_input": True, "regulatory_jurisdiction": "EAEU"},
        ).status_code
        == 200
    )
    tasks = client.post(f"/api/projects/{pid}/research-case/tasks/generate", json={})
    assert tasks.status_code == 200, tasks.text
    assert len(tasks.json()) >= 1
    timings["research_tasks_s"] = time.perf_counter() - t0

    # --- 3. Document upload (template excerpts only) ---
    t0 = time.perf_counter()
    fixture = build_fixture_docx_bytes()
    (run_dir / "fixture-excerpts.docx").write_bytes(fixture)
    up = client.post(
        f"/api/projects/{pid}/documents",
        files={
            "file": (
                "bosutinib_template_excerpts.docx",
                fixture,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"source_type": "SmPC"},
    )
    assert up.status_code == 201, up.text
    doc_id = up.json()["id"]
    source_id = up.json().get("source_id")
    assert client.post(f"/api/projects/{pid}/documents/{doc_id}/ingest").status_code == 200
    timings["document_ingest_s"] = time.perf_counter() - t0

    if not source_id:
        detail = client.get(f"/api/projects/{pid}").json()
        sources = detail.get("sources") or []
        assert sources, "expected source from upload"
        source_id = sources[0]["id"]

    # Search
    hits = client.post(
        f"/api/projects/{pid}/research/search",
        json={"query": "tmax t1/2 Cmax вариабельность бозутиниб"},
    )
    assert hits.status_code == 200

    # --- 4. Evidence from template excerpts (manual research API) ---
    t0 = time.perf_counter()
    ev_tmax = client.post(
        f"/api/projects/{pid}/research/evidence",
        json={
            "source_id": source_id,
            "document_id": doc_id,
            "page": 1,
            "field_code": "tmax",
            "extracted_text": "tmax бозутиниба в среднем составляет 6 часов",
            "auto_detect_conflicts": True,
        },
    )
    assert ev_tmax.status_code == 201, ev_tmax.text
    tmax_claim_id = ev_tmax.json()["claims"][0]["id"]

    ev_hl = client.post(
        f"/api/projects/{pid}/research/evidence",
        json={
            "source_id": source_id,
            "document_id": doc_id,
            "page": 1,
            "field_code": "half_life",
            "extracted_text": "t1/2 бозутиниба составляет 35,5 часов",
            "auto_detect_conflicts": True,
        },
    )
    assert ev_hl.status_code == 201, ev_hl.text
    hl_claim_id = ev_hl.json()["claims"][0]["id"]

    ev_cv = client.post(
        f"/api/projects/{pid}/research/evidence",
        json={
            "source_id": source_id,
            "document_id": doc_id,
            "page": 1,
            "field_code": "cv_cmax",
            "extracted_text": "Cmax составляет >30%",
            "auto_detect_conflicts": True,
        },
    )
    assert ev_cv.status_code == 201, ev_cv.text
    notes.append(
        "CV source text states '>30%'; planning CV=30 used as floor with expert note — not a precise published CVintra."
    )
    timings["evidence_create_s"] = time.perf_counter() - t0

    # Conflicts refresh
    client.post(f"/api/projects/{pid}/research/conflicts", json={})

    # --- 5. Study build (known template facts) ---
    t0 = time.perf_counter()
    missing.append("sponsor (study sponsor not stated as project sponsor in excerpts)")
    missing.append("investigator / clinical site contacts")
    missing.append("exact registration number")
    missing.append("exact CVintra point estimate (only >30%)")

    client.put(
        f"/api/projects/{pid}/study",
        json={
            "protocol_number": "BE-BOS-400-GOLDEN",
            "title": "Исследование биоэквивалентности бозутиниба 400 мг",
            "short_title": "Bosutinib 400 mg BE",
            "version": "0.1-golden",
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    )
    # Do NOT invent sponsor → FINAL must block later
    client.put(
        f"/api/projects/{pid}/product",
        json={
            "trade_name": "Бозутиниб",
            "inn": "bosutinib",
            "dosage": "400 mg",
            "dosage_form": "tablet",
            "provenance": {
                "origin": "SOURCE_DERIVED",
                "status": "PROPOSED",
                "source_ids": [source_id],
                "notes": "From template excerpt",
            },
        },
    )
    client.put(
        f"/api/projects/{pid}/reference-product",
        json={
            "trade_name": "Бозулиф",
            "inn": "bosutinib",
            "dosage": "400 mg",
            "purchased_status": "UNKNOWN",
            "provenance": {
                "origin": "SOURCE_DERIVED",
                "status": "PROPOSED",
                "source_ids": [source_id],
                "notes": "Bosulif named in template excerpt; purchased_status unknown",
            },
        },
    )
    client.post(
        f"/api/projects/{pid}/design",
        json={
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
            "treatments": ["T", "R"],
            "food_condition": "FED",
            "decision_status": "PROPOSED",
            "provenance": {"origin": "SOURCE_DERIVED", "status": "PROPOSED", "source_ids": [source_id]},
        },
    )
    client.post(
        f"/api/projects/{pid}/food",
        json={
            "condition": "FED",
            "meal_type": "HIGH_CALORIE",
            "decision_status": "PROPOSED",
            "provenance": {"origin": "SOURCE_DERIVED", "status": "PROPOSED", "source_ids": [source_id]},
        },
    )
    client.put(
        f"/api/projects/{pid}/eligibility",
        json={
            "inclusion": [{"number": 1, "text": "Healthy volunteers"}],
            "non_inclusion": [],
            "exclusion": [],
        },
    )
    analyte = client.post(
        f"/api/projects/{pid}/analytes",
        json={
            "name": "bosutinib",
            "type": "PARENT",
            "tmax_min": 6,
            "tmax_max": 6,
            "tmax_unit": "h",
            "half_life_min": 35.5,
            "half_life_max": 35.5,
            "half_life_unit": "h",
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    ).json()
    aid = analyte["id"]

    # Expert verify + apply via foundation EvidenceCreate (VERIFIED claims)
    verified_bundle = client.post(
        f"/api/projects/{pid}/research-case/evidence",
        json={
            "source_id": source_id,
            "evidence_type": "PK",
            "claim": "tmax+half_life verified from template excerpts",
            "extracted_text": TEMPLATE_EXCERPTS[2] + " " + TEMPLATE_EXCERPTS[3],
            "page": 1,
            "confidence": 0.95,
            "verification_status": "VERIFIED",
            "claims": [
                {
                    "field_name": "tmax",
                    "value": "6 h",
                    "normalized_value": {"min": 6, "max": 6},
                    "unit": "h",
                    "confidence": 0.95,
                    "status": "VERIFIED",
                    "origin": "SOURCE_DERIVED",
                    "source_ids": [source_id],
                },
                {
                    "field_name": "half_life",
                    "value": "35.5 h",
                    "normalized_value": {"min": 35.5, "max": 35.5},
                    "unit": "h",
                    "confidence": 0.95,
                    "status": "VERIFIED",
                    "origin": "SOURCE_DERIVED",
                    "source_ids": [source_id],
                },
            ],
        },
    )
    assert verified_bundle.status_code == 201, verified_bundle.text
    claim_ids = [c["id"] for c in verified_bundle.json()["claims"]]
    applied = client.post(
        f"/api/projects/{pid}/research-case/apply-verified",
        json={"claim_ids": claim_ids, "analyte_id": aid},
    )
    assert applied.status_code == 200, applied.text
    applied_ids = applied.json().get("applied") or []
    assert len(applied_ids) >= 1, applied.text

    client.post(
        f"/api/projects/{pid}/pk/parameters",
        json={
            "analyte_id": aid,
            "parameter_code": "Tmax",
            "unit": "h",
            "range_min": 6,
            "range_max": 6,
            "provenance": {"origin": "SOURCE_DERIVED", "status": "PROPOSED", "source_ids": [source_id]},
        },
    )
    client.post(
        f"/api/projects/{pid}/washout/calculate",
        json={"selected_value": 14, "selected_unit": "day"},
    )  # template mentions 14 days / >5 t1/2 context
    client.post(
        f"/api/projects/{pid}/observation/calculate",
        json={"selected_duration": 72, "selected_unit": "h"},
    )  # template: 72 h sampling
    client.post(f"/api/projects/{pid}/sampling/recommend", json={"persist": True})
    client.post(
        f"/api/projects/{pid}/blood-volume/calculate",
        json={
            "blood_volume_per_pk_sample_ml": 5,
            "screening_volume_ml": 15,
            "safety_laboratory_volume_ml": 10,
            "subjects": 28,
        },
    )
    client.post(
        f"/api/projects/{pid}/subjects",
        json={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
            "reserve_n": 0,
        },
    )
    cv = client.post(
        f"/api/projects/{pid}/statistics/cv",
        json={
            "source_id": source_id,
            "analyte_id": aid,
            "parameter": "Cmax",
            "design": "CROSSOVER_2X2",
            "condition": "FED",
            "n_total": None,
            "n_be_analysis": None,
            "cv_value": 30,
            "evidence": [
                {
                    "source_id": source_id,
                    "extracted_text": TEMPLATE_EXCERPTS[1],
                    "status": "PROPOSED",
                }
            ],
            "provenance": {
                "origin": "SOURCE_DERIVED",
                "status": "PROPOSED",
                "source_ids": [source_id],
                "notes": "Planning floor from template '>30%'; not exact CVintra",
            },
        },
    )
    assert cv.status_code == 201, cv.text
    client.post(
        f"/api/projects/{pid}/statistics/cv/select",
        json={
            "selection_method": "SINGLE_STUDY",
            "selected_study_id": cv.json()["id"],
            "cv_study_ids": [cv.json()["id"]],
        },
    )
    client.post(
        f"/api/projects/{pid}/statistics/sample-size",
        json={"design_type": "CROSSOVER_2X2", "dropout_pct": 10},
    )
    timings["study_build_s"] = time.perf_counter() - t0

    # --- 6. Validation ---
    t0 = time.perf_counter()
    val = client.post(f"/api/projects/{pid}/validate")
    assert val.status_code == 200, val.text
    summary = val.json()["summary"]
    snap = val.json()["canonical_snapshot"]
    write_json(run_dir / "validation-report.json", val.json())
    write_json(run_dir / "study-snapshot.json", snap)
    timings["validation_s"] = time.perf_counter() - t0
    # May have warnings (sponsor etc.) but preferably no CRITICAL/ERROR blocking if complete
    # If blocking — still continue; FINAL will reflect gate
    validation_blocking = bool(summary.get("blocking"))

    # --- 7. Protocol assembly (deterministic) ---
    t0 = time.perf_counter()
    p1 = client.post(f"/api/projects/{pid}/protocol/build", json={})
    assert p1.status_code == 200, p1.text
    draft1 = p1.json()
    p2 = client.post(f"/api/projects/{pid}/protocol/build", json={})
    assert p2.status_code == 200
    draft2 = p2.json()
    assert draft1["canonical_fingerprint"] == draft2["canonical_fingerprint"]
    write_json(
        run_dir / "protocol-build-report.json",
        draft1.get("build_report") or {},
    )
    write_json(
        run_dir / "protocol-draft-meta.json",
        {
            "status": draft1["status"],
            "fingerprint": draft1["canonical_fingerprint"],
            "sections": len(draft1["sections"]),
            "tables": len(draft1["tables"]),
            "generator_version": draft1["generator_version"],
        },
    )
    timings["protocol_assembly_s"] = time.perf_counter() - t0

    # --- 8. DOCX modes ---
    t0 = time.perf_counter()
    draft_docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "DRAFT", "ensure_protocol": False},
    )
    assert draft_docx.status_code == 200, draft_docx.text
    draft_body = draft_docx.json()
    timings["docx_draft_s"] = time.perf_counter() - t0

    # Save DRAFT artifact immediately (later REVIEW/FINAL may become "latest")
    assert draft_body["status"] == "READY", draft_body
    dl = client.get(f"/api/projects/{pid}/protocol/docx/download")
    assert dl.status_code == 200, dl.text
    draft_path = run_dir / "protocol-draft.docx"
    draft_path.write_bytes(dl.content)
    Document(io.BytesIO(dl.content))
    qa = validate_docx_file(draft_path, allow_unresolved=True)
    write_json(run_dir / "docx-draft-validation.json", qa.to_dict())
    assert qa.table_count == 33
    assert qa.paragraph_count > 0

    t0 = time.perf_counter()
    review_docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "REVIEW", "ensure_protocol": False},
    ).json()
    timings["docx_review_s"] = time.perf_counter() - t0

    t0 = time.perf_counter()
    final_docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "FINAL", "ensure_protocol": False},
    ).json()
    timings["docx_final_s"] = time.perf_counter() - t0

    # FINAL must block without sponsor (correct behavior)
    assert final_docx["status"] == "BLOCKED", final_docx
    notes.append(f"FINAL correctly BLOCKED: {final_docx.get('blocking_reasons')}")
    write_json(
        run_dir / "docx-final-status.json",
        {
            "status": final_docx["status"],
            "blocking_reasons": final_docx.get("blocking_reasons"),
            "note": "protocol-final.docx not written — correct gate without sponsor",
        },
    )
    write_json(
        run_dir / "docx-review-status.json",
        {"status": review_docx["status"], "blocking_reasons": review_docx.get("blocking_reasons")},
    )

    # Structural comparison vs template
    tpl = get_template_profile().template_path()
    assert sha256_hex(tpl.read_bytes()) == TEMPLATE_CHECKSUM_SHA256
    tpl_doc = Document(str(tpl))
    gen_doc = Document(str(draft_path))
    comparison = {
        "template_tables": len(tpl_doc.tables),
        "generated_tables": len(gen_doc.tables),
        "template_paragraphs": len(tpl_doc.paragraphs),
        "generated_paragraphs": len(gen_doc.paragraphs),
        "template_checksum_unchanged": True,
        "critical_values_present": {
            "protocol_number": "BE-BOS-400-GOLDEN" in "\n".join(p.text for p in gen_doc.paragraphs)
            or any(
                "BE-BOS-400-GOLDEN" in (c.text or "")
                for t in gen_doc.tables
                for r in t.rows
                for c in r.cells
            ),
            "bosutinib_or_product": any(
                "бозутиниб" in (p.text or "").lower() or "bosutinib" in (p.text or "").lower()
                for p in gen_doc.paragraphs
            ),
        },
        "byte_equality_required": False,
    }
    write_json(run_dir / "structural-comparison.json", comparison)

    # Visual QA limitation — no soffice/docx2pdf
    visual = {
        "pdf_render": False,
        "reason": "No LibreOffice/soffice/docx2pdf available in environment",
        "structural_proxy": {
            "opens_in_python_docx": True,
            "table_count": qa.table_count,
            "heading_count": qa.heading_count,
            "unresolved_count": len(qa.unresolved_found),
        },
    }
    write_json(run_dir / "visual-qa.json", visual)

    # Performance + run summary
    write_json(
        run_dir / "performance.json",
        {
            **timings,
            "docx_draft_threshold_s": 30,
            "docx_draft_exceeds_threshold": timings.get("docx_draft_s", 0) > 30,
            "background_job_recommended": timings.get("docx_draft_s", 0) > 30,
        },
    )
    write_json(
        run_dir / "e2e-summary.json",
        {
            "project_id": pid,
            "ai_enabled": False,
            "validation_blocking": validation_blocking,
            "validation_summary": summary,
            "protocol_status": draft1["status"],
            "docx_draft": draft_body["status"],
            "docx_review": review_docx["status"],
            "docx_final": final_docx["status"],
            "missing_data": missing,
            "notes": notes,
            "golden_dir": str(run_dir.relative_to(GOLDEN_ROOT.parent)),
        },
    )

    # Pointer to latest run (overwrite only the pointer file, not artifacts)
    write_json(
        GOLDEN_ROOT / "LATEST_AI_OFF.json",
        {"run_dir": str(run_dir.name), "project_id": pid},
    )


def test_phase10_e2e_ai_mock_proposed_only(client: TestClient, tmp_path, monkeypatch) -> None:
    """AI mock proposes evidence only — never auto-verifies Study."""
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    monkeypatch.setenv("AI_ENABLED", "false")
    from app.core.config import get_settings

    get_settings.cache_clear()

    run_dir = new_run_dir("bosutinib-ai-mock")
    project = client.post("/api/projects", json={"name": "Golden AI mock"}).json()
    pid = project["id"]
    client.post(
        f"/api/projects/{pid}/client-input",
        json={
            "requested_product_name": "Бозутиниб",
            "inn": "bosutinib",
            "dosage": "400 mg",
            "dosage_form": "tablet",
            "requested_subject_count": 28,
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    )
    client.post(f"/api/projects/{pid}/research-case", json={})
    fixture = build_fixture_docx_bytes()
    up = client.post(
        f"/api/projects/{pid}/documents",
        files={
            "file": (
                "excerpts.docx",
                fixture,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"source_type": "SmPC"},
    )
    assert up.status_code == 201
    doc_id = up.json()["id"]
    assert client.post(f"/api/projects/{pid}/documents/{doc_id}/ingest").status_code == 200

    analyte = client.post(
        f"/api/projects/{pid}/analytes",
        json={
            "name": "bosutinib",
            "type": "PARENT",
            "tmax_min": 1,
            "tmax_max": 1,
            "half_life_min": 10,
            "half_life_max": 10,
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    ).json()

    before = client.get(f"/api/projects/{pid}").json()["analytes"][0]

    extract = client.post(
        f"/api/projects/{pid}/ai/extract",
        json={"task_type": "EXTRACT_PK", "force_mock": True, "query": "Tmax half-life"},
    )
    # May be 200 with claims or empty — mock regex may not match Cyrillic template excerpts
    assert extract.status_code == 200, extract.text
    body = extract.json()
    assert body["run"]["provider"] == "mock"
    for c in body.get("claims") or []:
        assert c["status"] == "PROPOSED"
        assert c["origin"] == "AI_PROPOSED"

    after = client.get(f"/api/projects/{pid}").json()["analytes"][0]
    assert after["tmax_min"] == before["tmax_min"]
    assert after["half_life_min"] == before["half_life_min"]

    write_json(
        run_dir / "ai-mock-summary.json",
        {
            "claims": body.get("claims") or [],
            "study_unchanged_without_verify": True,
            "analyte_before": {"tmax_min": before["tmax_min"], "half_life_min": before["half_life_min"]},
            "analyte_after": {"tmax_min": after["tmax_min"], "half_life_min": after["half_life_min"]},
            "note": "Mock may yield 0 claims on Cyrillic-only excerpts; still proves AI path does not mutate Study",
        },
    )
    write_json(GOLDEN_ROOT / "LATEST_AI_MOCK.json", {"run_dir": str(run_dir.name)})
