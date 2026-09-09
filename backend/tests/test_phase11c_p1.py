"""Phase 11C — P1 administrative / organizational DOCX coverage."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.domain.docx_profile import ACTIVE_SECTION_CODES, TABLE_KEY_TO_INDEX
from app.domain.docx_renderer import _find_heading_indexes, render_protocol_docx
from app.domain.protocol_assembly import assemble_protocol
from app.domain.validation_engine import validate_administration


def _admin_ctx(**overrides):
    ctx = {
        "project_id": "p11c",
        "study": {
            "protocol_number": "BE-ADM-001",
            "title": "Admin coverage test",
            "country": "RU",
        },
        "sponsor": {
            "id": "sp1",
            "name": "ООО ФармаСпонсор",
            "address": "Москва, ул. Пример 1",
            "country": "RU",
            "contact_phone": "+7-495-000-00-01",
            "contact_email": "sponsor@example.com",
        },
        "organizations": [
            {
                "id": "o1",
                "name": "Клинический центр №1",
                "role": "CLINICAL_CENTER",
                "address": "Санкт-Петербург",
                "country": "RU",
            },
            {
                "id": "o2",
                "name": "BioLab Analytics",
                "role": "BIOANALYTICAL_LAB",
                "address": "Москва",
                "contact_email": "lab@example.com",
            },
            {"id": "o3", "name": "CRO Partner", "role": "CRO", "country": "RU"},
        ],
        "persons": [
            {
                "id": "p1",
                "role": "SPONSOR_AUTHORIZED",
                "full_name": "Иванов И.И.",
                "title": "Директор по клиническим исследованиям",
                "phone": "+7-495-111-11-11",
                "email": "ivanov@example.com",
            },
            {
                "id": "p2",
                "role": "MEDICAL_EXPERT",
                "full_name": "Петров П.П.",
                "title": "Медицинский эксперт",
            },
            {
                "id": "p3",
                "role": "PRINCIPAL_INVESTIGATOR",
                "full_name": "Сидоров С.С.",
                "title": "Главный исследователь",
                "organization_id": "o1",
                "organization_name": "Клинический центр №1",
            },
        ],
        "study_administration": {
            "id": "sa1",
            "insurance_provider": "Страховая компания Альфа",
            "insurance_policy": "POL-2026-001",
            "insurance_details": "Страхование добровольцев на период исследования.",
            "financing_source": "ООО ФармаСпонсор",
            "financing_details": "Финансирование исследования осуществляет спонсор.",
            "publication_policy": "Публикация результатов — по согласованию со спонсором и исследователем.",
            "publication_contacts": "sponsor@example.com",
        },
        "product": {"trade_name": "Test", "inn": "x", "dosage": "100 mg"},
        "reference_product": {"trade_name": "Ref", "inn": "x", "dosage": "100 mg"},
        "design": {"type": "CROSSOVER_2X2", "periods": 2},
        "food": {"condition": "FED"},
        "subjects": {"target_evaluable_n": 24, "planned_randomized_n": 28},
        "sample_size": {"evaluable_n": 24, "randomized_n": 28},
        "sampling": {"points": [{"time_h": 0, "reason": "BASELINE"}]},
        "eligibility": {"inclusion": [{"number": 1, "text": "HV"}], "non_inclusion": [], "exclusion": []},
        "analytes": [{"id": "a1", "name": "drug"}],
        "sources": [],
        "evidence_claims": [],
        "evidence_summary": {"count": 0},
    }
    ctx.update(overrides)
    return ctx


def _section_text(sec: dict) -> str:
    parts: list[str] = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            parts.append(str(b["text"]))
        if b.get("items"):
            parts.extend(str(x) for x in b["items"])
    return "\n".join(parts)


def _sec(out: dict, code: str) -> dict:
    return next(s for s in out["sections"] if s["section_code"] == code)


def test_complete_sponsor_section() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.2"))
    assert "ООО ФармаСпонсор" in text
    assert "{{SPONSOR" not in text


def test_authorized_sponsor_persons() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.3"))
    assert "Иванов И.И." in text
    assert "{{SPONSOR_PERSONS" not in text


def test_medical_expert() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.4"))
    assert "Петров П.П." in text
    assert "{{MEDICAL_EXPERT" not in text


def test_investigators_and_clinical_centers() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.5"))
    assert "Сидоров С.С." in text
    assert "Клинический центр №1" in text
    assert "{{INVESTIGATORS" not in text


def test_analytical_lab() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.6"))
    assert "BioLab Analytics" in text
    assert "{{ANALYTICAL_LAB" not in text


def test_key_organizations_cro() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.7"))
    assert "CRO Partner" in text
    assert "{{KEY_ORGS" not in text


def test_signatures_resolved() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.8"))
    assert "ООО ФармаСпонсор" in text
    assert "Сидоров С.С." in text
    assert "{{SIGNATURE" not in text
    sig = next(t for t in out["tables"] if t["table_key"] == "SIGNATURES")
    assert sig["fill_mode"] == "LABEL"
    assert any("Иванов" in str(c) for row in sig["rows"] for c in row)


def test_investigator_agreement() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "1.9"))
    assert "Сидоров С.С." in text
    assert "{{INVESTIGATOR_AGREEMENT" not in text


def test_insurance_and_financing_section_14() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "14"))
    assert "Страховая компания Альфа" in text
    assert "ООО ФармаСпонсор" in text
    assert "{{INSURANCE" not in text
    assert "{{FINANCING" not in text


def test_publication_section_15() -> None:
    out = assemble_protocol(_admin_ctx(), blocking_validation=False)
    text = _section_text(_sec(out, "15"))
    assert "Публикация результатов" in text
    assert "{{PUBLICATION" not in text


def test_missing_required_organization_unresolved() -> None:
    ctx = _admin_ctx(
        sponsor=None,
        persons=[],
        organizations=[{"name": "Only CRO", "role": "CRO"}],
        study_administration=None,
    )
    out = assemble_protocol(ctx, blocking_validation=False)
    assert "{{SPONSOR.NAME}}" in _section_text(_sec(out, "1.2"))
    assert "{{INVESTIGATORS.DETAILS}}" in _section_text(_sec(out, "1.5"))
    assert "{{INSURANCE.DETAILS}}" in _section_text(_sec(out, "14"))


def test_validation_administration_complete_passes() -> None:
    issues = validate_administration(_admin_ctx())
    assert not any(i.rule_id == "VAL.ADMIN.SPONSOR_MISSING.v1" for i in issues)


def test_validation_administration_missing_sponsor() -> None:
    issues = validate_administration(_admin_ctx(sponsor=None, organizations=[]))
    assert any(i.rule_id == "VAL.ADMIN.SPONSOR_MISSING.v1" for i in issues)


def test_p1_sections_active_in_docx_profile() -> None:
    for code in ("1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9", "14", "15"):
        assert code in ACTIVE_SECTION_CODES


def _docx_section_text(doc: Document, code: str) -> str:
    heading_map = _find_heading_indexes(doc)
    if code not in heading_map:
        return ""
    start = heading_map[code]
    paras = list(doc.paragraphs)
    sorted_heads = sorted(heading_map.items(), key=lambda kv: kv[1])
    end = len(paras)
    for i, (c, idx) in enumerate(sorted_heads):
        if c == code and i + 1 < len(sorted_heads):
            end = sorted_heads[i + 1][1]
            break
    parts: list[str] = []
    for p in paras[start + 1 : end]:
        style = p.style.name if p.style else ""
        if style.startswith("Heading") or style in {"1 абзац"}:
            continue
        if p.text:
            parts.append(p.text)
    return "\n".join(parts)


def _table_blob(doc: Document, key: str) -> str:
    idx = TABLE_KEY_TO_INDEX.get(key)
    if idx is None or idx >= len(doc.tables):
        return ""
    return "\n".join(c.text for r in doc.tables[idx].rows for c in r.cells if c.text)


def test_multiple_organizations() -> None:
    ctx = _admin_ctx(
        organizations=[
            {"id": "c1", "name": "Центр A", "role": "CLINICAL_CENTER", "country": "RU"},
            {"id": "c2", "name": "Центр B", "role": "CLINICAL_SITE", "country": "RU"},
            {"id": "l1", "name": "Lab One", "role": "ANALYTICAL_LAB", "country": "RU"},
            {"id": "l2", "name": "Lab Two", "role": "BIOANALYTICAL_LAB", "country": "RU"},
            {"id": "k1", "name": "CRO Alpha", "role": "CRO", "country": "RU"},
            {"id": "k2", "name": "Monitor Co", "role": "MONITORING", "country": "RU"},
        ]
    )
    out = assemble_protocol(ctx, blocking_validation=False)
    centers = _section_text(_sec(out, "1.5"))
    assert "Центр A" in centers and "Центр B" in centers
    labs = _section_text(_sec(out, "1.6"))
    assert "Lab One" in labs and "Lab Two" in labs
    orgs = _section_text(_sec(out, "1.7"))
    assert "CRO Alpha" in orgs and "Monitor Co" in orgs


def test_synopsis_section1_sponsor_consistency(tmp_path: Path) -> None:
    ctx = _admin_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    sponsor_sec = _section_text(_sec(out, "1.2"))
    assert "ООО ФармаСпонсор" in sponsor_sec
    meta = next(t for t in out["tables"] if t["table_key"] == "STUDY_METADATA")
    meta_blob = " ".join(str(c) for row in meta["rows"] for c in row)
    assert "BE-ADM-001" in meta_blob
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11c-admin",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))
    sec12 = _docx_section_text(doc, "1.2")
    assert "ООО ФармаСпонсор" in sec12
    assert "{{SPONSOR" not in sec12
    scoped = sec12 + "\n" + _table_blob(doc, "STUDY_METADATA") + "\n" + _table_blob(doc, "SYNOPSIS_N")
    assert "ООО ФармаСпонсор" in scoped or "BE-ADM-001" in scoped
    assert not re.search(r"46\s+(?:субъект|рандомиз)", scoped, re.IGNORECASE)


def test_canonical_admin_consistency() -> None:
    ctx = _admin_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    sponsor = _section_text(_sec(out, "1.2"))
    financing = _section_text(_sec(out, "14"))
    assert "ООО ФармаСпонсор" in sponsor
    assert "ООО ФармаСпонсор" in financing
    cons = out.get("consistency_snapshot") or {}
    assert cons.get("randomized_n") == 28
    assert cons.get("evaluable_n") == 24


def test_no_legacy_placeholders_in_p1_sections(tmp_path: Path) -> None:
    ctx = _admin_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    p1_codes = ("1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9", "14", "15")
    draft_blob = "\n".join(_section_text(_sec(out, c)) for c in p1_codes)
    assert "{{SPONSOR" not in draft_blob
    assert "{{INVESTIGATORS" not in draft_blob
    assert "{{SIGNATURE" not in draft_blob
    assert "{{INSURANCE" not in draft_blob
    assert "{{FINANCING" not in draft_blob
    assert "{{PUBLICATION" not in draft_blob

    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11c-legacy",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))
    doc_blob = "\n".join(_docx_section_text(doc, c) for c in p1_codes)
    doc_blob += "\n" + _table_blob(doc, "SIGNATURES")
    assert "{{SPONSOR" not in doc_blob
    assert "N/A" not in doc_blob


def test_docx_no_legacy_sponsor_placeholder_when_complete(tmp_path: Path) -> None:
    ctx = _admin_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11c-docx",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))
    sec12 = _docx_section_text(doc, "1.2")
    assert "1.2" in _find_heading_indexes(doc)
    assert "ООО ФармаСпонсор" in sec12
    assert "{{SPONSOR" not in sec12
    sig_text = _table_blob(doc, "SIGNATURES")
    assert "ivanov@example.com" in sig_text
    assert "+7-495-111-11-11" in sig_text
