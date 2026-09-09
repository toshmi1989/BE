"""Phase 12B.3 — sections 9–15 controlled content generation tests."""

from __future__ import annotations

from pathlib import Path

from app.domain.canonical_subjects import get_canonical_subject_counts
from app.domain.content_core_templates import CORE_12B3_SECTION_CODES
from app.domain.content_generation_qa import run_sections_9_15_qa
from app.domain.content_resolver import filter_decision_for_render
from app.domain.protocol_assembly import GENERATORS, assemble_protocol
from app.domain.protocol_generators_12b3 import CORE_12B3_GENERATORS


def _ctx(**overrides):
    base = {
        "project_id": "p12b3",
        "study": {"protocol_number": "BE-12B3-001", "title": "BE 12B3"},
        "sponsor": {"name": "Sponsor Co", "id": "sponsor-1"},
        "product": {
            "trade_name": "TestDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
            "manufacturer": "Maker",
        },
        "reference_product": {
            "trade_name": "RefDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
        },
        "design": {
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
        },
        "food": {"condition": "FASTING"},
        "subjects": {
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
            "reserve_n": 0,
        },
        "sample_size": {
            "evaluable_n": 24,
            "randomized_n": 28,
            "method": "standard",
            "id": "ss-1",
        },
        "statistical_config": {
            "id": "stat-cfg-1",
            "alpha": 0.05,
            "power": 0.8,
            "be_lower": 0.8,
            "be_upper": 1.25,
            "analysis_method": "ANOVA on log-transformed PK parameters",
            "transformation": "LN",
            "software": "SAS",
            "source_ids": ["stat-src-1"],
        },
        "cv_selection": {
            "selected_cv": 25.0,
            "parameter": "Cmax",
            "source_study_ids": ["cv-study-1"],
            "id": "cv-1",
        },
        "study_administration": {},
        "expert_decisions": [],
        "evidence_claims": [],
        "sources": [],
    }
    base.update(overrides)
    return base


def _assemble(ctx=None, only=None):
    return assemble_protocol(
        ctx or _ctx(),
        blocking_validation=False,
        only_sections=only or CORE_12B3_SECTION_CODES,
    )


def _section(assembled, code):
    return next(s for s in assembled["sections"] if s["section_code"] == code)


def _texts(sec) -> str:
    out = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            out.append(str(b["text"]))
        out.extend(str(x) for x in (b.get("items") or []))
    return "\n".join(out)


def _blocks(sec):
    return list(sec.get("content_blocks") or [])


# ---- Section 9: statistics / sample size ----


def test_statistical_method_uses_config():
    a = _assemble()
    text = _texts(_section(a, "9.1"))
    assert "ANOVA" in text
    assert "α=" in text or "alpha" in text.lower()


def test_alpha_from_config_only_no_invent_when_empty():
    a = _assemble(_ctx(statistical_config={}))
    sec = _section(a, "9.3")
    assert "{{STATISTICS.ALPHA}}" in _texts(sec)
    assert "0,05" not in _texts(sec)
    assert "0.05" not in _texts(sec)
    assert any(b.get("resolution_status") == "UNRESOLVED" for b in _blocks(sec))


def test_alpha_formats_from_canonical_config():
    a = _assemble()
    text = _texts(_section(a, "9.3"))
    assert "α=0,05" in text
    assert _blocks(_section(a, "9.3"))[0]["display_as_final"] is True


def test_be_criteria_from_config_only_no_invent_when_empty():
    a = _assemble(_ctx(statistical_config={"id": "empty-be"}))
    sec = _section(a, "9.7.2")
    assert "{{STATISTICS.BE_LIMITS}}" in _texts(sec)
    assert "80%" not in _texts(sec)
    assert "125%" not in _texts(sec)


def test_be_criteria_formats_limits():
    a = _assemble()
    text = _texts(_section(a, "9.7.2"))
    assert "80%" in text
    assert "125%" in text


def test_power_alpha_display_formatting():
    a = _assemble()
    text = _texts(_section(a, "9.2"))
    assert "мощность 80%" in text
    assert "α=0,05" in text


def test_sample_size_uses_subject_plan_n():
    a = _assemble()
    text = _texts(_section(a, "9.2"))
    assert "SubjectPlan" in text
    assert "N=24" in text or "оцениваемых N=24" in text
    assert "N=28" in text or "рандомизированных N=28" in text
    counts = get_canonical_subject_counts(_ctx())
    assert counts.source == "subject_plan"
    assert counts.randomized_n == 28


def test_sample_size_shows_both_when_subjectplan_and_calculation_diverge():
    ctx = _ctx(
        subjects={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
            "reserve_n": 0,
        },
        sample_size={"evaluable_n": 24, "randomized_n": 36, "method": "standard", "id": "ss-div"},
    )
    counts = get_canonical_subject_counts(ctx)
    assert counts.diverges_from_sample_size is True
    a = _assemble(ctx)
    text = _texts(_section(a, "9.2"))
    assert "SubjectPlan" in text
    assert "SampleSizeCalculation" in text
    assert "28" in text and "36" in text
    assert "Семантическое различие" in text or "≠" in text


def test_calculated_n_not_used_as_randomized_n():
    ctx = _ctx(
        subjects={},
        sample_size={"evaluable_n": 24, "randomized_n": 34, "method": "standard", "id": "ss-calc"},
    )
    counts = get_canonical_subject_counts(ctx)
    assert counts.source == "sample_size_fallback"
    a = _assemble(ctx)
    ss_blocks = [b for b in _blocks(_section(a, "9.2")) if b.get("block_code") == "STAT.SAMPLE_SIZE"]
    assert ss_blocks
    b = ss_blocks[0]
    text = str(b.get("text") or "")
    assert "{{SUBJECTS.RANDOMIZED_N}}" in text
    assert "SampleSizeCalculation" in text
    assert "34" in text
    assert b.get("subject_plan_randomized_n") is None
    assert b.get("sample_size_randomized_n") == 34
    assert b.get("display_as_final") is False
    # Protocol randomized must not silently equal calculated-only N
    assert "рандомизированных N=34" not in text.replace("SampleSizeCalculation", "")


def test_missing_randomized_n_creates_gap():
    ctx = _ctx(
        subjects={"target_evaluable_n": 24, "planned_screened_n": 40, "reserve_n": 0},
        sample_size={"evaluable_n": 24, "method": "standard"},
    )
    assert get_canonical_subject_counts(ctx).randomized_n is None
    a = _assemble(ctx)
    sec = _section(a, "9.2")
    text = _texts(sec)
    assert "{{SUBJECTS.RANDOMIZED_N}}" in text
    gaps = []
    for b in _blocks(sec):
        gaps.extend(b.get("knowledge_gaps") or [])
    assert any(
        g.get("related_rule_id") == "STAT.N_RANDOMIZED_MISSING" or "randomized" in str(g.get("question") or "").lower()
        for g in gaps
    )
    findings = run_sections_9_15_qa(sections=a["sections"], ctx=ctx, mode="DRAFT")
    assert any(f.code in {"STAT.N_RANDOMIZED_MISSING", "QA.STAT.RANDOMIZED_N_UNRESOLVED"} for f in findings)


def test_unverified_potvin_not_rendered_final():
    ctx = _ctx(
        sample_size={
            "evaluable_n": 24,
            "randomized_n": 28,
            "method": "POTVIN Method B two-stage",
            "id": "ss-potvin",
        },
        expert_decisions=[],
    )
    a = _assemble(ctx)
    sec = _section(a, "9.2")
    ss_blocks = [b for b in _blocks(sec) if b.get("block_code") == "STAT.SAMPLE_SIZE"]
    assert ss_blocks
    assert ss_blocks[0]["display_as_final"] is False
    assert "{{STATISTICS.POTVIN}}" in str(ss_blocks[0].get("text") or "") or any(
        "{{STATISTICS.POTVIN}}" in (b.get("unresolved") or []) for b in ss_blocks
    ) or any(
        (g.get("related_rule_id") == "STAT.POTVIN_UNVERIFIED")
        for b in ss_blocks
        for g in (b.get("knowledge_gaps") or [])
    )


def test_cv_without_provenance_fails_validation():
    ctx = _ctx(cv_selection={"selected_cv": 30.0, "parameter": "AUC", "source_study_ids": []})
    a = _assemble(ctx)
    findings = run_sections_9_15_qa(sections=a["sections"], ctx=ctx, mode="DRAFT")
    codes = {f.code for f in findings}
    assert "STAT.CV_MISSING_SOURCE" in codes or "QA.STAT.MISSING_PROVENANCE" in codes
    ss = next(b for b in _blocks(_section(a, "9.2")) if b.get("block_code") == "STAT.SAMPLE_SIZE")
    assert ss.get("display_as_final") is False


def test_proposed_statistic_not_final():
    ctx = _ctx(
        statistical_config={},
        stats_policies={"stopping_rules": "Stop if futility criterion met (draft wording)"},
    )
    a = _assemble(ctx)
    sec = _section(a, "9.4")
    blocks = _blocks(sec)
    assert blocks
    for b in blocks:
        if b.get("resolution_status") == "PROPOSED":
            assert b.get("display_as_final") is False
    assert any(b.get("resolution_status") == "PROPOSED" for b in blocks) or any(
        b.get("display_as_final") is False for b in blocks
    )


def test_analysis_populations_from_subject_plan():
    a = _assemble()
    text = _texts(_section(a, "9.7"))
    assert "24" in text
    assert "28" in text
    assert "SubjectPlan" in text or "рандомиз" in text.lower()


def test_stopping_rules_missing_creates_gap():
    a = _assemble(_ctx(statistical_config={"id": "x", "alpha": 0.05}))
    sec = _section(a, "9.4")
    assert "{{STATISTICS.STOPPING_RULES}}" in _texts(sec)
    assert any(b.get("knowledge_gaps") for b in _blocks(sec))


def test_missing_data_policy_missing_creates_gap():
    a = _assemble()
    assert "{{STATISTICS.MISSING_DATA}}" in _texts(_section(a, "9.5"))


def test_outliers_policy_missing_creates_gap():
    a = _assemble()
    assert "{{STATISTICS.OUTLIERS}}" in _texts(_section(a, "9.7.3"))


def test_section9_method_be_limits_not_invented():
    a = _assemble(_ctx(statistical_config={"analysis_method": "ANOVA", "id": "no-be"}))
    text = _texts(_section(a, "9.1"))
    assert "{{STATISTICS.BE_LIMITS}}" in text or any(
        "BE limits" in str(g.get("question") or "") for b in _blocks(_section(a, "9.1")) for g in b.get("knowledge_gaps") or []
    )


# ---- Sections 10–15 ----


def test_sections_10_15_generators_registered():
    for key in (
        "data_access",
        "standard_text",
        "financing_insurance",
        "publications",
        "sample_size",
        "alpha",
        "be_criteria",
    ):
        assert key in CORE_12B3_GENERATORS
        assert GENERATORS[key] is CORE_12B3_GENERATORS[key]


def test_section10_data_access_static_verified():
    a = _assemble()
    blocks = _blocks(_section(a, "10"))
    assert any(b.get("origin") == "STATIC_VERIFIED" or b.get("content_type") == "STATIC_VERIFIED" for b in blocks)
    blob = _texts(_section(a, "10")).lower()
    assert "мониторинг" in blob or "доступ" in blob or "отклонен" in blob


def test_section11_quality_overlay():
    a = _assemble()
    blocks = _blocks(_section(a, "11"))
    assert blocks
    assert any(b.get("content_type") == "STATIC_VERIFIED" or b.get("origin") == "STATIC_VERIFIED" for b in blocks)


def test_missing_ethics_committee_not_invented():
    a = _assemble(_ctx(study_administration={}, organizations=[], persons=[]))
    text = _texts(_section(a, "12"))
    assert "{{ETHICS.COMMITTEE}}" in text
    fake_names = ("Acme IRB", "Dummy Ethics", "Test Committee LLC", "Вымышленный комитет")
    for name in fake_names:
        assert name not in text


def test_ethics_committee_from_admin_when_present():
    a = _assemble(
        _ctx(study_administration={"ethics_committee": "IEC Central Hospital", "id": "admin-1"})
    )
    text = _texts(_section(a, "12"))
    assert "IEC Central Hospital" in text
    assert "{{ETHICS.COMMITTEE}}" not in text


def test_missing_insurance_not_invented():
    a = _assemble(_ctx(study_administration={}, organizations=[], persons=[]))
    text = _texts(_section(a, "14"))
    assert "{{INSURANCE.DETAILS}}" in text
    assert "Страховая компания Рога и Копыта" not in text
    assert "Allianz" not in text
    assert "policy #" not in text.lower()


def test_missing_retention_period_not_invented():
    a = _assemble(_ctx(study_administration={}))
    text = _texts(_section(a, "13"))
    assert "{{DATA.RETENTION_YEARS}}" in text
    assert "15 лет" not in text
    assert "25 years" not in text.lower()


def test_retention_renders_when_canonical():
    a = _assemble(_ctx(study_administration={"retention_years": 15, "id": "adm-ret"}))
    text = _texts(_section(a, "13"))
    assert "15" in text
    assert "{{DATA.RETENTION_YEARS}}" not in text


def test_missing_financial_data_not_invented():
    a = _assemble(_ctx(study_administration={}, organizations=[], persons=[]))
    text = _texts(_section(a, "14"))
    assert "{{FINANCING.DETAILS}}" in text
    assert "Grant #12345" not in text
    assert "бюджет" not in text.lower() or "{{" in text


def test_missing_publication_clause_not_invented():
    a = _assemble(_ctx(study_administration={}))
    text = _texts(_section(a, "15"))
    assert "{{PUBLICATION.POLICY}}" in text
    assert "sponsor shall own all publications" not in text.lower()
    assert any(b.get("resolution_status") == "UNRESOLVED" for b in _blocks(_section(a, "15")))


def test_publication_clause_from_admin():
    a = _assemble(
        _ctx(
            study_administration={
                "publication_policy": "Публикации согласуются со спонсором до подачи.",
                "id": "adm-pub",
            }
        )
    )
    assert "согласуются со спонсором" in _texts(_section(a, "15"))


def test_static_verified_block_preserved():
    a = _assemble()
    for code in ("11", "12", "10", "13"):
        blocks = _blocks(_section(a, code))
        static = [
            b
            for b in blocks
            if b.get("origin") == "STATIC_VERIFIED" or b.get("content_type") == "STATIC_VERIFIED"
        ]
        if static:
            assert static[0]["display_as_final"] is True
            assert static[0]["resolution_status"] == "RESOLVED"
            return
    raise AssertionError("expected STATIC_VERIFIED block in section 10–13")


def test_financing_shows_canonical_sponsor():
    a = _assemble()
    text = _texts(_section(a, "14"))
    assert "Sponsor Co" in text


# ---- Expert decisions / QA ----


def test_rejected_decision_not_rendered():
    assert filter_decision_for_render({"status": "REJECTED", "decision_type": "POTVIN_METHOD"}) is None
    findings = run_sections_9_15_qa(
        sections=[
            {
                "section_code": "9.2",
                "content_blocks": [
                    {
                        "text": "x",
                        "decision_status": "REJECTED",
                        "used_in_content": True,
                        "display_as_final": True,
                    }
                ],
            }
        ],
        ctx=_ctx(),
    )
    assert any(f.code in {"CONTENT.REJECTED_DECISION_USED", "QA.CONTENT.REJECTED_DECISION"} for f in findings)


def test_superseded_decision_not_rendered():
    assert filter_decision_for_render({"status": "SUPERSEDED"}) is None
    findings = run_sections_9_15_qa(
        sections=[
            {
                "section_code": "9.1",
                "content_blocks": [
                    {"text": "x", "decision_status": "SUPERSEDED", "used_in_content": True}
                ],
            }
        ],
        ctx=_ctx(),
    )
    assert any(f.code in {"CONTENT.SUPERSEDED_DECISION_USED", "QA.CONTENT.SUPERSEDED_DECISION"} for f in findings)


def test_legacy_admin_value_detected():
    findings = run_sections_9_15_qa(
        sections=[],
        ctx=_ctx(),
        legacy_hints={"admin_legacy": "old template ethics committee from prior protocol"},
    )
    assert any(f.code == "QA.CONTENT.LEGACY_VALUE" for f in findings)


def test_proposed_as_final_blocked_by_qa():
    findings = run_sections_9_15_qa(
        sections=[
            {
                "section_code": "9.4",
                "content_blocks": [
                    {
                        "text": "proposed policy",
                        "resolution_status": "PROPOSED",
                        "display_as_final": True,
                    }
                ],
            }
        ],
        ctx=_ctx(),
        mode="DRAFT",
    )
    assert any(
        f.code in {"QA.CONTENT.PROPOSED_AS_FINAL", "STAT.UNAPPROVED_PARAMETERS", "QA.STAT.UNAPPROVED_VALUE"}
        for f in findings
    )


def test_approved_potvin_can_render():
    ctx = _ctx(
        sample_size={
            "evaluable_n": 24,
            "randomized_n": 28,
            "method": "POTVIN Method B",
            "id": "ss-p",
        },
        expert_decisions=[
            {
                "status": "APPROVED",
                "decision_type": "POTVIN_METHOD",
                "payload": {"method": "POTVIN Method B", "text": "Approved two-stage Potvin B"},
                "decided_at": "2026-01-01",
            }
        ],
    )
    a = _assemble(ctx)
    ss = next(b for b in _blocks(_section(a, "9.2")) if b.get("block_code") == "STAT.SAMPLE_SIZE")
    assert "POTVIN" in str(ss.get("text") or "").upper() or "Утверждённый метод" in str(ss.get("text") or "")
    assert "{{STATISTICS.POTVIN}}" not in (ss.get("unresolved") or [])


def test_cv_with_provenance_ok():
    ctx = _ctx()
    a = _assemble(ctx)
    findings = run_sections_9_15_qa(sections=a["sections"], ctx=ctx, mode="DRAFT")
    assert not any(f.code == "STAT.CV_MISSING_SOURCE" for f in findings)
    text = _texts(_section(a, "9.2"))
    assert "cv-study-1" in text or "CVintra=25" in text


def test_core_12b3_section_codes_cover_tree():
    expected = {
        "9",
        "9.1",
        "9.2",
        "9.3",
        "9.4",
        "9.5",
        "9.6",
        "9.7",
        "9.7.1",
        "9.7.1.1",
        "9.7.1.2",
        "9.7.2",
        "9.7.3",
        "9.7.4",
        "10",
        "11",
        "12",
        "13",
        "14",
        "15",
    }
    assert expected <= set(CORE_12B3_SECTION_CODES)
    # Actual SECTION_TREE: 10 is Direct access only (no invented 10.2 code)
    assert "10.2" not in CORE_12B3_SECTION_CODES


def test_assemble_emits_all_12b3_codes():
    a = _assemble()
    codes = {s["section_code"] for s in a["sections"]}
    for c in ("9.1", "9.2", "9.3", "10", "11", "12", "13", "14", "15"):
        assert c in codes


def test_no_n_a_fallback():
    a = _assemble(_ctx(study_administration={}, statistical_config={}))
    blob = "\n".join(_texts(s) for s in a["sections"]).lower()
    assert "n/a" not in blob
    assert "not applicable" not in blob


def test_generators_registered_overlay():
    for k, fn in CORE_12B3_GENERATORS.items():
        assert GENERATORS[k] is fn
    assert "heading" not in CORE_12B3_GENERATORS


# ---- DOCX ----


def test_targeted_docx_sections_9_12_14(tmp_path):
    from app.domain.docx_renderer import render_protocol_docx

    ctx = _ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="12b3",
        protocol_version="1",
        only_sections=["9.2", "12", "14"],
    )
    assert result.status == "READY", result.blocking_reasons
    assert result.filename


def test_unselected_sections_preserved(tmp_path):
    import shutil

    from docx import Document

    from app.domain.docx_profile import get_template_profile
    from app.domain.docx_renderer import _find_heading_indexes, _section_body_range, render_protocol_docx

    ctx = _ctx()
    assembled = assemble_protocol(ctx, blocking_validation=False)
    profile = get_template_profile()
    baseline = tmp_path / "base.docx"
    shutil.copy2(profile.template_path(), baseline)
    base_doc = Document(str(baseline))
    base_map = _find_heading_indexes(base_doc)
    static_code = next((c for c in ("5", "5.1", "6", "7", "8") if c in base_map), None)
    if static_code is None:
        static_code = next(c for c in base_map if not str(c).startswith(("9", "1")))
    br = _section_body_range(base_map, static_code, len(base_doc.paragraphs))
    before = [base_doc.paragraphs[i].text for i in range(br[0], br[1])]
    result = render_protocol_docx(
        protocol_payload=assembled,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="12b3-static",
        protocol_version="1",
        only_sections=["9.2"],
    )
    assert result.status == "READY", result.blocking_reasons
    out = Document(str(result.output_path))
    out_map = _find_heading_indexes(out)
    orng = _section_body_range(out_map, static_code, len(out.paragraphs))
    after = [out.paragraphs[i].text for i in range(orng[0], orng[1])]
    assert before == after


def test_no_global_replace_in_12b3_source():
    src = Path(__file__).resolve().parents[1] / "app" / "domain" / "protocol_generators_12b3.py"
    text = src.read_text(encoding="utf-8")
    # Formatting helpers may use str.replace for decimal commas; forbid DOCX/content overlays
    assert "find_and_replace" not in text.lower()
    assert "global_replace" not in text.lower()
    assert "document.Replace" not in text
    assert "paragraphs[" not in text
    docx_src = Path(__file__).resolve().parents[1] / "app" / "domain" / "docx_renderer.py"
    docx_text = docx_src.read_text(encoding="utf-8")
    assert "only_sections" in docx_text


def test_no_positional_mapping_in_12b3_source():
    src = Path(__file__).resolve().parents[1] / "app" / "domain" / "protocol_generators_12b3.py"
    text = src.read_text(encoding="utf-8")
    assert "cells[0]" not in text
    assert "cell[1]" not in text
    assert "paragraphs[" not in text


def test_full_pipeline_golden_sections_9_15(client):
    pid = client.post("/api/projects", json={"name": "12B3-golden"}).json()["id"]
    client.put(
        f"/api/projects/{pid}/study",
        json={"protocol_number": "BE-12B3-G", "title": "Golden 12B3", "version": "1.0"},
    )
    client.put(
        f"/api/projects/{pid}/product",
        json={"trade_name": "TestDrug", "inn": "x", "dosage": "100 mg", "dosage_form": "tablet"},
    )
    client.put(
        f"/api/projects/{pid}/reference-product",
        json={"trade_name": "RefDrug", "inn": "x", "dosage": "100 mg", "dosage_form": "tablet"},
    )
    client.post(
        f"/api/projects/{pid}/design",
        json={"type": "CROSSOVER_2X2", "periods": 2, "sequences": [["T", "R"], ["R", "T"]]},
    )
    client.post(f"/api/projects/{pid}/food", json={"condition": "FASTING"})
    client.post(
        f"/api/projects/{pid}/subjects",
        json={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "reserve_n": 0,
            "planned_screened_n": 40,
        },
    )
    gen = client.post(
        f"/api/projects/{pid}/content/generate-sections-9-15",
        json={"persist": True, "mode": "DRAFT"},
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()
    assert body["mutates_study"] is False
    assert body["sections"]
    codes = {s.get("section_code") for s in body["sections"]}
    assert "9.2" in codes or any(str(c).startswith("9") for c in codes)
    docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "DRAFT", "only_sections": ["9.2", "12", "14"]},
    )
    assert docx.status_code == 200, docx.text
    assert docx.json()["status"] == "READY", docx.json()
