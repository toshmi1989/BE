"""Phase 11C-H — Synopsis N/sponsor hardening + REVIEW/FINAL legacy subject-N gate."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.domain.canonical_subjects import get_canonical_subject_counts
from app.domain.docx_profile import TABLE_KEY_TO_INDEX
from app.domain.docx_renderer import _find_heading_indexes, _legacy_subject_n_gate, render_protocol_docx
from app.domain.protocol_assembly import assemble_protocol
from app.domain.static_blocks import legacy_uncontrolled_blocks
from app.domain.study_snapshot import build_canonical_snapshot
from app.domain.synopsis_fill import (
    SUBJECT_COUNT_46_RE,
    classify_legacy_46_hit,
    find_legacy_subject_count_46,
    rewrite_subject_count_text,
)


def _complete_ctx(**overrides):
    ctx = {
        "project_id": "p11ch",
        "study": {
            "protocol_number": "BE-HARD-001",
            "title": "Hardening BE",
            "country": "RU",
            "version": "1.0",
            "primary_objective": "Оценить биоэквивалентность",
        },
        "sponsor": {
            "id": "sp1",
            "name": "ООО ФармаСпонсор",
            "address": "Москва",
            "country": "RU",
            "contact_email": "sponsor@example.com",
        },
        "organizations": [
            {
                "id": "o1",
                "name": "Клинический центр №1",
                "role": "CLINICAL_CENTER",
                "address": "Санкт-Петербург",
                "country": "RU",
            },
            {
                "id": "o2",
                "name": "BioLab Analytics",
                "role": "BIOANALYTICAL_LAB",
                "address": "Москва",
            },
            {"id": "o3", "name": "CRO Partner", "role": "CRO", "country": "RU"},
        ],
        "persons": [
            {
                "id": "p1",
                "role": "SPONSOR_AUTHORIZED",
                "full_name": "Иванов И.И.",
                "title": "Директор",
                "phone": "+7-495-111-11-11",
                "email": "ivanov@example.com",
            },
            {
                "id": "p2",
                "role": "MEDICAL_EXPERT",
                "full_name": "Петров П.П.",
                "title": "Медицинский эксперт",
            },
            {
                "id": "p3",
                "role": "PRINCIPAL_INVESTIGATOR",
                "full_name": "Сидоров С.С.",
                "title": "Главный исследователь",
                "organization_id": "o1",
                "organization_name": "Клинический центр №1",
            },
        ],
        "study_administration": {
            "insurance_provider": "СК Альфа",
            "insurance_policy": "POL-1",
            "financing_source": "ООО ФармаСпонсор",
            "publication_policy": "Публикация по согласованию со спонсором.",
        },
        "product": {
            "trade_name": "TestDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
            "manufacturer": "ООО Тест",
            "manufacturer_country": "RU",
            "composition": "testdrug 100 mg",
            "route": "oral",
            "registration_number": "LP-001",
            "storage_conditions": "15-25C",
            "shelf_life": "24 months",
            "batch": "T-001",
        },
        "reference_product": {
            "trade_name": "RefDrug",
            "inn": "testdrug",
            "dosage": "100 mg",
            "dosage_form": "tablet",
            "manufacturer": "Ref Pharma",
            "manufacturer_country": "DE",
            "composition": "testdrug 100 mg",
            "route": "oral",
            "registration_number": "EU-001",
            "storage_conditions": "15-25C",
            "shelf_life": "36 months",
            "batch": "R-001",
            "purchased_status": "UNKNOWN",
            "source_ids": ["s1"],
        },
        "design": {"type": "CROSSOVER_2X2", "periods": 2, "sequences": [["T", "R"], ["R", "T"]]},
        "food": {"condition": "FED", "meal_type": "HIGH_CALORIE"},
        "subjects": {
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
        },
        "sample_size": {"evaluable_n": 24, "randomized_n": 28, "selected_cv": 25},
        "sampling": {
            "points": [
                {"time_h": 0.0, "reason": "BASELINE"},
                {"time_h": 2.0, "reason": "TMAX_CAPTURE"},
                {"time_h": 24.0, "reason": "FINAL"},
            ],
            "total_points_per_period": 3,
        },
        "observation": {"selected_duration": 24, "unit": "h"},
        "washout": {"selected_value": 7, "unit": "day"},
        "analytes": [{"id": "a1", "name": "testdrug", "type": "PARENT"}],
        "pk_parameters": [{"analyte_id": "a1", "parameter_code": "Cmax"}],
        "eligibility": {
            "inclusion": [{"number": 1, "text": "HV"}],
            "non_inclusion": [],
            "exclusion": [{"number": 1, "text": "AE"}],
        },
        "sources": [{"id": "s1", "type": "SmPC", "title": "Label", "year": 2024}],
        "evidence_claims": [
            {
                "field_name": "clinical_summary",
                "value": "HV PK available.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
            {
                "field_name": "risk_benefit",
                "value": "Acceptable for BE.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
            {
                "field_name": "pharmacology",
                "value": "Oral absorption.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
            {
                "field_name": "literature",
                "value": "SmPC and published PK.",
                "status": "VERIFIED",
                "source_ids": ["s1"],
            },
        ],
        "statistical_config": {"alpha": 0.05, "be_lower": 0.8, "be_upper": 1.25},
        "cv_selection": {"selected_cv": 25, "selection_method": "SINGLE_STUDY"},
        "evidence_summary": {"count": 4},
    }
    ctx.update(overrides)
    return ctx


def _section_text(sec: dict) -> str:
    parts: list[str] = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            parts.append(str(b["text"]))
        if b.get("items"):
            parts.extend(str(x) for x in b["items"])
    return "\n".join(parts)


def test_synopsis_n_exact_canonical() -> None:
    ctx = _complete_ctx()
    subjects = get_canonical_subject_counts(ctx)
    assert subjects.randomized_n == 28
    assert subjects.evaluable_n == 24
    snap = build_canonical_snapshot(ctx)
    assert snap.randomized_n == 28
    out = assemble_protocol(ctx, blocking_validation=False)
    syn = next(s for s in out["sections"] if s["section_code"] == "SYNOPSIS")
    text = _section_text(syn)
    assert "рандомизированных: 28" in text
    assert "Оцениваемых субъектов: 24" in text
    assert "46" not in text
    sec26 = next(s for s in out["sections"] if s["section_code"] == "2.6")
    t26 = _section_text(sec26)
    assert "24" in t26 and "28" in t26
    sec92 = next(s for s in out["sections"] if s["section_code"] == "9.2")
    assert "рандомизированных N=28" in _section_text(sec92)
    syn_n = next(t for t in out["tables"] if t["table_key"] == "SYNOPSIS_N")
    blob = " ".join(str(c) for row in syn_n["rows"] for c in row)
    assert "28" in blob and "24" in blob


def test_synopsis_sponsor_exact_canonical() -> None:
    ctx = _complete_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    syn = _section_text(next(s for s in out["sections"] if s["section_code"] == "SYNOPSIS"))
    sec12 = _section_text(next(s for s in out["sections"] if s["section_code"] == "1.2"))
    assert "ООО ФармаСпонсор" in syn
    assert "ООО ФармаСпонсор" in sec12
    syn_n = next(t for t in out["tables"] if t["table_key"] == "SYNOPSIS_N")
    admin_blob = " ".join(str(c) for row in syn_n["rows"] for c in row)
    assert "ООО ФармаСпонсор" in admin_blob
    assert "Сидоров" in admin_blob or "Клинический центр" in admin_blob


def test_toc_page_number_46_does_not_trigger_subject_n_error() -> None:
    hit = classify_legacy_46_hit(
        location="para[60]",
        text="5. Отбор и исключение субъектов\t46",
        style="toc 1",
    )
    assert hit["kind"] == "TOC_PAGE_NUMBER"
    assert hit["must_replace"] is False
    hits = find_legacy_subject_count_46(
        paragraphs=[
            ("para[60]", "5. Отбор и исключение субъектов\t46", "toc 1"),
            ("para[61]", "5.1 Критерии включения субъектов\t46", "Normal"),
        ],
        table_cells=[("T16R1C2", "46")],
        canonical_randomized_n=28,
    )
    # TOC tab-page and bare CV n ignored
    assert hits == []


def test_legacy_synopsis_n_cell_cannot_survive_final(tmp_path: Path) -> None:
    # Unit: rewrite removes subject-count 46
    old = "Рандомизированные добровольцы: 46\nВсего у 46 добровольцев."
    new = rewrite_subject_count_text(old, randomized_n=28, screened_n=40)
    assert "46" not in new
    assert "28" in new
    assert SUBJECT_COUNT_46_RE.search(old)
    assert not SUBJECT_COUNT_46_RE.search(new)

    # Gate: if residual subject-count 46 present → FINAL blocked
    residual = find_legacy_subject_count_46(
        paragraphs=[],
        table_cells=[
            (
                "T2R12C1",
                "Рандомизированные добровольцы: 46",
            )
        ],
        canonical_randomized_n=28,
    )
    assert residual

    ctx = _complete_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    # After normal render, FINAL must not see residual subject-count 46
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="FINAL",
        output_dir=tmp_path,
        project_slug="p11ch-final-n",
        protocol_version="1",
    )
    # May still block on other unresolved (CRF etc.) — but never due to missing N rewrite
    if result.status == "BLOCKED":
        assert not any("legacy subject-count 46" in r for r in (result.blocking_reasons or []))
    else:
        doc = Document(str(result.output_path))
        assert _legacy_subject_n_gate(doc, out.get("consistency_snapshot") or {"randomized_n": 28}) == []


def test_complete_p1_admin_review_ready(tmp_path: Path) -> None:
    ctx = _complete_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    # No P1 unresolved markers
    p1_codes = ("1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9", "14", "15")
    blob = "\n".join(
        _section_text(next(s for s in out["sections"] if s["section_code"] == c)) for c in p1_codes
    )
    assert "{{SPONSOR" not in blob
    assert "{{INVESTIGATORS" not in blob
    assert "{{SIGNATURE" not in blob
    assert "{{INSURANCE" not in blob

    result = render_protocol_docx(
        protocol_payload={**out, "status": "READY_FOR_REVIEW"},
        study_ctx=ctx,
        mode="REVIEW",
        output_dir=tmp_path,
        project_slug="p11ch-review",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))
    # Synopsis sponsor + N
    t03 = doc.tables[TABLE_KEY_TO_INDEX["SYNOPSIS_N"]]
    syn_blob = "\n".join(c.text for r in t03.rows for c in r.cells)
    assert "ООО ФармаСпонсор" in syn_blob
    assert "Рандомизированные добровольцы: 28" in syn_blob or "у 28 добровольцев" in syn_blob
    assert not re.search(r"Рандомизированные добровольцы:\s*46", syn_blob)
    assert not re.search(r"у\s+46\s+добровольц", syn_blob)
    # TOC page 46 may remain — must not trip gate
    assert _legacy_subject_n_gate(doc, {"randomized_n": 28}) == []
    # §1.2 present
    assert "1.2" in _find_heading_indexes(doc)


def test_no_p1_legacy_uncontrolled() -> None:
    ids = {b.block_id for b in legacy_uncontrolled_blocks()}
    assert "LEGACY.SYNOPSIS_N46" not in ids
    assert not any("P1" in i or "SPONSOR" in i for i in ids)


def test_docx_complete_draft_admin_sections(tmp_path: Path) -> None:
    ctx = _complete_ctx()
    out = assemble_protocol(ctx, blocking_validation=False)
    result = render_protocol_docx(
        protocol_payload=out,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="p11ch-draft",
        protocol_version="1",
    )
    assert result.status == "READY", result.blocking_reasons
    doc = Document(str(result.output_path))
    heads = _find_heading_indexes(doc)
    for code in ("1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9"):
        assert code in heads
    t03 = doc.tables[TABLE_KEY_TO_INDEX["SYNOPSIS_N"]]
    syn_blob = "\n".join(c.text for r in t03.rows for c in r.cells)
    assert "ООО ФармаСпонсор" in syn_blob
    assert "BE-HARD-001" in syn_blob
    assert "Сидоров" in syn_blob or "Петров" in syn_blob or "Иванов" in syn_blob
    assert "BioLab" in syn_blob or "Клинический центр" in syn_blob
    sig = doc.tables[TABLE_KEY_TO_INDEX["SIGNATURES"]]
    sig_blob = "\n".join(c.text for r in sig.rows for c in r.cells)
    assert "ivanov@example.com" in sig_blob
