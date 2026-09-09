"""Phase 6 research engine integration — no external network."""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient


def test_research_profile_tasks_documents_evidence_flow(client: TestClient, tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "docs"))
    from app.core.config import get_settings

    get_settings.cache_clear()

    project = client.post("/api/projects", json={"name": "Phase6 research"}).json()
    pid = project["id"]

    client.post(
        f"/api/projects/{pid}/client-input",
        json={
            "requested_product_name": "Bosutinib",
            "inn": "bosutinib",
            "dosage": "400 mg",
            "dosage_form": "tablet",
            "route": "oral",
            "requested_subject_count": 46,
            "provenance": {"origin": "USER", "status": "PROPOSED"},
        },
    )

    profile = client.post(
        f"/api/projects/{pid}/research-profile",
        json={"regulatory_jurisdiction": "EAEU", "sync_from_client_input": True},
    )
    assert profile.status_code == 200, profile.text
    assert profile.json()["inn"] == "bosutinib"
    assert profile.json()["search_profile"]["pk_search_terms"]

    case = client.post(f"/api/projects/{pid}/research-case", json={"status": "NEW"})
    assert case.status_code == 201

    tasks = client.post(f"/api/projects/{pid}/research-case/tasks/generate")
    assert tasks.status_code == 200, tasks.text
    body = tasks.json()
    assert any(t["task_type"] == "REFERENCE_PRODUCT" for t in body)
    pk = next(t for t in body if t["task_type"] == "PK")
    assert "PRODUCT_LABEL" in pk["depends_on_tasks"]

    # DOCX upload + ingest
    buf = BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Tmax was 2–3 h under fed conditions.")
    doc.add_paragraph("CVintra for Cmax was approximately 25%.")
    doc.save(buf)
    files = {
        "file": (
            "label.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    up = client.post(
        f"/api/projects/{pid}/documents",
        files=files,
        data={"source_type": "SmPC"},
    )
    assert up.status_code == 201, up.text
    doc_id = up.json()["id"]
    assert up.json()["checksum"]
    assert up.json()["source_id"]

    ingested = client.post(f"/api/projects/{pid}/documents/{doc_id}/ingest")
    assert ingested.status_code == 200, ingested.text
    assert ingested.json()["status"] == "INGESTED"

    pages = client.get(f"/api/projects/{pid}/documents/{doc_id}/pages")
    assert pages.status_code == 200
    assert pages.json()
    assert "Tmax" in pages.json()[0]["text"]

    search = client.post(
        f"/api/projects/{pid}/research/search",
        json={"query": "Tmax", "phrase": False},
    )
    assert search.status_code == 200, search.text
    assert search.json()
    assert search.json()[0]["document_id"] == doc_id
    assert search.json()[0]["snippet"]

    source_id = up.json()["source_id"]
    ev1 = client.post(
        f"/api/projects/{pid}/research/evidence",
        json={
            "source_id": source_id,
            "document_id": doc_id,
            "page": 1,
            "field_code": "tmax",
            "extracted_text": "Tmax was 2–3 h",
            "auto_detect_conflicts": True,
        },
    )
    assert ev1.status_code == 201, ev1.text
    claim = ev1.json()["claims"][0]
    assert claim["normalized_value"]["min"] == 2
    assert claim["normalized_value"]["max"] == 3
    assert claim["origin"] == "USER"
    assert claim["status"] == "PROPOSED"

    # conflicting tmax
    client.post(
        f"/api/projects/{pid}/research/evidence",
        json={
            "source_id": source_id,
            "field_code": "tmax",
            "extracted_text": "Tmax 4 h",
            "auto_detect_conflicts": True,
        },
    )
    conflicts = client.post(f"/api/projects/{pid}/research/conflicts")
    assert conflicts.status_code == 200
    assert any(c["field_name"] == "tmax" for c in conflicts.json())

    completeness = client.post(f"/api/projects/{pid}/research/completeness")
    assert completeness.status_code == 200
    assert "score" in completeness.json()
    assert completeness.json()["notes"]

    fields = client.get("/api/reference-data/evidence-fields")
    assert fields.status_code == 200
    assert any(f["code"] == "tmax" for f in fields.json())

    # TXT upload
    txt = client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("notes.txt", b"Food effect: fed state recommended.\n", "text/plain")},
        data={"source_type": "OTHER"},
    )
    assert txt.status_code == 201
    client.post(f"/api/projects/{pid}/documents/{txt.json()['id']}/ingest")

    get_profile = client.get(f"/api/projects/{pid}/research-profile")
    assert get_profile.status_code == 200
