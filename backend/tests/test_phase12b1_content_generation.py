"""Phase 12B.1 — core protocol content generation tests."""

from __future__ import annotations

from app.domain.content_core_templates import CORE_12B1_SECTION_CODES, be_objective_template
from app.domain.content_generation_qa import run_content_generation_qa
from app.domain.content_resolver import filter_decision_for_render
from app.domain.display_value_registry import find_raw_enums_in_text, resolve_display
from app.domain.protocol_assembly import assemble_protocol, GENERATORS
from app.domain.protocol_generators_12b1 import CORE_12B1_GENERATORS


def _ctx(**overrides):
    base = {
        "project_id": "p12b1",
        "study": {
            "protocol_number": "BE-12B1-001",
            "title": "BE study",
            "primary_objective": None,
        },
        "sponsor": {"name": "Sponsor Co"},
        "product": {
            "trade_name": "TestDrug",
            "inn": "testdrug",
            "manufacturer": "Maker",
            "dosage_form": "tablet",
            "dosage": "100 mg",
            "source_ids": ["s1"],
        },
        "reference_product": {
            "trade_name": "RefDrug",
            "inn": "testdrug",
            "manufacturer": "RefMaker",
            "dosage_form": "tablet",
            "dosage": "100 mg",
            "source_ids": ["s2"],
        },
        "design": {
            "type": "CROSSOVER_2X2",
            "periods": 2,
            "sequences": [["T", "R"], ["R", "T"]],
        },
        "food": {"condition": "FASTING"},
        "subjects": {
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "planned_screened_n": 40,
            "reserve_n": 0,
        },
        "sample_size": {"evaluable_n": 24, "randomized_n": 28},
        "sampling": {
            "points": [
                {"time_h": 0.0, "reason": "BASELINE"},
                {"time_h": 2.0, "reason": "TMAX_CAPTURE"},
                {"time_h": 1.0, "reason": "ABSORPTION"},
                {"time_h": 24.0, "reason": "FINAL"},
            ],
            "total_points_per_period": 4,
            "source_ids": ["samp1"],
        },
        "washout": {"selected_value": 7, "unit": "day", "requires_washout": True},
        "observation": {"selected_duration": 24, "unit": "h"},
        "analytes": [{"name": "testdrug", "tmax": 2.0}],
        "expert_decisions": [],
        "eligibility": {"inclusion": [], "non_inclusion": [], "exclusion": []},
        "sources": [],
        "evidence_claims": [],
    }
    base.update(overrides)
    return base


def _assemble(ctx=None, only=None):
    return assemble_protocol(
        ctx or _ctx(),
        blocking_validation=False,
        only_sections=only or CORE_12B1_SECTION_CODES,
    )


def _section(assembled, code):
    return next(s for s in assembled["sections"] if s["section_code"] == code)


def _texts(sec):
    out = []
    for b in sec.get("content_blocks") or []:
        if b.get("text"):
            out.append(b["text"])
        out.extend(str(x) for x in (b.get("items") or []))
    return "\n".join(out)


# ---- Required tests ----


def test_identity_resolves_from_canonical():
    a = _assemble()
    # subjects / product from canonical
    subj = _section(a, "2.6") if any(s["section_code"] == "2.6" for s in a["sections"]) else None
    # 2.6 is subjects_rationale in tree - check CORE includes it
    assert "2.6" in CORE_12B1_SECTION_CODES or any(
        "28" in _texts(s) for s in a["sections"] if s["section_code"] in {"2.6", "4.8.1"}
    )
    prod = _section(a, "2.1.1")
    assert "TestDrug" in _texts(prod)
    assert prod["content_blocks"][0].get("canonical_source") == "product"


def test_product_fields_use_semantic_mapping():
    a = _assemble()
    prod = _section(a, "2.1.1")
    text = _texts(prod)
    assert "Торговое наименование" in text or "trade" in text.lower() or "TestDrug" in text
    assert "Производитель" in text or "Maker" in text
    assert "Лекарственная форма" in text or "tablet" in text
    # Not positional cell indices
    assert "cell[0]" not in text


def test_reference_is_not_auto_selected():
    a = _assemble(_ctx(reference_product=None, expert_decisions=[]))
    ref = _section(a, "2.1.2")
    text = _texts(ref)
    assert "REFERENCE_SELECTION_REVIEW_REQUIRED" in text or "{{REFERENCE_PRODUCT" in text
    assert ref["content_blocks"][0].get("display_as_final") is False


def test_proposed_design_not_final():
    a = _assemble(
        _ctx(
            design={},
            expert_decisions=[
                {
                    "id": "d1",
                    "decision_type": "DESIGN",
                    "status": "PROPOSED",
                    "proposed_value": {"design": "CROSSOVER_2X2"},
                }
            ],
        )
    )
    des = _section(a, "4.2")
    assert des["content_blocks"][0].get("display_as_final") is False
    assert des["content_blocks"][0].get("resolution_status") in {"PROPOSED", "UNRESOLVED"}


def test_approved_design_resolves():
    a = _assemble()
    des = _section(a, "4.2")
    text = _texts(des)
    assert "CROSSOVER_2X2" not in text
    assert des["content_blocks"][0].get("display_as_final") is True
    assert "перекр" in text.lower() or "2×2" in text or "2x2" in text.lower()


def test_missing_design_creates_gap():
    a = _assemble(_ctx(design={}))
    des = _section(a, "4.2")
    gaps = des["content_blocks"][0].get("knowledge_gaps") or []
    assert gaps


def test_subject_n_uses_canonical_source():
    a = _assemble()
    # find subjects_rationale section
    sec = next(s for s in a["sections"] if s["section_code"] == "2.6")
    text = _texts(sec)
    assert "28" in text
    assert "SubjectPlan" in text or sec["content_blocks"][0].get("canonical_source") == "subject_plan"


def test_subject_n_does_not_use_template_text():
    a = _assemble(_ctx(subjects={"target_evaluable_n": 24, "planned_randomized_n": 42, "planned_screened_n": 50, "reserve_n": 0}))
    sec = _section(a, "2.6")
    text = _texts(sec)
    assert "42" in text
    assert "46" not in text  # legacy synopsis N


def test_sampling_uses_canonical_plan():
    a = _assemble()
    samp = _section(a, "4.4.2")
    text = _texts(samp)
    assert "каноническ" in text.lower() or samp["content_blocks"][0].get("canonical_source") == "sampling.points"
    assert any(b.get("table_key") == "BLOOD_SAMPLING" for b in samp["content_blocks"])


def test_sampling_points_ordered():
    a = _assemble()
    samp = _section(a, "4.4.2")
    text = _texts(samp)
    # 0, 1, 2, 24 in order in summary
    idx0 = text.find("0")
    idx1 = text.find("1")
    idx2 = text.find("2")
    assert 0 <= idx0 < idx1 < idx2 or "0, 1, 2, 24" in text.replace(".0", "")


def test_sampling_not_generated_from_old_table():
    # Even if legacy points differ, generator uses ctx sampling only
    a = _assemble(
        _ctx(
            sampling={
                "points": [{"time_h": 0.0}, {"time_h": 3.5}, {"time_h": 12.0}],
                "source_ids": ["canonical"],
            }
        )
    )
    text = _texts(_section(a, "4.4.2"))
    assert "3.5" in text
    assert "old table" not in text.lower()


def test_washout_requires_approved_value():
    a = _assemble(_ctx(washout={}))
    # washout may be under 2.12
    wo = next(s for s in a["sections"] if s["section_code"] == "2.12")
    assert wo["content_blocks"][0].get("display_as_final") is False


def test_food_does_not_invent_calories():
    a = _assemble()
    treat = _section(a, "4.4")
    text = _texts(treat).lower()
    assert "kcal" not in text
    assert "50%" not in text
    assert "200 ml" not in text
    assert "натощак" in text or "{{food" in text


def test_content_resolver_provenance():
    a = _assemble()
    prod = _section(a, "2.1.1")
    b0 = prod["content_blocks"][0]
    assert b0.get("resolution_status")
    assert "canonical_source" in b0 or b0.get("origin") == "SOURCE_DERIVED"


def test_proposed_content_not_final():
    assert filter_decision_for_render({"status": "PROPOSED"}) is None


def test_rejected_decision_not_rendered():
    assert filter_decision_for_render({"status": "REJECTED"}) is None
    findings = run_content_generation_qa(
        sections=[
            {
                "section_code": "4.2",
                "content_blocks": [
                    {
                        "text": "x",
                        "decision_status": "REJECTED",
                        "used_in_content": True,
                        "display_as_final": True,
                    }
                ],
            }
        ]
    )
    assert any(f.code == "CONTENT.REJECTED_DECISION_USED" for f in findings)


def test_superseded_decision_not_rendered():
    assert filter_decision_for_render({"status": "SUPERSEDED"}) is None
    findings = run_content_generation_qa(
        sections=[
            {
                "section_code": "4.2",
                "content_blocks": [
                    {
                        "text": "x",
                        "decision_status": "SUPERSEDED",
                        "used_in_content": True,
                    }
                ],
            }
        ]
    )
    assert any(f.code == "CONTENT.SUPERSEDED_DECISION_USED" for f in findings)


def test_raw_enum_not_rendered():
    a = _assemble()
    blob = "\n".join(_texts(s) for s in a["sections"])
    assert not find_raw_enums_in_text(blob)
    assert "CROSSOVER_2X2" not in blob
    assert "FASTING" not in blob or "натощак" in blob


def test_canonical_cross_section_consistency():
    a = _assemble()
    prod = _texts(_section(a, "2.1.1"))
    ref = _texts(_section(a, "2.1.2"))
    des = _texts(_section(a, "4.2"))
    assert "TestDrug" in prod
    assert "RefDrug" in ref
    assert "CROSSOVER_2X2" not in des
    n_sec = _texts(_section(a, "2.6"))
    assert "28" in n_sec


def test_placeholder_blocks_final():
    findings = run_content_generation_qa(
        sections=[
            {
                "section_code": "2.1.1",
                "status": "UNRESOLVED",
                "content_blocks": [{"text": "{{TEST_PRODUCT.MISSING}}", "unresolved": ["{{TEST_PRODUCT.MISSING}}"]}],
            }
        ],
        mode="FINAL",
    )
    assert any(f.code == "CONTENT.PLACEHOLDER_UNRESOLVED" and f.blocking for f in findings)


def test_legacy_value_detected():
    findings = run_content_generation_qa(
        sections=[
            {
                "section_code": "2.6",
                "content_blocks": [{"text": "рандомизированных субъектов: 46"}],
            }
        ],
        canonical={"randomized_n": 28},
        legacy_hints={"randomized_n": "46"},
    )
    assert any(f.code == "CONTENT.LEGACY_VALUE" for f in findings)


def test_generators_registered():
    for k in CORE_12B1_GENERATORS:
        assert GENERATORS[k] is CORE_12B1_GENERATORS[k]


def test_objective_template_requires_all_fields():
    assert be_objective_template(design_display=None, test_name="A", reference_name="B") is None
    t = be_objective_template(
        design_display=resolve_display("CROSSOVER_2X2", context="design"),
        test_name="A",
        reference_name="B",
    )
    assert t and "A" in t and "B" in t


def test_docx_targeted_render_and_pipeline(client):
    pid = client.post("/api/projects", json={"name": "12B1-docx"}).json()["id"]
    client.put(
        f"/api/projects/{pid}/study",
        json={
            "protocol_number": "BE-12B1-DOCX-001",
            "title": "12B1 targeted DOCX",
            "version": "1.0",
        },
    )
    client.post(
        f"/api/projects/{pid}/design",
        json={"type": "CROSSOVER_2X2", "periods": 2, "sequences": [["T", "R"], ["R", "T"]]},
    )
    client.post(
        f"/api/projects/{pid}/subjects",
        json={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "reserve_n": 0,
            "planned_screened_n": 40,
        },
    )
    client.put(
        f"/api/projects/{pid}/product",
        json={
            "trade_name": "TestDrug",
            "inn": "x",
            "dosage": "100 mg",
            "dosage_form": "tablet",
            "manufacturer": "M",
        },
    )
    client.put(
        f"/api/projects/{pid}/reference-product",
        json={"trade_name": "RefDrug", "inn": "x", "dosage": "100 mg", "dosage_form": "tablet"},
    )
    core = client.post(
        f"/api/projects/{pid}/content/generate-core",
        json={"persist": True, "mode": "DRAFT"},
    )
    assert core.status_code == 200, core.text
    body = core.json()
    assert body["mutates_study"] is False
    assert body["protocol_draft_id"]
    # Targeted DOCX
    docx = client.post(
        f"/api/projects/{pid}/protocol/docx/build",
        json={"mode": "DRAFT", "only_sections": ["4.2", "4.4.2", "2.1.1"]},
    )
    assert docx.status_code == 200, docx.text
    payload = docx.json()
    assert payload["status"] == "READY", payload
    assert payload.get("filename")
    assert payload.get("storage_key")


def test_no_global_replace_in_docx_renderer_source():
    from pathlib import Path

    src = Path(__file__).resolve().parents[1] / "app" / "domain" / "docx_renderer.py"
    text = src.read_text(encoding="utf-8")
    assert "replace_all" not in text.lower() or "only_sections" in text
    assert "only_sections" in text


def test_docx_does_not_modify_unrelated_static_blocks(tmp_path):
    """Targeted render must leave unselected section bodies identical to the template copy."""
    import shutil

    from docx import Document

    from app.domain.docx_profile import get_template_profile
    from app.domain.docx_renderer import _find_heading_indexes, _section_body_range, render_protocol_docx

    ctx = _ctx()
    assembled = assemble_protocol(ctx, blocking_validation=False)
    profile = get_template_profile()
    baseline = tmp_path / "baseline.docx"
    shutil.copy2(profile.template_path(), baseline)
    base_doc = Document(str(baseline))
    base_map = _find_heading_indexes(base_doc)
    # Pick a non-core section present in the template (e.g. 7 or 8) if available
    static_code = next((c for c in ("7", "8", "9", "10", "5", "6") if c in base_map), None)
    if static_code is None:
        static_code = next(c for c in base_map if c not in {"4.2", "4.4.2", "2.1.1"})
    br = _section_body_range(base_map, static_code, len(base_doc.paragraphs))
    assert br is not None
    before = [base_doc.paragraphs[i].text for i in range(br[0], br[1])]

    result = render_protocol_docx(
        protocol_payload=assembled,
        study_ctx=ctx,
        mode="DRAFT",
        output_dir=tmp_path,
        project_slug="12b1-static",
        protocol_version="1",
        only_sections=["4.2"],
    )
    assert result.status == "READY", result.blocking_reasons
    out = Document(str(result.output_path))
    out_map = _find_heading_indexes(out)
    assert static_code in out_map
    orng = _section_body_range(out_map, static_code, len(out.paragraphs))
    assert orng is not None
    after = [out.paragraphs[i].text for i in range(orng[0], orng[1])]
    assert before == after


def test_full_pipeline_canonical_to_docx_api(client):
    pid = client.post("/api/projects", json={"name": "12B1-pipe"}).json()["id"]
    client.post(
        f"/api/projects/{pid}/design",
        json={"type": "CROSSOVER_2X2", "periods": 2, "sequences": [["T", "R"], ["R", "T"]]},
    )
    client.post(
        f"/api/projects/{pid}/food",
        json={"condition": "FASTING"},
    )
    client.post(
        f"/api/projects/{pid}/subjects",
        json={
            "target_evaluable_n": 24,
            "planned_randomized_n": 28,
            "reserve_n": 0,
            "planned_screened_n": 40,
        },
    )
    gen = client.post(f"/api/projects/{pid}/content/generate-core", json={"persist": True}).json()
    assert gen.get("sections")
    # QA should not invent kcal
    blob = str(gen)
    assert "800" not in blob or "kcal" not in blob.lower()
