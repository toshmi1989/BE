"""Phase 30.3 — DOCX semantic integrity hardening (fail-closed FINAL)."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document

from app.domain.cover_mapping import build_cover_rows, detect_wrong_cover_mapping, resolve_cover_value
from app.domain.docx_profile import DOCX_GENERATOR_VERSION, TABLE_KEY_TO_INDEX, get_template_profile
from app.domain.docx_renderer import _fill_label_value_table, render_protocol_docx
from app.domain.docx_semantic_integrity import (
    assess_docx_semantic_integrity,
    semantic_preflight_from_context,
)
from app.domain.docx_toc import clear_stale_toc_page_numbers, refresh_toc_fields
from app.domain.protocol_constants import PROTOCOL_GENERATOR_VERSION
from app.domain.semantic_field_types import FieldTypeError, coerce_typed
from app.domain.template_block_registry import TEMPLATE_BLOCKS, required_final_blocks


def test_version_0353():
    assert PROTOCOL_GENERATOR_VERSION == "0.35.3"
    assert DOCX_GENERATOR_VERSION == "0.35.3"


def test_cover_not_study_metadata_ordinal():
    assert TABLE_KEY_TO_INDEX.get("COVER_METADATA") == 0
    assert "STUDY_METADATA" not in TABLE_KEY_TO_INDEX


def test_wrong_subject_dosage_mapping_blocked():
    """P0.1 — dosage_form must not accept subject count integer."""
    with pytest.raises(FieldTypeError):
        coerce_typed("dosage_form", 56)
    assert detect_wrong_cover_mapping("Лекарственная форма", "56") == (
        "WRONG_MAPPING_DOSAGE_FORM_SUBJECT_COUNT"
    )
    ctx = {
        "study": {"protocol_number": "UPDCB-02-BE-2026", "title": "T", "version": "1.0", "version_date": "01.01.2026"},
        "product": {"trade_name": "Upadacitinib", "inn": "upadacitinib", "dosage_form": "таблетки", "dosage": "15 mg"},
        "sponsor": {"name": "Sponsor"},
        "subjects": {"planned_randomized_n": 56, "target_evaluable_n": 6},
    }
    _field, value, err = resolve_cover_value("Лекарственная форма", ctx)
    assert err is None
    assert value is not None
    assert "56" not in value
    assert "15" in value or "таблетки" in value


def test_no_ordinal_fallback_fills_wrong_cover_cell(tmp_path: Path):
    """STUDY_METADATA-style rows must not ordinal-dump into cover labels."""
    tpl = get_template_profile().template_path()
    out = tmp_path / "cover.docx"
    shutil.copy2(tpl, out)
    doc = Document(str(out))
    # Simulate old bug payload: Evaluable/Randomized as rows 4-5
    bad_rows = [
        ["Protocol number", "UPDCB"],
        ["Title", "Title"],
        ["Design", "XO"],
        ["Food", "FED"],
        ["Evaluable N", "6"],
        ["Randomized N", "56"],
    ]
    audit = _fill_label_value_table(doc.tables[0], bad_rows, allow_ordinal_fallback=False)
    assert audit["ordinal_fallback"] is False
    # Cover dosage_form cell must not become 56
    form_cell = None
    for row in doc.tables[0].rows:
        if "лекарственная форма" in (row.cells[0].text or "").lower():
            form_cell = (row.cells[1].text or "").strip()
            break
    assert form_cell is not None
    assert form_cell != "56"
    assert "400" in form_cell or "таблет" in form_cell.lower()  # still template until typed fill


def test_typed_dose_rejects_bare_integer():
    with pytest.raises(FieldTypeError):
        coerce_typed("dose", 15)
    assert coerce_typed("dose", "15 mg") == "15 mg"


def test_stale_400_mg_blocks_final(tmp_path: Path):
    tpl = get_template_profile().template_path()
    out = tmp_path / "stale400.docx"
    shutil.copy2(tpl, out)
    ctx = {
        "product": {"trade_name": "Upadacitinib", "inn": "upadacitinib", "dosage": "15 mg"},
        "study": {"protocol_number": "UPDCB-02-BE-2026", "version_date": "01.02.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(out, ctx, mode="FINAL", toc_refreshed=True)
    assert any(i.code == "STALE_TEMPLATE_DOSE_400" for i in report.issues)
    assert report.ok is False


def test_stale_washout_inconsistency_blocks_final(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Период отмывки длительностью 14 дней.")
    doc.add_paragraph("Washout 7 day between periods.")
    path = tmp_path / "wo.docx"
    doc.save(str(path))
    ctx = {
        "product": {"trade_name": "X", "inn": "x", "dosage": "15 mg"},
        "study": {"version_date": "01.01.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(path, ctx, mode="FINAL", toc_refreshed=True)
    assert any(i.code == "WASHOUT_INCONSISTENT" for i in report.issues)


def test_stale_date_blocks_final(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Версия 1.0 от 18.04.2025 г.")
    path = tmp_path / "date.docx"
    doc.save(str(path))
    ctx = {
        "product": {"trade_name": "X", "inn": "x", "dosage": "15 mg"},
        "study": {"version_date": "01.09.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(path, ctx, mode="FINAL", toc_refreshed=True)
    assert any(i.code == "STALE_TEMPLATE_DATE" for i in report.issues)


def test_unresolved_placeholder_blocks_final(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Method: {{BIOANALYSIS.METHOD}}")
    path = tmp_path / "ph.docx"
    doc.save(str(path))
    ctx = {
        "product": {"trade_name": "X", "inn": "x", "dosage": "15 mg"},
        "study": {"version_date": "01.01.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(path, ctx, mode="FINAL", toc_refreshed=True)
    assert any(i.code == "UNRESOLVED_TEMPLATE_PLACEHOLDER" for i in report.issues)
    assert report.placeholder_count >= 1


def test_missing_sampling_preflight_blocks_final():
    ctx = {
        "study": {"version_date": "01.01.2026", "protocol_number": "U"},
        "product": {"dosage_form": "tab", "dosage": "15 mg"},
        "subjects": {"target_evaluable_n": 6, "planned_randomized_n": 56},
        "washout": {"selected_value": 7, "unit": "day"},
        "observation": {"selected_duration": 72},
        "eligibility": {"inclusion": ["age 18-45"]},
        "bioanalysis_plan": {"method": "LC-MS"},
        "statistical_config": {"status": "APPROVED"},
        "sampling": {},
    }
    pf = semantic_preflight_from_context(ctx, mode="FINAL")
    assert pf["ok"] is False
    assert any(b["code"] == "MISSING_SAMPLING_PLAN" for b in pf["blockers"])


def test_missing_eligibility_preflight_blocks_final():
    ctx = {
        "study": {"version_date": "01.01.2026"},
        "product": {"dosage_form": "tab", "dosage": "15 mg"},
        "subjects": {"target_evaluable_n": 6, "planned_randomized_n": 56},
        "washout": {"selected_value": 7},
        "observation": {"selected_duration": 72},
        "eligibility": {"inclusion": []},
        "bioanalysis_plan": {"method": "LC-MS"},
        "statistical_config": {"status": "APPROVED"},
        "sampling": {"points": [{"time_h": 0}]},
    }
    pf = semantic_preflight_from_context(ctx, mode="FINAL")
    assert any(b["code"] == "MISSING_ELIGIBILITY" for b in pf["blockers"])


def test_missing_bioanalysis_preflight_blocks_final():
    ctx = {
        "study": {"version_date": "01.01.2026"},
        "product": {"dosage_form": "tab", "dosage": "15 mg"},
        "subjects": {"target_evaluable_n": 6, "planned_randomized_n": 56},
        "washout": {"selected_value": 7},
        "observation": {"selected_duration": 72},
        "eligibility": {"inclusion": ["ok"]},
        "bioanalysis_plan": {},
        "statistical_config": {"status": "APPROVED"},
        "sampling": {"points": [{"time_h": 0}]},
    }
    pf = semantic_preflight_from_context(ctx, mode="FINAL")
    assert any(b["code"] == "MISSING_BIOANALYSIS" for b in pf["blockers"])


def test_missing_statistics_preflight_blocks_final():
    ctx = {
        "study": {"version_date": "01.01.2026"},
        "product": {"dosage_form": "tab", "dosage": "15 mg"},
        "subjects": {"target_evaluable_n": 6, "planned_randomized_n": 56},
        "washout": {"selected_value": 7},
        "observation": {"selected_duration": 72},
        "eligibility": {"inclusion": ["ok"]},
        "bioanalysis_plan": {"method": "LC-MS"},
        "statistical_config": {},
        "sampling": {"points": [{"time_h": 0}]},
    }
    pf = semantic_preflight_from_context(ctx, mode="FINAL")
    assert any(b["code"] == "MISSING_STATISTICS_PLAN" for b in pf["blockers"])


def test_stale_toc_cleared(tmp_path: Path):
    doc = Document()
    p = doc.add_paragraph("1. Introduction\t12")
    p.style = doc.styles["Normal"]
    # emulate TOC style if available
    try:
        p.style = doc.styles["TOC 1"]
    except KeyError:
        p.add_run("")  # keep text
    path = tmp_path / "toc.docx"
    doc.save(str(path))
    status = refresh_toc_fields(path)
    assert status.get("toc_refreshed") is True


def test_cross_section_product_consistency_helper():
    rows = build_cover_rows(
        {
            "study": {"title": "T", "protocol_number": "U", "version": "1", "version_date": "d"},
            "product": {"trade_name": "Upa", "inn": "upa", "dosage_form": "tab", "dosage": "15 mg"},
            "sponsor": {"name": "S"},
        }
    )
    form_row = next(r for r in rows if "форма" in r[0].lower())
    assert "56" not in form_row[1]
    assert "15 mg" in form_row[1]


def test_stale_composition_markers_in_template(tmp_path: Path):
    """Template composition residues are product-specific — FINAL integrity flags 400 mg."""
    tpl = get_template_profile().template_path()
    out = tmp_path / "comp.docx"
    shutil.copy2(tpl, out)
    ctx = {
        "product": {
            "trade_name": "Upadacitinib",
            "inn": "upadacitinib",
            "dosage": "15 mg",
            "composition": None,
        },
        "study": {"version_date": "01.09.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(out, ctx, mode="FINAL", toc_refreshed=True)
    assert report.stale_content_count >= 1 or any(
        i.code in {"STALE_TEMPLATE_DOSE_400", "STALE_TEMPLATE_DATE"} for i in report.issues
    )


def test_internal_pk_table_id_blocks_final(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("См. таблица PK_PARAMETERS")
    path = tmp_path / "pkid.docx"
    doc.save(str(path))
    ctx = {
        "product": {"trade_name": "X", "inn": "x", "dosage": "15 mg"},
        "study": {"version_date": "01.01.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(path, ctx, mode="FINAL", toc_refreshed=True)
    assert any(i.code == "INTERNAL_TABLE_ID_LEAK" for i in report.issues)


def test_template_block_registry_covers_required_sections():
    assert len(TEMPLATE_BLOCKS) >= 30
    ids = {b.block_id for b in TEMPLATE_BLOCKS}
    for need in (
        "cover.metadata",
        "sampling.table",
        "eligibility.inclusion",
        "bioanalysis.method",
        "statistics",
        "washout",
        "toc",
    ):
        assert need in ids
    assert required_final_blocks()


def test_missing_sampling_placeholder_in_docx(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("{{SAMPLING.POINTS}} {{SAMPLING.TIME_H}}")
    path = tmp_path / "samp.docx"
    doc.save(str(path))
    ctx = {
        "product": {"trade_name": "X", "inn": "x", "dosage": "15 mg"},
        "study": {"version_date": "01.01.2026"},
        "washout": {"selected_value": 7, "unit": "day"},
    }
    report = assess_docx_semantic_integrity(path, ctx, mode="FINAL", toc_refreshed=True)
    assert any(
        i.code in {"MISSING_SAMPLING_PLAN", "UNRESOLVED_TEMPLATE_PLACEHOLDER"} for i in report.issues
    )
