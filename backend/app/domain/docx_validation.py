"""DOCX validation — presentation QA only."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZipFile, BadZipFile

from docx import Document

UNRESOLVED_RE = re.compile(r"\{\{[A-Z0-9_.]+\}\}")


@dataclass
class DocxIssue:
    severity: str  # CRITICAL | ERROR | WARNING | INFO
    code: str
    message: str
    details: dict[str, Any] = field(default_factory=dict)


@dataclass
class DocxValidationReport:
    ok: bool
    blocking: bool
    issues: list[DocxIssue] = field(default_factory=list)
    paragraph_count: int = 0
    table_count: int = 0
    unresolved_found: list[str] = field(default_factory=list)
    heading_count: int = 0

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "blocking": self.blocking,
            "paragraph_count": self.paragraph_count,
            "table_count": self.table_count,
            "heading_count": self.heading_count,
            "unresolved_found": self.unresolved_found,
            "issues": [
                {
                    "severity": i.severity,
                    "code": i.code,
                    "message": i.message,
                    "details": i.details,
                }
                for i in self.issues
            ],
        }


def validate_docx_file(
    path: Path,
    *,
    expected_min_tables: int = 30,
    expected_min_headings: int = 20,
    allow_unresolved: bool = False,
) -> DocxValidationReport:
    issues: list[DocxIssue] = []
    unresolved: list[str] = []

    if not path.exists():
        return DocxValidationReport(
            ok=False,
            blocking=True,
            issues=[DocxIssue("CRITICAL", "FILE_MISSING", f"Missing file: {path}")],
        )

    try:
        with ZipFile(path) as zf:
            names = zf.namelist()
            if "word/document.xml" not in names:
                issues.append(
                    DocxIssue("CRITICAL", "CORRUPT_XML", "word/document.xml missing")
                )
    except BadZipFile:
        return DocxValidationReport(
            ok=False,
            blocking=True,
            issues=[DocxIssue("CRITICAL", "CORRUPT_ZIP", "Not a valid DOCX/ZIP")],
        )

    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        return DocxValidationReport(
            ok=False,
            blocking=True,
            issues=[DocxIssue("CRITICAL", "OPEN_FAILED", f"Cannot open DOCX: {exc}")],
        )

    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    heading_count = 0
    blob_parts: list[str] = []

    for p in doc.paragraphs:
        text = p.text or ""
        blob_parts.append(text)
        style = p.style.name if p.style else ""
        if style.startswith("Heading"):
            heading_count += 1
        unresolved.extend(UNRESOLVED_RE.findall(text))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                blob_parts.append(cell.text or "")
                unresolved.extend(UNRESOLVED_RE.findall(cell.text or ""))

    unresolved = sorted(set(unresolved))
    blob = "\n".join(blob_parts)

    if table_count < expected_min_tables:
        issues.append(
            DocxIssue(
                "ERROR",
                "TABLE_COUNT",
                f"Expected at least {expected_min_tables} tables, found {table_count}",
                {"table_count": table_count},
            )
        )
    if heading_count < expected_min_headings:
        issues.append(
            DocxIssue(
                "ERROR",
                "HEADING_COUNT",
                f"Expected at least {expected_min_headings} headings, found {heading_count}",
                {"heading_count": heading_count},
            )
        )
    if unresolved and not allow_unresolved:
        issues.append(
            DocxIssue(
                "CRITICAL",
                "UNRESOLVED_PLACEHOLDER",
                f"Unresolved placeholders in DOCX: {unresolved[:20]}",
                {"count": len(unresolved)},
            )
        )
    elif unresolved and allow_unresolved:
        issues.append(
            DocxIssue(
                "WARNING",
                "UNRESOLVED_PLACEHOLDER_DRAFT",
                f"Draft mode allows unresolved: {unresolved[:20]}",
                {"count": len(unresolved)},
            )
        )

    # Required section presence (heading text contains codes)
    for code in ("1.1", "2.1.1", "4.4.2", "5.1", "9.2"):
        if code not in blob:
            issues.append(
                DocxIssue(
                    "CRITICAL",
                    "MISSING_SECTION",
                    f"Required section marker missing: {code}",
                )
            )

    # Broken internal refs
    from app.domain.reference_registry import detect_broken_reference_text
    from app.domain.display_value_registry import find_raw_enums_in_text

    broken = detect_broken_reference_text(blob)
    if broken:
        issues.append(
            DocxIssue(
                "CRITICAL",
                "BROKEN_REFERENCE",
                f"Broken references in DOCX: {broken}",
                {"hits": broken},
            )
        )

    raw_enums = find_raw_enums_in_text(blob)
    if raw_enums:
        severity = "WARNING" if allow_unresolved else "ERROR"
        issues.append(
            DocxIssue(
                severity,
                "RAW_ENUM_IN_DOCUMENT",
                f"Raw enum values in DOCX: {raw_enums}",
                {"enums": raw_enums},
            )
        )

    if para_count == 0:
        issues.append(DocxIssue("CRITICAL", "EMPTY_DOCUMENT", "No paragraphs"))

    blocking = any(i.severity in {"CRITICAL", "ERROR"} for i in issues)
    ok = not blocking
    return DocxValidationReport(
        ok=ok,
        blocking=blocking,
        issues=issues,
        paragraph_count=para_count,
        table_count=table_count,
        unresolved_found=unresolved,
        heading_count=heading_count,
    )
