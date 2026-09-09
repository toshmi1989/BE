"""Phase 10 helpers — golden artifact IO (no new engines)."""

from __future__ import annotations

import json
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[2]
GOLDEN_ROOT = REPO_ROOT / "golden"


def new_run_dir(prefix: str = "phase10") -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    path = GOLDEN_ROOT / f"{prefix}-{stamp}"
    path.mkdir(parents=True, exist_ok=False)
    return path


def write_json(path: Path, data: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def copy_file(src: Path, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return dest


TEMPLATE_EXCERPTS = [
    "Расчет количества субъектов основан на данных о внутрииндивидуальной вариабельности фармакокинетического параметра Cmax бозутиниба, как наиболее вариабельного фармакокинетического параметра.",
    "Для бозутиниба ожидаемая внутрииндивидуальная вариабельность фармакокинетического параметра Cmax составляет >30% [1].",
    "Согласно данным оригинальных препаратов tmax бозутиниба в среднем составляет 6 часов [2].",
    "Согласно данным оригинальных препаратов конечный период t1/2 бозутиниба составляет 35,5 часов [2], поэтому забор образцов крови в течение 72 часов является достаточным.",
    "Сравниваемые препараты: Бозулиф (бозутиниб, 400 мг, таблетки, покрытые пленочной оболочкой, Pfiser Inc Inc, США) и Бозутиниб (400 мг, таблетки, покрытые пленочной оболочкой).",
    "Для исследования биоэквивалентности у добровольцев предполагается использовать сравниваемые препараты в каждом периоде однократно бозутиниб в дозе равной 400 мг после еды.",
]


def build_fixture_docx_bytes() -> bytes:
    """Golden research fixture: excerpts copied from BE_Protocol_Template_v2.0 body (not invented)."""
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_heading("Golden fixture excerpts from BE_Protocol_Template_v2.0", level=1)
    doc.add_paragraph(
        "Source: templates/protocol/BE_Protocol_Template_v2.0.docx body text. "
        "Used only for offline E2E evidence extraction tests."
    )
    for text in TEMPLATE_EXCERPTS:
        doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
