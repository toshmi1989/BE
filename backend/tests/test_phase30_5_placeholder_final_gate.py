"""Phase 30.5 — Deterministic DRAFT/FINAL placeholder gate (no new medical engines)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

from app.core.config import Settings, get_settings
from app.domain.docx_profile import DOCX_GENERATOR_VERSION, get_template_profile
from app.domain.docx_semantic_integrity import assess_docx_semantic_integrity
from app.domain.placeholder_registry import (
    OPTIONAL_DYNAMIC,
    REQUIRED_DYNAMIC,
    TOC_VISUAL_VALIDATION,
    UNKNOWN,
    assess_placeholders_for_mode,
    classify_placeholder,
    registry_as_list,
)
from app.domain.protocol_constants import PROTOCOL_GENERATOR_VERSION


def test_version_0355():
    assert Settings().app_version == "0.35.5"
    assert PROTOCOL_GENERATOR_VERSION == "0.35.5"
    assert DOCX_GENERATOR_VERSION == "0.35.5"
    assert get_settings().app_version == "0.35.5"


def test_placeholder_registry_complete_and_classified():
    rows = registry_as_list()
    assert len(rows) >= 30
    classes = {r["classification"] for r in rows}
    assert REQUIRED_DYNAMIC in classes
    assert OPTIONAL_DYNAMIC in classes
    # Every catalog entry must declare FINAL policy
    for r in rows:
        assert r["final_policy"] in {"BLOCK", "ALLOW"}
        assert r["classification"] in {
            REQUIRED_DYNAMIC,
            OPTIONAL_DYNAMIC,
            "STATIC_TEMPLATE",
        }


def test_required_placeholder_blocks_final():
    gate = assess_placeholders_for_mode(["{{SAMPLING.POINTS}}"], mode="FINAL")
    assert gate["ok"] is False
    assert gate["code"] == "UNRESOLVED_TEMPLATE_PLACEHOLDER"
    assert gate["required_count"] == 1
    assert gate["blocking"][0]["code"] == "SAMPLING.POINTS"
    assert gate["blocking"][0]["section"] == "4.4.2"


def test_optional_allowlisted_placeholder_permitted_in_final():
    gate = assess_placeholders_for_mode(["{{PUBLICATION.POLICY}}"], mode="FINAL")
    assert gate["ok"] is True
    assert gate["code"] == "PLACEHOLDERS_OK"
    assert gate["optional_allowed"]
    assert classify_placeholder("{{PUBLICATION.POLICY}}").classification == OPTIONAL_DYNAMIC


def test_unknown_placeholder_blocks_final():
    gate = assess_placeholders_for_mode(["{{TOTALLY.UNKNOWN.FIELD}}"], mode="FINAL")
    assert gate["ok"] is False
    assert gate["unknown_count"] == 1
    assert gate["blocking"][0]["classification"] == UNKNOWN
    assert classify_placeholder("{{TOTALLY.UNKNOWN.FIELD}}").classification == UNKNOWN


def test_draft_allows_explicit_placeholders():
    markers = ["{{SAMPLING.POINTS}}", "{{TOTALLY.UNKNOWN.FIELD}}", "{{PUBLICATION.POLICY}}"]
    gate = assess_placeholders_for_mode(markers, mode="DRAFT")
    assert gate["ok"] is True
    assert gate["code"] == "DRAFT_PLACEHOLDERS_ALLOWED"
    assert len(gate["blocking_for_final"]) >= 2  # required + unknown


def test_docx_final_gate_scans_body_and_blocks(tmp_path: Path):
    tpl = get_template_profile().template_path()
    out = tmp_path / "ph.docx"
    shutil.copy2(tpl, out)
    doc = Document(str(out))
    doc.add_paragraph("Sampling: {{SAMPLING.POINTS}}")
    doc.save(str(out))

    ctx = {
        "study": {"version_date": "01.01.2026"},
        "product": {"dosage_form": "tablet", "dosage": "15 mg", "inn": "upadacitinib"},
        "subjects": {"planned_randomized_n": 56},
    }
    draft = assess_docx_semantic_integrity(out, ctx, mode="DRAFT", toc_refreshed=True)
    assert draft.ok or any(
        i.code == "UNRESOLVED_TEMPLATE_PLACEHOLDER" and i.severity == "WARNING"
        for i in draft.issues
    )
    final = assess_docx_semantic_integrity(out, ctx, mode="FINAL", toc_refreshed=True)
    assert final.ok is False
    assert any(i.code == "UNRESOLVED_TEMPLATE_PLACEHOLDER" for i in final.issues)


def test_docx_optional_allowlist_does_not_alone_block_final(tmp_path: Path):
    out = tmp_path / "opt.docx"
    doc = Document()
    doc.add_paragraph("Policy: {{PUBLICATION.POLICY}}")
    doc.add_paragraph("Sources: {{SOURCES.LIST}}")
    doc.save(str(out))
    ctx = {
        "study": {"version_date": "01.01.2026"},
        "product": {"dosage_form": "tablet", "dosage": "15 mg"},
        "subjects": {"planned_randomized_n": 56},
    }
    report = assess_docx_semantic_integrity(out, ctx, mode="FINAL", toc_refreshed=True)
    ph_issues = [i for i in report.issues if i.code == "UNRESOLVED_TEMPLATE_PLACEHOLDER"]
    assert not ph_issues


def test_toc_visual_validation_not_executed():
    assert TOC_VISUAL_VALIDATION == "NOT_EXECUTED"
    gate = assess_placeholders_for_mode([], mode="FINAL")
    assert gate["toc_visual_validation"] == "NOT_EXECUTED"
