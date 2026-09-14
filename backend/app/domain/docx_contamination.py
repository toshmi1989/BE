"""Phase 29.2 — Clear product-specific template contamination from rendered DOCX.

Primary protection: semantic registry (`template_contamination`).
Secondary: token scrub (`docx_stale_scrub`).
Tertiary: this post-render scan.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from lxml import etree

from app.domain.protocol_template_registry import study_is_bosutinib_sample
from app.domain.template_contamination import (
    CLEARED_PLACEHOLDER_RU,
    fingerprint_match,
    is_product_specific_text,
    scan_text_blob_for_contamination,
    text_contamination_hits,
)


def _paragraph_set_text(paragraph: Any, text: str) -> None:
    for r in list(paragraph.runs):
        parent = r._element.getparent()
        if parent is not None:
            parent.remove(r._element)
    paragraph.add_run(text)


def _iter_header_footer_parts(section: Any) -> list[Any]:
    parts: list[Any] = []
    for attr in (
        "header",
        "footer",
        "first_page_header",
        "first_page_footer",
        "even_page_header",
        "even_page_footer",
    ):
        part = getattr(section, attr, None)
        if part is not None:
            parts.append(part)
    return parts


def _walk_tables(tables: Any) -> Any:
    for table in tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _walk_tables(cell.tables)


def _collect_blob(doc: DocumentObject) -> str:
    chunks: list[str] = []
    for p in doc.paragraphs:
        chunks.append(p.text or "")
    for table in _walk_tables(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text or "")
    for section in doc.sections:
        for hf in _iter_header_footer_parts(section):
            for p in hf.paragraphs:
                chunks.append(p.text or "")
            for table in _walk_tables(hf.tables):
                for row in table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text or "")
    return "\n".join(chunks)


def _should_clear_paragraph_text(text: str) -> bool:
    if not text or not text.strip():
        return False
    if fingerprint_match(text):
        return True
    if is_product_specific_text(text):
        return True
    return False


def clear_product_specific_contamination(
    doc: DocumentObject,
    study_ctx: dict[str, Any],
    *,
    placeholder: str = CLEARED_PLACEHOLDER_RU,
) -> dict[str, Any]:
    """Remove Bosutinib-example semantic content for non-bosutinib studies.

    Does NOT invent Upadacitinib pharmacology — clears or blocks.
    """
    product = study_ctx.get("product") or {}
    if study_is_bosutinib_sample(product):
        return {"cleared": 0, "skipped": "study_is_bosutinib_sample"}

    cleared = 0
    cleared_locations: list[str] = []

    def handle(p: Any, loc: str) -> None:
        nonlocal cleared
        text = p.text or ""
        if not _should_clear_paragraph_text(text):
            return
        # Headings that only name the section (e.g. "2.8 ... бозутиниба") — rewrite generically
        style = (p.style.name if p.style is not None else "") or ""
        if style.startswith("Heading") or style.startswith("Заголовок"):
            new = re_sub_product_heading(text)
            _paragraph_set_text(p, new)
        else:
            _paragraph_set_text(p, placeholder)
        cleared += 1
        cleared_locations.append(loc)

    for i, p in enumerate(doc.paragraphs):
        handle(p, f"body.p[{i}]")

    for ti, table in enumerate(_walk_tables(doc.tables)):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    handle(p, f"body.table[{ti}].r{ri}.c{ci}.p[{pi}]")

    for si, section in enumerate(doc.sections):
        for hf in _iter_header_footer_parts(section):
            for pi, p in enumerate(hf.paragraphs):
                handle(p, f"section[{si}].hf.p[{pi}]")
            for ti, table in enumerate(_walk_tables(hf.tables)):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, p in enumerate(cell.paragraphs):
                            # Prefer identity line rewrite for header product cell
                            text = p.text or ""
                            if "Исследуемый препарат" in text and _should_clear_paragraph_text(text):
                                trade = str(
                                    (study_ctx.get("product") or {}).get("trade_name")
                                    or (study_ctx.get("product") or {}).get("inn")
                                    or ""
                                ).strip()
                                if trade:
                                    _paragraph_set_text(p, f"Исследуемый препарат: {trade}")
                                    cleared += 1
                                    cleared_locations.append(
                                        f"section[{si}].hf.table[{ti}].r{ri}.c{ci}.p[{pi}]"
                                    )
                                    continue
                            handle(p, f"section[{si}].hf.table[{ti}].r{ri}.c{ci}.p[{pi}]")

    return {
        "cleared": cleared,
        "locations": cleared_locations[:80],
        "placeholder_used": placeholder[:60],
    }


def re_sub_product_heading(text: str) -> str:
    """Neutralize product name inside section headings without inventing clinical text."""
    out = text
    replacements = (
        ("бозутиниба", "исследуемого препарата"),
        ("Бозутиниба", "исследуемого препарата"),
        ("бозутиниб", "исследуемый препарат"),
        ("Бозутиниб", "Исследуемый препарат"),
        ("Бозулиф", "референтный препарат"),
        ("бозулиф", "референтный препарат"),
        ("Bosutinib", "investigational product"),
        ("Bosulif", "reference product"),
    )
    for a, b in replacements:
        out = out.replace(a, b)
    return out


def scan_docx_contamination(
    path: Path | str,
    study_ctx: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Tertiary scan of a saved DOCX for semantic contamination."""
    study_ctx = study_ctx or {}
    product = study_ctx.get("product") or {}
    if study_is_bosutinib_sample(product):
        return {"contaminated": False, "skipped": "study_is_bosutinib_sample"}

    path = Path(path)
    doc = Document(str(path))
    blob = _collect_blob(doc)
    visible = scan_text_blob_for_contamination(blob)

    # XML-level (comments / split runs)
    xml_hits: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(z.read(name))
            part_blob = "".join((t.text or "") for t in root.xpath("//*[local-name()='t']"))
            part_scan = scan_text_blob_for_contamination(part_blob)
            if part_scan["contaminated"]:
                xml_hits.append({"part": name, **part_scan})

    contaminated = visible["contaminated"] or bool(xml_hits)
    # Cleared placeholder is allowed
    return {
        "contaminated": contaminated,
        "visible": visible,
        "xml_parts_contaminated": xml_hits,
        "fail_categories": sorted(
            {
                *(visible.get("by_category") or {}).keys(),
                *[c for part in xml_hits for c in (part.get("by_category") or {})],
            }
        ),
    }


def scrub_contamination_xml_package(path: Path | str, study_ctx: dict[str, Any]) -> dict[str, Any]:
    """Secondary XML-level clear for comments / split runs not visible to python-docx."""
    from app.domain.template_contamination import iter_contamination_markers

    product = study_ctx.get("product") or {}
    if study_is_bosutinib_sample(product):
        return {"scrubbed_nodes": 0, "skipped": "study_is_bosutinib_sample"}

    path = Path(path)
    replacement = str(
        (product.get("trade_name") or product.get("inn") or "")
    ).strip() or "«исследуемый препарат»"

    # Prefer emptying chemistry / disease markers; replace product names with study product.
    name_tokens = set(
        __import__(
            "app.domain.template_contamination", fromlist=["CONTAMINATION_MARKER_GROUPS"]
        ).CONTAMINATION_MARKER_GROUPS["PRODUCT_NAME"]
    )
    other_tokens = [t for t in iter_contamination_markers() if t not in name_tokens]

    scrubbed = 0
    tmp = path.with_suffix(".scrubxml.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                try:
                    root = etree.fromstring(data)
                except etree.XMLSyntaxError:
                    zout.writestr(item, data)
                    continue
                changed = False
                for node in root.xpath("//*[local-name()='t']"):
                    text = node.text or ""
                    if not text:
                        continue
                    new = text
                    for tok in sorted(name_tokens, key=len, reverse=True):
                        if tok in new:
                            new = new.replace(tok, replacement)
                    for tok in sorted(other_tokens, key=len, reverse=True):
                        if tok in new:
                            new = new.replace(tok, "")
                    if new != text:
                        node.text = new
                        scrubbed += 1
                        changed = True
                if changed:
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            zout.writestr(item, data)
    tmp.replace(path)
    return {"scrubbed_nodes": scrubbed, "replacement": replacement}


def assert_docx_clean_or_raise(path: Path | str, study_ctx: dict[str, Any]) -> dict[str, Any]:
    from app.domain.exceptions import ValidationError

    scan = scan_docx_contamination(path, study_ctx)
    if scan.get("contaminated"):
        raise ValidationError(
            "Template contains product-specific content that is not supported by the current study.",
            field="CRITICAL_TEMPLATE_CONTAMINATION",
            details={
                "action": "Resolve template/content mapping before DOCX generation.",
                "scan": scan,
            },
        )
    return scan
